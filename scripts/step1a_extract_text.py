"""第 1 步 a：把 02_LiS_electrolyte 的正文 + SI 抽成纯文本缓存（可续跑，已有的跳过）。

输出：output/text_cache/<key>.main.txt、<key>.si.txt；output/text_cache/_index.csv
用法：D:\\anaconda\\envs\\papertrail\\python.exe step1a_extract_text.py
"""
import csv
import io
import os
import re
import sys
from concurrent.futures import ProcessPoolExecutor

LIB = r"D:\Desktop\锂电催化\机器学习\skills\Open-Papertrail\Open-Papertrail\corpus\library"
SUB = os.path.join(LIB, "02_LiS_electrolyte")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, "output", "text_cache")
LONG_PREFIX = "\\\\?\\"  # Windows 长路径前缀 \\?\


def safe_key(doi):
    return re.sub(r"[^\w.-]+", "_", doi)


def pdf_text(raw):
    import pymupdf
    out = []
    with pymupdf.open(stream=raw, filetype="pdf") as doc:
        for i, page in enumerate(doc):
            out.append(f"\n<<PAGE {i + 1}>>\n")
            out.append(page.get_text("text"))
    return "".join(out)


def docx_text(raw):
    import docx
    d = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def doc_text(raw):
    # 老式 Word 97 .doc：没有转换工具，粗暴地从二进制里捞 UTF-16 与单字节可打印串
    runs = re.findall(r"(?:[\x20-\x7e\xb0-\xff]\x00){4,}", raw.decode("latin-1"))
    txt16 = "".join(r.encode("latin-1").decode("utf-16le", "ignore") + "\n" for r in runs)
    runs8 = re.findall(rb"[\x20-\x7e\r\n\t\xb5]{6,}", raw)
    txt8 = "\n".join(r.decode("latin-1") for r in runs8)
    return txt16 if len(txt16) > len(txt8) else txt8


def any_text(path):
    # 部分 SI 完整路径超过 Windows 260 字符上限，解析库打不开：Python 带长路径前缀读字节，再以内存流交给解析库
    with open(LONG_PREFIX + os.path.abspath(path), "rb") as fh:
        raw = fh.read()
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return pdf_text(raw)
    if ext == ".docx":
        return docx_text(raw)
    if ext == ".doc":
        return doc_text(raw)
    return ""


_LISTING = None


def si_files(row):
    # 以磁盘上"正文名_SI*"为准，索引里的 si_file 兜底
    global _LISTING
    if _LISTING is None:
        _LISTING = os.listdir(SUB)
    stem = os.path.splitext(row["pdf_file"])[0] + "_SI"
    found = sorted(f for f in _LISTING if f.startswith(stem))
    return found or [f.strip() for f in (row["si_file"] or "").split(";") if f.strip()]


def work(row):
    key = safe_key(row["doi"])
    res = {"key": key, "doi": row["doi"], "main_chars": 0, "si_chars": 0, "error": ""}
    for kind, files in (("main", [row["pdf_file"]]), ("si", si_files(row))):
        out = os.path.join(CACHE, f"{key}.{kind}.txt")
        if os.path.exists(out):
            res[f"{kind}_chars"] = os.path.getsize(out)
            continue
        if not files:
            continue
        chunks = []
        for f in files:
            try:
                chunks.append(any_text(os.path.join(SUB, f)))
            except Exception as ex:  # 坏文件记下来，不中断
                res["error"] += f"{kind}:{f}:{ex!r}; "
        txt = "\n".join(chunks)
        if not txt.strip():  # 全部失败就不写缓存，下次重试
            continue
        with open(out, "w", encoding="utf-8") as fh:
            fh.write(txt)
        res[f"{kind}_chars"] = len(txt)
    return res


def main():
    os.makedirs(CACHE, exist_ok=True)
    with open(os.path.join(LIB, "library_index.csv"), encoding="utf-8-sig") as fh:
        rows = [r for r in csv.DictReader(fh) if r["topic"] == "electrolyte"]
    print(f"{len(rows)} papers", flush=True)
    results = []
    with ProcessPoolExecutor(max_workers=10) as ex:
        for i, r in enumerate(ex.map(work, rows, chunksize=4), 1):
            results.append(r)
            if i % 100 == 0:
                print(f"{i}/{len(rows)}", flush=True)
    with open(os.path.join(CACHE, "_index.csv"), "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(results[0].keys()))
        w.writeheader()
        w.writerows(results)
    bad = [r for r in results if r["error"] or r["main_chars"] < 5000]
    print(f"done; {len(bad)} with error or main<5000 chars; with SI text: {sum(r['si_chars'] > 0 for r in results)}", flush=True)
    for r in bad:
        print(r, flush=True)


if __name__ == "__main__":
    sys.exit(main())
