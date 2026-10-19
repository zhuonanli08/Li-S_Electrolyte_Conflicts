"""第 1 步 d：由 step1_rows.pkl 生成体检汇总，写 output/第1步_体检表.xlsx（汇总 / 逐篇 / 说明 三页）。"""
import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "output")


def pct(n, d):
    return f"{n} ({n / d:.0%})" if d else "0"


def derive(df):
    df = df.copy()
    ok = lambda c: df[c] == "✅"  # noqa: E731
    df["E/S直接"] = ok("E/S")
    df["E/S可算"] = ok("E/S") | (ok("电解液用量µL") & ok("硫载量"))
    df["锂量可知"] = ok("锂厚") | ok("N/P") | ok("锂过量%") | (df["负极"] == "无负极")
    df["非室温测试"] = ok("低温测试(粗)") | ok("高温测试(粗)")
    es = df["E/S最小"].where(df["E/S直接"])
    df["E/S档"] = pd.cut(es, [0, 5, 10, 1e9], labels=["贫液(≤5)", "中等(5–10]", "富液(>10)"]).astype(object).where(lambda x: x.notna(), "未报告").astype(str)
    df.loc[(df["E/S档"] == "未报告") & ok("电解液用量µL"), "E/S档"] = "只给用量"
    li = df["锂厚最小µm"].where(ok("锂厚"))
    np_ = df["N/P最小"].where(ok("N/P"))
    thin = (li <= 50) | (np_ <= 3) | (df["负极"] == "无负极")
    df["锂档"] = pd.cut(li, [0, 50, 100, 250, 1e9], labels=["薄(≤50µm)", "较薄(50–100]", "中(100–250]", "厚(>250)"]).astype(object).where(lambda x: x.notna(), "未报告").astype(str)
    df.loc[(df["锂档"] == "未报告") & ok("N/P"), "锂档"] = "只给N/P"
    df.loc[df["负极"] == "无负极", "锂档"] = "无负极"
    df["负极受限(薄锂/N/P≤3/无负极)"] = thin
    ld = df["硫载量最大"].where(ok("硫载量"))
    df["载量档"] = pd.cut(ld, [0, 2, 4, 1e9], labels=["低(<2)", "中[2–4)", "高(≥4)"], right=False).astype(object).where(lambda x: x.notna(), "未报告").astype(str)
    df["贫液实验"] = es <= 5
    df["高载量实验"] = ld >= 4
    df["主层"] = df["有电池测试"] & (df["正极层"] == "硫碳") & (df["负极"] == "锂金属") & ~df["凝胶/固态嫌疑"]
    df["时段"] = pd.cut(df["年份"], [2015, 2019, 2022, 2026], labels=["2016–2019", "2020–2022", "2023–2026"]).astype(str)
    return df


def tables(df):
    T = []
    cell = df[df["有电池测试"]]
    main = df[df["主层"]]
    N = len(df)

    # 1 范围
    rows = [("全部论文", N, ""), ("有电池测试", len(cell), "其余多为纯计算 / 表征 / 液流"),
            ("主层：硫碳正极 + 锂金属负极 + 液态", len(main), "分析三的主战场")]
    for k in ["SPAN", "Li2S", "溶液正极(catholyte)", "硫碳+SPAN"]:
        rows.append((f"  其他正极层：{k}", int(((cell["正极层"] == k)).sum()), "单独成层，不和硫碳混比"))
    for k in ["石墨/硅(锂离子-硫)", "无负极"]:
        rows.append((f"  负极：{k}", int((cell["负极"] == k).sum()), ""))
    rows.append(("  题名像凝胶/固态（有电池测试中）", int(cell["凝胶/固态嫌疑"].sum()), "按方案不收，主层已剔除"))
    T.append(("表1 范围拆分", pd.DataFrame(rows, columns=["层", "篇数", "备注"])))

    # 2 报告率
    items = [("E/S 直接给出", "E/S直接"), ("E/S 可算（直接给，或给了用量+载量）", "E/S可算"),
             ("锂量可知（锂厚 / N/P / 锂过量 / 无负极）", "锂量可知"), ("  其中锂厚", None), ("  其中 N/P", None),
             ("硫载量", None), ("LiNO3 浓度", None), ("非室温测试（粗）", "非室温测试"), ("软包", None),
             ("Li||Cu 库仑效率", None), ("多硫化物溶解度测定", None), ("溶剂化表征（Raman/MD）", None),
             ("常规醚基准 1M LiTFSI DOL/DME", None)]
    colmap = {"  其中锂厚": "锂厚", "  其中 N/P": "N/P", "硫载量": "硫载量", "LiNO3 浓度": "LiNO3浓度", "软包": "软包",
              "Li||Cu 库仑效率": "Li||Cu", "多硫化物溶解度测定": "多硫化物溶解度测定", "溶剂化表征（Raman/MD）": "溶剂化表征(Raman/MD)",
              "常规醚基准 1M LiTFSI DOL/DME": "常规醚基准(1M LiTFSI DOL/DME)"}
    rows = []
    for name, col in items:
        f = (lambda d, c=col: d[c]) if col else (lambda d, c=colmap[name]: d[c] == "✅")
        rows.append((name, pct(int(f(main).sum()), len(main)), pct(int(f(cell).sum()), len(cell)),
                     *[pct(int(f(main[main["时段"] == p]).sum()), int((main["时段"] == p).sum())) for p in ["2016–2019", "2020–2022", "2023–2026"]]))
    T.append(("表2 字段报告率（✅ = 有本文自述证据）", pd.DataFrame(rows, columns=["字段", f"主层 n={len(main)}", f"有电池测试 n={len(cell)}",
                                                                   "主层 2016–2019", "主层 2020–2022", "主层 2023–2026"])))

    # 3 条件分布（主层）
    for title, col, order in [("表3a 主层 E/S 分布（每篇取最小值）", "E/S档", ["贫液(≤5)", "中等(5–10]", "富液(>10)", "只给用量", "未报告"]),
                              ("表3b 主层 锂量分布（每篇取最薄）", "锂档", ["薄(≤50µm)", "较薄(50–100]", "中(100–250]", "厚(>250)", "只给N/P", "未报告"]),
                              ("表3c 主层 硫载量分布（每篇取最高）", "载量档", ["高(≥4)", "中[2–4)", "低(<2)", "未报告"])]:
        vc = main[col].value_counts()
        T.append((title, pd.DataFrame([(k, pct(int(vc.get(k, 0)), len(main))) for k in order], columns=["档", "篇数"])))

    # 4 E/S × 锂 交叉（主层）
    ct = pd.crosstab(main["E/S档"], main["锂档"]).reindex(index=["贫液(≤5)", "中等(5–10]", "富液(>10)", "只给用量", "未报告"],
                                                         columns=["薄(≤50µm)", "较薄(50–100]", "中(100–250]", "厚(>250)", "只给N/P", "未报告"], fill_value=0)
    ct = ct.reset_index().rename(columns={"E/S档": "E/S \\ 锂"})
    T.append(("表4 主层 E/S × 锂量 交叉（篇）——贫液×薄锂那一格是分析三最缺的", ct))

    # 5 类别 × 条件
    fams = ["强溶剂化", "适中溶剂化", "弱溶剂化", "难溶/固相转化", "局部高浓", "高浓", "低浓", "介体/均相催化", "LiNO3", "盐/阴离子", "氟化/无氟", "离子液体"]
    rows = []
    for f in fams:
        has = df["电解液类别(题名+摘要)"].str.split("; ").apply(lambda xs, f=f: f in xs)
        d, m = df[has], df[has & df["主层"]]
        rows.append((f, len(d), len(m), int(m["贫液实验"].sum()), int(m["负极受限(薄锂/N/P≤3/无负极)"].sum()),
                     int((m["贫液实验"] & m["负极受限(薄锂/N/P≤3/无负极)"]).sum()), int(m["高载量实验"].sum()),
                     int((m["软包"] == "✅").sum()), int((m["低温测试(粗)"] == "✅").sum()), int(((df["正极层"] == "SPAN") & has).sum())))
    T.append(("表5 电解液类别 × 条件（类别按题名+摘要粗标、可多标签；条件列只数主层）",
              pd.DataFrame(rows, columns=["类别", "全部", "主层", "有贫液实验", "有负极受限实验", "贫液且负极受限", "有高载量(≥4)", "有软包", "有低温", "SPAN层"])))
    return T


def write_xlsx(df, T):
    path = os.path.join(OUT, "第1步_体检表.xlsx")
    order = ["doi", "年份", "期刊", "题名", "主层", "有电池测试", "正极层", "负极", "凝胶/固态嫌疑", "电解液类别(题名+摘要)", "条件标签(题名+摘要)",
             "E/S", "E/S最小", "E/S值", "E/S原句", "电解液用量µL", "用量值", "用量原句",
             "锂厚", "锂厚最小µm", "锂厚值", "锂厚原句", "N/P", "N/P最小", "N/P值", "N/P原句", "锂过量%", "无负极提及",
             "硫载量", "硫载量最大", "硫载量值", "硫载量原句", "LiNO3提及", "LiNO3浓度", "LiNO3浓度值", "LiNO3原句", "无LiNO3对照",
             "低温测试(粗)", "低温值", "低温原句", "高温测试(粗)", "高温值", "高温原句", "软包", "软包Ah", "软包Ah值",
             "Li||Cu", "多硫化物溶解度测定", "溶剂化表征(Raman/MD)", "常规醚基准(1M LiTFSI DOL/DME)", "基准原句",
             "E/S档", "锂档", "载量档", "负极受限(薄锂/N/P≤3/无负极)", "正极计数(SPAN/Li2S/catholyte/硫碳)", "正文字数", "有SI文本", "摘要来源"]
    per = df[order].sort_values(["主层", "年份"], ascending=[False, True])
    # PDF 文本里混有 Excel 不接受的控制字符
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    per = per.map(lambda v: ILLEGAL_CHARACTERS_RE.sub(" ", v) if isinstance(v, str) else v)
    notes = pd.DataFrame([
        ("✅", "找到本文自述证据：命中在实验部分或 SI，或正文句子带 we/our/Fig/cells were… 等自述标志"),
        ("🔶", "只有泛泛提及或来源存疑：正文里的'应当/实用要求/文献报道'、SI 文献对比表、能量密度测算参数"),
        ("❌", "全文（含 SI）没找到"),
        ("值的取法", "E/S、锂厚、N/P 每篇取最小（有没有做过苛刻条件）；载量取最大；同篇多条件见'值'列"),
        ("主层", "有电池测试 + 硫碳正极 + 锂金属负极 + 题名不像凝胶/固态"),
        ("负极受限", "锂厚 ≤50 µm、或 N/P ≤3、或无负极（阈值是方案里的建议起点，第 3 步按数据再定）"),
        ("电解液类别", "按题名+摘要关键词粗标、可多标签；摘要里拿来作对比的类别也会被标上，第 2 步论断抽取才是正式判定"),
        ("温度", "粗判：负温度、≥45 °C 且上下文是循环/测试，排除干燥、合成、熔点等"),
        ("正则验证", "见 output/第1步_体检结果.md 的'正则验证'一节"),
        ("原句前缀", "[EXP] 实验部分；[SI] 支撑信息；[MAIN] 其余正文"),
    ], columns=["项", "说明"])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        r = 0
        for title, t in T:
            pd.DataFrame([[title]]).to_excel(xw, sheet_name="汇总", startrow=r, index=False, header=False)
            t.to_excel(xw, sheet_name="汇总", startrow=r + 1, index=False)
            r += len(t) + 4
        per.to_excel(xw, sheet_name="逐篇", index=False)
        notes.to_excel(xw, sheet_name="说明", index=False)
        ws = xw.sheets["汇总"]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("表"):
                    c.font = Font(bold=True, size=12)
        ws.column_dimensions["A"].width = 44
        for col in "BCDEFGHIJK":
            ws.column_dimensions[col].width = 16
        ws = xw.sheets["逐篇"]
        ws.freeze_panes = "E2"
        hdr = PatternFill("solid", fgColor="DDEBF7")
        for i, name in enumerate(per.columns, 1):
            ws.cell(1, i).font = Font(bold=True)
            ws.cell(1, i).fill = hdr
            ws.cell(1, i).alignment = Alignment(wrap_text=True, vertical="top")
            w = 60 if "原句" in name else (50 if name == "题名" else (24 if name in ("doi", "电解液类别(题名+摘要)", "期刊") else 11))
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str) and len(c.value) > 60:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        ws = xw.sheets["说明"]
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 110
    return path


def main():
    df = derive(pd.read_pickle(os.path.join(OUT, "step1_rows.pkl")))
    df.to_pickle(os.path.join(OUT, "step1_rows_derived.pkl"))
    T = tables(df)
    print(write_xlsx(df, T))
    for title, t in T:
        print("\n" + title)
        print(t.to_string(index=False))


if __name__ == "__main__":
    main()
