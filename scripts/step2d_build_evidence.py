"""方案2 §5.4–5.7：把 output/第一阶段_证据表/json/*.json（每篇一份，Claude 精读填写）拼成证据表 Excel。

四页：主表（含复核列）、矛盾对、汇总、新旧对照。计算口径见方案2 §5.7（公式沿用毕设 v221"性能计算"页）。
10-18 第二轮评审修正：不可比描述符只展示、不参与方向和区间判定；删除按策略标签推断强弱；
作者依据须点名本行两种电解液；翻转只指强弱胜方互换；具体指标不同标③，取值点不同另标；
质量基 E/S 档位标来源（密度区间推定）并做剔除后敏感性分析；随机抽查固定沿用基线 9 条。
用法：D:\\anaconda\\envs\\papertrail\\python.exe -X utf8 -I step2d_build_evidence.py
"""
import glob
import json
import math
import os
import random
import re
import sys
import unicodedata
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "output", "第一阶段_证据表", "json")
CACHE = os.path.join(ROOT, "output", "text_cache")
OUT = os.path.join(ROOT, "output", "第一阶段_证据表.xlsx")
# 新旧对照的基线：v0 = 第二轮规则修正前；v1 = 第二轮修正后、用户第三轮复核时的版本（当前对照用 v1）
BASELINE = os.path.join(ROOT, "output", "第一阶段_证据表", "基线_v1_判定.json")
SEED = 20261018
RHO = (0.8, 2.0)  # 质量基 E/S 换算时假定的电解液密度范围 g/mL

# v221 默认值（只进"实际计算输入"，并打默认标记）
DEF_D_MM, DEF_V_UL, DEF_LI_UM = 12.0, 50.0, 1000.0
PI = 3.14  # v221 取 3.14
METRIC_GROUP = {"首圈容量": "容量与倍率", "倍率容量": "容量与倍率", "平均放电容量": "容量与倍率", "能量密度": "容量与倍率",
                "循环保持率或寿命": "循环", "指定圈数放电容量": "循环", "每圈衰减率": "循环", "全电池库仑效率": "循环",
                "自放电（容量变化）": "循环", "自放电（容量损失）": "循环",
                "Li‖Cu 库仑效率": "负极侧", "锂对称电池寿命": "负极侧",
                "过电位或动力学": "其他", "阻抗": "其他", "多硫化物溶解或穿梭": "其他", "其他": "其他"}
LOWER_BETTER = ("每圈衰减率", "阻抗", "自放电（容量损失）", "多硫化物溶解或穿梭")  # 穿梭类指标（穿梭电流、负极含硫量等）越低越好
# 自放电：容量变化 = 静置后相对静置前的变化（正为增加，越高越好）；容量损失 = 损失的比例（越低越好）
LABEL_ORDER = ["强溶剂化", "常规", "适中溶剂化", "弱溶剂化", "局部高浓", "高浓", "难溶", "其他"]  # 只用于显示排序，不参与强弱判定
TOL = {"E/S": (0.67, 1.5), "载量": (0.67, 1.5), "倍率": (0.5, 2.0), "温度": 5.0, "LiNO3": (0.5, 2.0)}


# ---------- 原文核对 ----------
def nz(s):
    s = unicodedata.normalize("NFKC", s or "").lower().replace("µ", "u").replace("μ", "u")
    return "".join(ch for ch in s if ch.isascii() and ch.isalnum())


_DOC = {}


def doc(doi):
    if doi not in _DOC:
        key = re.sub(r"[^\w.-]+", "_", doi)
        t = ""
        for part in ("main", "si"):
            p = os.path.join(CACHE, f"{key}.{part}.txt")
            if os.path.exists(p):
                t += open(p, encoding="utf-8", errors="ignore").read().replace("\x00", "") + "\n"
        _DOC[doi] = nz(t)
    return _DOC[doi]


def found(doi, quote):
    """引文按 … 分段，每段都要在原文（正文 + SI）里找到"""
    if not quote:
        return True
    segs = [nz(x) for x in re.split(r"…", quote)]
    return all(s in doc(doi) for s in segs if len(s) >= 8)


# ---------- 数值小工具 ----------
def num(v):
    """返回 (数值, 限定符)；'<243' → (243, '<')"""
    if v is None:
        return None, ""
    if isinstance(v, (int, float)):
        return float(v), ""
    m = re.match(r"\s*([<>≤≥~约]?)\s*(-?\d+(?:\.\d+)?)\s*$", str(v))
    return (float(m.group(2)), m.group(1)) if m else (None, "")


def rid_key(s):
    """行号自然排序：S09-C2 排在 S09-C10 前"""
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s)]


def fmt(x, nd=2):
    if x is None:
        return ""
    if isinstance(x, float):
        return f"{x:.{nd}f}".rstrip("0").rstrip(".")
    return str(x)


def band_es(x):
    if x is None:
        return "未知"
    return "贫液(≤5)" if x <= 5 else ("中等(5–15]" if x <= 15 else "富液(>15)")


def band_load(x):
    if x is None:
        return "未知"
    return "<2" if x < 2 else ("2–4" if x < 4 else "≥4")


def band_rate(x):
    if x is None:
        return "未知"
    return "≤0.3C" if x <= 0.3 else ("0.3–2C" if x <= 2 else ">2C")


def band_temp(t):
    if t is None:
        return "未知"
    return "室温(20–30)" if 20 <= t <= 30 else ("低温" if t < 20 else "高温")


# ---------- 条件计算（方案2 §5.7） ----------
def cond_calc(c, row):
    r = {"notes": []}
    d = c.get("diameter_mm")
    if c.get("area_cm2"):
        area, area_def = c["area_cm2"], False
    elif d:
        area, area_def = PI / 4 * (d / 10) ** 2, False
    else:
        area, area_def = PI / 4 * (DEF_D_MM / 10) ** 2, True
    r["area"], r["area_def"] = area, area_def
    load = c.get("S_loading")
    V = c.get("V_uL")
    r["V_input"], r["V_def"] = (V, False) if V is not None else (DEF_V_UL, True)
    # E/S
    es_rep, unit = c.get("ES_reported"), c.get("ES_unit", "µL/mg")
    es_calc = es_calc_def = None
    if load:
        m_s = load * r["area"]
        val = r["V_input"] / m_s
        if r["V_def"] or r["area_def"]:
            es_calc_def = val
        else:
            es_calc = val
    r["es_rep"], r["es_unit"], r["es_calc"], r["es_calc_def"] = es_rep, unit, es_calc, es_calc_def
    r["es_band_note"] = ""
    r["es_gg"] = r["es_conv"] = None
    if es_rep is not None and unit == "µL/mg":
        r["es_used"], r["es_src"], r["es_band_src"] = es_rep, "直接", "直接报告"
        if es_calc is not None and abs(es_calc - es_rep) / es_rep > 0.10:
            r["notes"].append(f"⚠️ E/S 报告 {es_rep} 与按体积/载量/面积算出的 {es_calc:.2f} 相差 >10%")
        r["es_band"] = band_es(es_rep)
    elif es_rep is not None and unit == "g/g" and c.get("density_g_mL"):
        rho = c["density_g_mL"]
        r["es_gg"], r["es_used"] = es_rep, es_rep / rho
        r["es_src"], r["es_band_src"] = f"可靠换算（ρ {rho} g/mL，{c.get('density_src') or '来源未写'}）", "可靠换算"
        r["es_conv"] = f"{es_rep} g/g ÷ {rho} g/mL = {es_rep / rho:.2f} µL/mg"
        r["es_band"] = band_es(r["es_used"])
    elif es_rep is not None and unit == "g/g":
        r["es_gg"], r["es_used"] = es_rep, None  # 不进入连续值匹配
        r["es_src"] = "直接（质量基 g/g，无可靠密度）"
        lo, hi = es_rep / RHO[1], es_rep / RHO[0]
        r["es_conv"] = f"假定 ρ {RHO[0]}–{RHO[1]} g/mL → {lo:.2f}–{hi:.2f} µL/mg"
        if hi <= 5 or lo > 15 or (lo > 5 and hi <= 15):
            r["es_band"], r["es_band_src"] = band_es(lo), "密度区间推定"
            r["es_band_note"] = f"{r['es_conv']}；在假定的 {RHO[0]}–{RHO[1]} g/mL 范围内档位不变"
        else:
            r["es_band"], r["es_band_src"] = "未知", "未知（换算区间跨档）"
            r["es_band_note"] = f"{r['es_conv']}，跨档"
    elif es_calc is not None:
        r["es_used"], r["es_src"], r["es_band"], r["es_band_src"] = es_calc, "计算", band_es(es_calc), "计算"
    elif es_calc_def is not None:
        r["es_used"], r["es_src"], r["es_band"], r["es_band_src"] = None, "计算（含默认输入），不进分析", "未知", "未知（含默认输入）"
    else:
        r["es_used"], r["es_src"], r["es_band"], r["es_band_src"] = None, "缺失", "未知", "未知（缺失）"
    # 锂
    li = c.get("Li_um")
    r["li_rep"] = li
    r["li_input"], r["li_def"] = (li, False) if li is not None else (DEF_LI_UM, True)
    r["li_cap"] = 0.2 * r["li_input"]
    fa, fb = row.get("first_cap_A"), row.get("first_cap_B")
    r["np_b"] = r["np_a"] = None
    if load and fb:
        r["np_b"] = r["li_cap"] / (fb * load * 0.001)
    if load and fa:
        r["np_a"] = r["li_cap"] / (fa * load * 0.001)
    r["np_def"] = r["li_def"]
    r["design"] = (r["li_cap"] / (load * 1.675)) if load else None
    if r["np_a"] and r["np_b"] and not r["np_def"] and ((r["np_a"] - 3) * (r["np_b"] - 3) < 0):
        r["notes"].append("🔶 N/P_A 与 N/P（按 B）落在 3 的两侧")
    # 瓶颈层（只用不含默认输入的值）
    np_author = c.get("NP_author")
    ctype = c.get("cathode_type", "")
    if ctype.startswith("无"):
        r["layer"] = r["layer_1012"] = "不适用"
        r["li_band"] = "不适用"
    else:
        anode = None  # True / False / None(未知)
        if li is not None and li <= 50:
            anode = True
        elif r["np_b"] is not None and not r["np_def"] and r["np_b"] <= 3:
            anode = True
        elif np_author is not None and np_author <= 3:
            anode = True
        elif li is not None:
            if r["np_b"] is not None and not r["np_def"]:
                anode = False
            elif r["design"] is not None and r["design"] > 3:
                anode = False  # 实测 N/P ≥ 设计比值
            elif np_author is not None:
                anode = np_author <= 3
            else:
                anode = None
        lean_hl = (load is not None and load >= 4 and r["es_band"] == "贫液(≤5)")
        if anode is True:
            r["layer"] = "负极受限"
        elif anode is False:
            r["layer"] = "正极受限" if lean_hl else "常规"
        elif li is None and np_author is None:
            r["layer"] = "锂库存未知"
        else:
            r["layer"] = "待定（锂厚已知、N/P 定不了）"
        if r["layer"] == "锂库存未知":
            r["layer_1012"] = "正极受限" if lean_hl else "常规"
        else:
            r["layer_1012"] = r["layer"]
        npk = r["np_b"] if (r["np_b"] is not None and not r["np_def"]) else np_author
        if li is None and np_author is None:
            r["li_band"] = "未报告"
        elif (li is not None and li <= 50) or (npk is not None and npk <= 3):
            r["li_band"] = "薄"
        elif li is not None and li <= 200:
            r["li_band"] = "中"
        elif li is not None:
            r["li_band"] = "厚"
        else:
            r["li_band"] = "未报告"
    # 倍率
    rate = row.get("rate_override") if row.get("rate_override") is not None else c.get("rate_author")
    d1c = c.get("rate_1C_def") or 1675
    r["rate_protocol"] = row.get("rate_protocol")
    if r["rate_protocol"]:  # 取值区间跨变倍率协议：不用单值倍率，不进匹配
        rate = None
    r["rate"], r["rate_def"] = rate, c.get("rate_1C_def")
    r["rate_u"] = rate * d1c / 1675 if rate is not None else None
    r["rate_band"] = "变倍率协议（不归档）" if r["rate_protocol"] else band_rate(r["rate_u"])
    r["temp_band"] = band_temp(c.get("temp_C"))
    r["load_band"] = band_load(load)
    return r


# ---------- 溶剂化依据 ----------
def cred(src, comparable, mapping="直接"):
    """体系内可比 ≠ 能映射到电池电解液：间接映射（纯溶剂 / 无盐体系 / 代理）的最高只到'中'"""
    if src == "实验测量":
        if not comparable:
            return "中"
        return "高" if mapping == "直接" else "中（体系内可比，映射到电池电解液是间接的）"
    if src == "模拟计算":
        return "中" if mapping == "直接" else "低（溶剂体系模拟，间接映射）"
    return "低"


SOLU_FIELDS = ("species", "medium", "method", "conc_def")


def common_key(d):
    """跨论文公共刻度：溶解度按 物种｜介质｜测法｜单位 自动生成（另需温差 ≤5 °C）；其他描述符用人工指定的 common_scale"""
    if d.get("kind") == "溶解度":
        if all(d.get(k) for k in SOLU_FIELDS) and "文献" not in d["method"] and "未写" not in d["method"]:
            return f"{d['species']}溶解度｜{d['medium']}｜{d['method']}｜{d['unit']}"
        return None
    return d.get("common_scale")


def solvation(P, row):
    """返回 依据条目、方向、方向来源、(公共刻度, 本篇刻度)、被排除的不可比描述符、参与判定的描述符"""
    A, B = row["A"], row["B"]
    entries, dirs, scales, allsc, excluded, used = [], [], {}, {}, [], []
    for d in P.get("descriptors", []):
        va, vb = d["values"].get(A), d["values"].get(B)
        if va is None or vb is None:
            continue
        comp = d["comparable"] and not ({A, B} & set(d.get("noncomparable_ids") or []))
        if va == vb:
            dr = "相同"
        else:
            stronger_a = (va > vb) if d["stronger_if"] == "higher" else (va < vb)
            dr = "A比B更强" if stronger_a else "A比B更弱"
        q = d.get("qual", {})
        mapping = d.get("mapping", "直接")
        c = cred(d["source"], comp, mapping)
        detail = "；".join(f"{lab}：{d[k]}" for lab, k in (("物种", "species"), ("介质", "medium"), ("测法", "method"), ("浓度定义", "conc_def")) if d.get(k))
        if d.get("kind"):
            detail += f"；温度：{fmt(d.get('temp_C')) + ' °C' if d.get('temp_C') is not None else (d.get('temp_note') or '未写')}"
        entries.append(f"【{d['id']}｜{d['object']}｜{d['source']}｜{d['scale']}】A={q.get(A, fmt(va))}，B={q.get(B, fmt(vb))} {d['unit']}；"
                       f"方向：{dr}；体系内可比：{'是' if comp else '否'}（{d['comparable_note']}）；"
                       f"映射到电池电解液：{mapping}{'（' + d['mapping_note'] + '）' if d.get('mapping_note') else ''}；可信度：{c}"
                       + (f"｜{detail}" if detail else "") + ("" if comp else "；⛔ 不可比，不参与方向和区间判定"))
        if not comp:
            if dr != "相同":
                excluded.append(d["id"])
            continue
        if dr != "相同":
            dirs.append(dr)
            used.append(d)
        ck = common_key(d)
        if ck:
            scales[ck] = (va, vb, d.get("temp_C"))
        allsc[f"本篇 {d['id']}（{d['scale']}）"] = (va, vb, d.get("temp_C"))
    if dirs:
        dirn = dirs[0] if len(set(dirs)) == 1 else "依据冲突"
        srcn = "描述符（可比）"
    elif row.get("solv_dir_author") and row.get("solv_author_type") != "类别称呼" and row.get("solv_quote"):
        dirn, srcn = row["solv_dir_author"], "作者明确依据（低可信）"
        entries.append(f"【作者明确依据】{row['solv_dir_author']}；原句点名本行两种电解液；可信度：低")
    else:
        if row.get("solv_dir_author"):
            entries.append(f"【作者类别称呼】{row['solv_dir_author']}；原句只是类别称呼（未点名本行两种电解液），不作判定依据")
        dirn, srcn = "无法判断", "无可比依据"
    return entries, dirn, srcn, (scales, allsc), excluded, used


def polarity(row):
    if row["metric"] == "过电位或动力学" and re.search(r"mV|Ω|ohm", row.get("unit", "")):
        return "lower"
    if row["metric"] in LOWER_BETTER:
        return "lower"
    return "higher"


def diff_text(row, a, b):
    """差值型指标：正 = A 更好；只有百分数相减才叫百分点"""
    d = (a - b) if polarity(row) == "higher" else (b - a)
    return f"差值 {d:+.1f} 个百分点" if row.get("unit") == "%" else f"差值 {d:+.1f} {row.get('unit') or ''}".rstrip()


def code_direction(row):
    """有两个直接数值时代码判方向；返回 (方向, 差异说明)"""
    (a, qa), (b, qb) = num(row.get("val_A")), num(row.get("val_B"))
    if a is None or b is None or not row.get("comparable_point", True):
        return None, ""
    hi = polarity(row) == "higher"
    if qa == "<" and a <= b:
        return ("A更差" if hi else "A更好"), "A 为上限值"
    if qb == "<" and b <= a:
        return ("A更好" if hi else "A更差"), "B 为上限值"
    if qa == ">" and a >= b:
        return ("A更好" if hi else "A更差"), "A 为下限值"
    if qb == ">" and b >= a:
        return ("A更差" if hi else "A更好"), "B 为下限值"
    approx = qa in ("~", "约") or qb in ("~", "约")
    if (qa or qb) and not approx:
        return None, "含限定符"
    if "库仑效率" in row["metric"] and row.get("unit") == "%":
        if a == b:
            return "差异小", "百分点差 0"
        return (("A更好" if a > b else "A更差") if hi else ("A更好" if a < b else "A更差")), f"百分点差 {a - b:+.2f}"
    if row["metric"].startswith("自放电"):
        better = (a > b) if hi else (a < b)
        return ("差异小" if a == b else ("A更好" if better else "A更差")), diff_text(row, a, b)
    base = max(abs(a), abs(b))
    if base and abs(a - b) / base < 0.05:
        return "差异小", "相对差 <5%" + ("（近似值）" if approx else "")
    better = (a > b) if hi else (a < b)
    return ("A更好" if better else "A更差"), ("近似值" if approx else "")


def effect(row):
    (a, qa), (b, qb) = num(row.get("val_A")), num(row.get("val_B"))
    srcs = (str(row.get("val_src_A", "")), str(row.get("val_src_B", "")))
    if a is None or b is None or qa or qb or not row.get("comparable_point", True):
        return ""
    if not all(s.startswith("直接") for s in srcs):
        return ""
    if row["metric"].startswith("自放电"):
        return diff_text(row, a, b)
    if "库仑效率" in row["metric"] and row.get("unit") == "%":
        out = f"百分点差 {a - b:+.2f}"
        dec = lambda v: len(str(v).split(".")[1]) if "." in str(v) else 0  # noqa: E731
        if a <= 99.9 and b <= 99.9 and dec(row["val_A"]) >= 2 and dec(row["val_B"]) >= 2 and a < 100 and b < 100:
            out += f"；ln 损失比 {math.log((100 - b) / (100 - a)):+.3f}"
        return out
    if a <= 0 or b <= 0:
        return ""
    return f"{math.log(a / b) if polarity(row) == 'higher' else math.log(b / a):+.3f}"


def winner(dirn, obs):
    if dirn == "依据冲突":
        return "依据冲突"
    if dirn not in ("A比B更强", "A比B更弱") or obs not in ("A更好", "A更差", "差异小"):
        return "无法判断"
    if obs == "差异小":
        return "差异小"
    weak_a = dirn == "A比B更弱"
    return "弱者胜" if (weak_a == (obs == "A更好")) else "强者胜"


def ctype_pair(P, row):
    la, lb = P["electrolytes"][row["A"]]["label"], P["electrolytes"][row["B"]]["label"]
    a, b = sorted([la, lb], key=LABEL_ORDER.index)
    return f"{a} ↔ {b}"


# ---------- 主表 ----------
def build_rows():
    rows = []
    for path in sorted(glob.glob(os.path.join(SRC, "*.json"))):
        P = json.load(open(path, encoding="utf-8"))
        pp = P["paper"]
        doi = pp["doi"]
        bad_el = {k for k, e in P["electrolytes"].items() if not found(doi, e.get("quote"))}
        bad_de = {d["id"] for d in P.get("descriptors", []) if not found(doi, d.get("quote"))}
        bad_co = {k for k, c in P["conditions"].items() if not found(doi, c.get("quote"))}
        for i, row in enumerate(P["comparisons"], 1):
            c = P["conditions"][row["cond"]]
            ea, eb = P["electrolytes"][row["A"]], P["electrolytes"][row["B"]]
            cc = cond_calc(c, row)
            entries, dirn, dsrc, (scales, allsc), excluded, used = solvation(P, row)
            cdir, cnote = code_direction(row)
            obs = row["obs_dir"]
            notes = list(cc["notes"])
            mark_warn = any(n.startswith("⚠️") for n in notes)
            if cdir and cdir != obs:
                notes.append(f"⚠️ 代码按数值判为'{cdir}'，填写为'{obs}'")
                mark_warn = True
            win = winner(dirn, obs)
            qok = found(doi, row.get("quote"))
            rqok = found(doi, row.get("reason_quote")) and found(doi, row.get("solv_quote"))
            flags = list(row.get("flags", []))
            if not rqok:
                notes.append("❌ 原因/方向原句在原文里找不到")
            used_de = {d["id"] for d in P.get("descriptors", []) if d["values"].get(row["A"]) is not None and d["values"].get(row["B"]) is not None}
            bad_row = [f"电解液 {k}" for k in sorted({row["A"], row["B"]} & bad_el)] + [f"描述符 {k}" for k in sorted(used_de & bad_de)] +                       ([f"条件 {row['cond']}"] if row["cond"] in bad_co else [])
            if bad_row:
                notes.append("❌ 以下原句在原文里找不到：" + "、".join(bad_row))
            if dirn == "依据冲突":
                flags.append("🔶 溶剂化依据冲突")
            if obs == "无法判断" or win == "无法判断":
                flags.append("🔶 方向或胜方无法判断")
            if row["metric"] == "平均放电容量" and not re.search(r"\d+\s*圈", row.get("point", "")):
                flags.append("🔶 平均放电容量未写平均的圈数范围")
            computed = cc["es_src"].startswith("计算") or cc["np_b"] is not None or cc["design"] is not None
            if not qok or bad_row or not rqok:
                mark = "❌"
            elif mark_warn:
                mark = "⚠️"
            elif flags or "❌" in "".join(notes):
                mark = "🔶"
            elif computed:
                mark = "🧮"
            else:
                mark = "✅"
            pid = pp["pool_id"]
            rid = f"{pid}-C{i}"
            mgroup = METRIC_GROUP.get(row["metric"], "其他")
            maps = [d.get("mapping", "直接") for d in used]
            if not used:
                bmap = ""
            elif all(m == "直接" for m in maps):
                bmap = "直接（依据测的就是电池电解液）"
            else:
                notes_m = list(dict.fromkeys(d.get("mapping_note", "") for d in used if d.get("mapping", "直接") != "直接"))
                bmap = ("直接 + 间接：" if "直接" in maps else "间接映射：") + "；".join(n for n in notes_m if n)
                if c.get("cathode_type", "").startswith("无"):
                    bmap += "；本行是无硫正极的对称 / 半电池测试"
            revise = "；".join(x for x in (row.get("revise"), c.get("revise")) if x)
            ch = row.get("changed", [])
            fa, fb = row.get("first_cap_A"), row.get("first_cap_B")
            rows.append({
                "_paper": pid, "_cond": row["cond"], "_A": row["A"], "_B": row["B"], "_metric": row["metric"], "_mgroup": mgroup,
                "_winner": win, "_ctype": ctype_pair(P, row), "_cc": cc, "_c": c, "_scales": scales, "_allsc": allsc, "_comparable_point": row.get("comparable_point", True),
                "_anions": {ea.get("salt_anion"), eb.get("salt_anion")}, "_lino3": (ea, eb), "_evtype": row.get("evidence_type"),
                "_main": pp.get("main_layer") and c.get("cathode_type") == "硫碳", "_point": row.get("point", ""),
                "_excluded": excluded, "_used": used, "_dsrc": dsrc, "_effect": effect(row), "_solv_quote": row.get("solv_quote"),
                "_revise": revise, "_bmap": bmap, "_stat": row.get("stat_basis"),
                "A1 行号": rid, "A2 论文号": pid, "A2 DOI": pp["doi"], "A3 文献": pp["ref"], "A4 通讯作者": pp["corr"] + (f"｜{pp['group_note']}" if pp.get("group_note") else ""),
                "A5 抽样来源": pp["source"], "A6 主层（本行）": "是" if (pp.get("main_layer") and c.get("cathode_type") == "硫碳") else f"否（正极：{c.get('cathode_type')}）",
                "B1 电解液A": f"{row['A']}：{ea['name']}", "B1 A配方": f"溶剂 {ea['solvent']}；盐 {ea['salt']}；稀释剂 {ea['diluent']}；LiNO₃ {ea['lino3']}；添加剂 {ea['additive']}",
                "B2 电解液B": f"{row['B']}：{eb['name']}", "B2 B配方": f"溶剂 {eb['solvent']}；盐 {eb['salt']}；稀释剂 {eb['diluent']}；LiNO₃ {eb['lino3']}；添加剂 {eb['additive']}",
                "B3 同时改变的因素": "、".join(ch), "B4 仅改变一个配方因素": "是" if len(ch) == 1 else "否",
                "B5 标签A（作者原词）": ea["label_author"], "B5 标签A": ea["label"], "B5 标签B（作者原词）": eb["label_author"], "B5 标签B": eb["label"],
                "B6 比较类型": ctype_pair(P, row),
                "B7 溶剂化依据条目": "\n".join(entries), "B9 溶剂化方向": dirn, "B9 方向来源": dsrc,
                "B9 未采用的不可比依据": "、".join(excluded), "B9 依据映射": bmap,
                "B10 公共刻度值": "\n".join(f"{k}：A={fmt(v[0])}，B={fmt(v[1])}" for k, v in scales.items()),
                "C1 正极类型": c.get("cathode_type"), "C1 正极材料": c.get("cathode_material"),
                "C2 硫面载量": fmt(c.get("S_loading")), "C2 值来源": c.get("S_loading_src"),
                "C3 硫含量（复合物/极片）": f"{fmt(c.get('S_content_composite'))} / {fmt(c.get('S_content_electrode'))}",
                "C4 电解液体积µL（报告｜计算输入）": f"{fmt(c.get('V_uL')) or '未写'}｜{fmt(cc['V_input'])}{'（默认）' if cc['V_def'] else ''}",
                "C4 正极面积cm²（计算输入）": f"{fmt(cc['area'], 3)}{'（默认 d=12 mm）' if cc['area_def'] else ''}",
                "C5 E/S 报告": (f"{fmt(cc['es_rep'])} {cc['es_unit']}" if cc["es_rep"] is not None else "未写"),
                "C5 E/S 计算": fmt(cc["es_calc"]) if cc["es_calc"] is not None else (f"{fmt(cc['es_calc_def'])}（含默认输入）" if cc["es_calc_def"] is not None else ""),
                "C5 E/S 采用值": fmt(cc["es_used"]), "C5 值来源": cc["es_src"],
                "C5 质量基原值 g/g": fmt(cc["es_gg"]), "C5 密度假定与换算区间": cc["es_band_note"] or (cc["es_conv"] or ""),
                "C5 档位来源": cc["es_band_src"],
                "C6 锂厚µm（报告）": fmt(cc["li_rep"]) if cc["li_rep"] is not None else "未写",
                "C6 锂厚（计算输入｜默认标记）": f"{fmt(cc['li_input'])}｜{'默认 1000' if cc['li_def'] else '—'}",
                "C6 锂面容量 mAh/cm²": fmt(cc["li_cap"]) + ("（含默认）" if cc["li_def"] else ""),
                "C7 首圈容量 A / B": f"{fmt(fa) or '—'} / {fmt(fb) or '—'}" + (f"（{row.get('first_cap_note')}）" if row.get("first_cap_note") else ""),
                "C8 N/P（实测，按B）": (fmt(cc["np_b"]) + ("（含默认输入，不进分析）" if cc["np_def"] else "")) if cc["np_b"] is not None else "",
                "C8 N/P_A": (fmt(cc["np_a"]) + ("（含默认）" if cc["np_def"] else "")) if cc["np_a"] is not None else "",
                "C9 设计比值（理论）": (fmt(cc["design"]) + ("（含默认）" if cc["li_def"] else "")) if cc["design"] is not None else "",
                "C10 作者N/P及基准": (f"{fmt(c.get('NP_author'))}｜{c.get('NP_author_basis') or ''}" if c.get("NP_author") is not None else (c.get("NP_author_basis") or "")),
                "C11 电池形式": c.get("cell"),
                "C12 倍率（作者）": (f"变倍率：{cc['rate_protocol']}" if cc["rate_protocol"]
                                  else fmt(cc["rate"]) + ("C" if cc["rate"] is not None else "") + (f"｜{c.get('rate_note')}" if c.get("rate_note") else "")),
                "C12 1C 定义 mA/g": fmt(cc["rate_def"]) if cc["rate_def"] else ("未写（按 1675）" if cc["rate"] is not None else ""),
                "C13 统一倍率": fmt(cc["rate_u"], 3),
                "C14 温度 °C": (fmt(c.get("temp_C")) if c.get("temp_C") is not None else "") + f"｜{c.get('temp_src')}"
                               + (f"；静置温度：{c['temp_rest']}" if c.get("temp_rest") else ""),
                "C15 盐阴离子 A/B": f"{ea.get('salt_anion')} / {eb.get('salt_anion')}", "C15 LiNO₃ A/B": f"{ea['lino3']} / {eb['lino3']}",
                "C16 瓶颈层": cc["layer"], "C16 瓶颈层（10-12 口径）": cc["layer_1012"],
                "C17 E/S 档": cc["es_band"], "C17 载量档": cc["load_band"], "C17 锂量档": cc["li_band"], "C17 倍率档": cc["rate_band"], "C17 温度档": cc["temp_band"],
                "C18 负极预处理": c.get("anode_pretreat"), "C19 正极孔隙率 %": fmt(c.get("porosity")),
                "D1 指标": row["metric"], "D1 指标组": mgroup,
                "D2 数值A": fmt(row.get("val_A")), "D2 数值B": fmt(row.get("val_B")), "D2 单位": row.get("unit"), "D2 取值点": row.get("point"),
                "D2 值来源A": row.get("val_src_A"), "D2 值来源B": row.get("val_src_B"),
                "D3 误差信息": row.get("error_info"), "D4 观察方向": obs, "D4 代码核对": (cdir or "") + (f"（{cnote}）" if cnote else ""),
                "D5 作者判断": row.get("author_judgement"), "D6 胜方": win, "D7 效应量": effect(row),
                "D8 证据位置": row.get("evidence_loc"), "D9 证据类型": row.get("evidence_type"),
                "E1 作者给的原因": "、".join(row.get("reasons", [])), "E1 原因原句": row.get("reason_quote"),
                "E2 作者声明的适用条件": row.get("scope_claim"), "E3 同篇控制比较": "",
                "F1 关键原句": row.get("quote"), "F2 核验标记": mark,
                "F2 核验说明": "；".join(notes + flags + ([row["note"]] if row.get("note") else [])) or "原句已在原文找到",
                "F3 随机抽查（固定原 9 条）": "", "F3b 定向复核（本轮修改涉及）": "", "F4 复核 ✍": row.get("review", ""),
            })
    return rows


def cond_diff(c1, c2):
    out = []
    for lab, k in [("E/S", "es_used"), ("载量", None), ("锂厚", "li_rep"), ("统一倍率", "rate_u")]:
        if k is None:
            a, b = c1["_c"].get("S_loading"), c2["_c"].get("S_loading")
        else:
            a, b = c1["_cc"][k], c2["_cc"][k]
        if a != b:
            out.append(f"{lab} {fmt(a) or '?'}→{fmt(b) or '?'}")
    for lab, k in [("温度", "temp_C"), ("静置温度", "temp_rest"), ("正极", "cathode_type"), ("电池", "cell"), ("孔隙率", "porosity")]:
        a, b = c1["_c"].get(k), c2["_c"].get(k)
        if a != b:
            out.append(f"{lab} {fmt(a) or '?'}→{fmt(b) or '?'}")
    if c1["_cc"]["es_band"] != c2["_cc"]["es_band"] and not any(x.startswith("E/S") for x in out):
        out.append(f"E/S 档 {c1['_cc']['es_band']}→{c2['_cc']['es_band']}")
    return out


def flip_kind(wr, wo):
    """同一对配方、同一指标跨条件的胜方关系；只有强弱胜方互换才叫翻转"""
    if "无法判断" in (wr, wo):
        return "方向不明"
    if "依据冲突" in (wr, wo):
        return "依据冲突"
    if wr == wo:
        return "均为差异小" if wr == "差异小" else "未翻转（胜方相同）"
    if {wr, wo} == {"强者胜", "弱者胜"}:
        return "翻转（强弱胜方互换）"
    return "一侧进入差异小区间"


def same_paper_controls(rows):
    ctrl = []
    by = defaultdict(list)
    for r in rows:
        by[(r["_paper"], frozenset([r["_A"], r["_B"]]), r["_metric"])].append(r)
    for (pid, pair, metric), rs in by.items():
        rs = [r for r in rs]
        if len({r["_cond"] for r in rs}) < 2:
            continue
        for r in rs:
            others, kinds = [], []
            for o in rs:
                if o is r or o["_cond"] == r["_cond"]:
                    continue
                diff = cond_diff(r, o)
                wr, wo = r["_winner"], o["_winner"]
                kind = flip_kind(wr, wo)
                kinds.append((o["A1 行号"], kind))
                pt = "取值点相同" if (r["_point"] == o["_point"] and r["_comparable_point"] and o["_comparable_point"]) else "取值点不同／待核可比性"
                pr, po = r["_c"].get("protocol"), o["_c"].get("protocol")
                if pr and po:
                    srcs = [x for x in (r["_c"].get("protocol_src"), o["_c"].get("protocol_src")) if x]
                    inferred = any("推定" in x for x in srcs)
                    pt += (f"；已报告协议相符{'，部分对应关系推定' if inferred else ''}（{'；'.join(dict.fromkeys(srcs))}）" if pr == po
                           else f"；协议不同：{pr} vs {po}")
                else:
                    pt += "；协议未逐项核"
                both_na = [lab for lab, k in (("载量", "S_loading"), ("锂厚", "Li_um")) if r["_c"].get(k) is None and o["_c"].get(k) is None]
                if both_na:
                    pt += f"；{'、'.join(both_na)}未知（两行均未报告）"
                basis = "、".join(dict.fromkeys(x for x in (r.get("_stat"), o.get("_stat")) if x))
                eff = (f"效应量 本行 {r['_effect'] or '—'}（A={r['_A']}）｜该行 {o['_effect'] or '—'}（A={o['_A']}）；正 = A 更好"
                       + (f"；{basis}的描述性差异，不作为已确认的单因素效应" if basis else "")
                       if (r["_effect"] or o["_effect"]) else "")
                others.append(f"{o['A1 行号']}（{o['_cond']}）：改变 {'、'.join(diff) or '（未识别差异）'}；胜方 本行 {wr}｜该行 {wo} → {kind}；{pt}"
                              + (f"；{eff}" if eff else ""))
            r["E3 同篇控制比较"] = "\n".join(others)
            r["_kinds"] = kinds
        ctrl.append((pid, pair, metric, rs))
    return ctrl


def in_tol(lab, a, b):
    if a is None or b is None:
        return None
    if lab == "温度":
        return abs(a - b) <= TOL["温度"]
    lo, hi = TOL[lab]
    if b == 0:
        return None
    return lo <= a / b <= hi


def key_conditions(x, y):
    """返回 [(条件, X 值, Y 值, 同/异/缺)]"""
    out = []
    pairs = [("E/S", x["_cc"]["es_used"], y["_cc"]["es_used"]), ("载量", x["_c"].get("S_loading"), y["_c"].get("S_loading")),
             ("倍率", x["_cc"]["rate_u"], y["_cc"]["rate_u"]), ("温度", x["_c"].get("temp_C"), y["_c"].get("temp_C"))]
    for lab, a, b in pairs:
        t = in_tol(lab, a, b)
        out.append((lab, a, b, "缺" if t is None else ("同" if t else "异")))
    la, lb = x["_cc"]["li_band"], y["_cc"]["li_band"]
    out.append(("锂量", la, lb, "缺" if "未报告" in (la, lb) or "不适用" in (la, lb) else ("同" if la == lb else "异")))
    ca, cb = x["_c"].get("cathode_type"), y["_c"].get("cathode_type")
    out.append(("正极类型", ca, cb, "同" if ca == cb else "异"))
    out.append(("盐阴离子", "/".join(sorted(x["_anions"])), "/".join(sorted(y["_anions"])), "同" if x["_anions"] == y["_anions"] else "异"))

    def lino(r):
        ea, eb = r["_lino3"]
        vs = [(e.get("lino3_value"), e.get("lino3_unit")) for e in (ea, eb)]
        if any(v is None for v, _ in vs):
            return None
        return vs
    la3, lb3 = lino(x), lino(y)
    if la3 is None or lb3 is None:
        st = "缺"
    else:
        pa = all(v > 0 for v, _ in la3)
        pb = all(v > 0 for v, _ in lb3)
        if not pa and not pb:
            st = "同"
        elif pa != pb:
            st = "异"
        elif {u for _, u in la3} != {u for _, u in lb3}:
            st = "缺"
        else:
            st = "同" if in_tol("LiNO3", la3[1][0], lb3[1][0]) else "异"
    out.append(("LiNO₃", x["C15 LiNO₃ A/B"], y["C15 LiNO₃ A/B"], st))
    return out


def interval_rel(x, y):
    """只用可比描述符（不可比的已在 solvation 里排除）；跨论文还要求温差 ≤5 °C"""
    rel = []
    same = x["_paper"] == y["_paper"]
    sx, sy = (x["_allsc"], y["_allsc"]) if same else (x["_scales"], y["_scales"])
    for k in set(sx) & set(sy):
        if not same:
            tx, ty = sx[k][2], sy[k][2]
            if tx is None or ty is None or abs(tx - ty) > TOL["温度"]:
                rel.append(f"{k}：温度未知或相差 >5 °C，不判")
                continue
        a0, a1 = sorted(sx[k][:2])
        b0, b1 = sorted(sy[k][:2])
        ov = min(a1, b1) - max(a0, b0)
        if a0 == a1 or b0 == b1:  # A、B 同档时区间退化为一点：点落在另一区间内（含端点）算重叠
            ok = ov >= 0
        else:
            ok = ov > 0
        rel.append(f"{k}：{'重叠' if ok else '不重叠'}{'（代理）' if '代理' in k else ''}")
    return "；".join(sorted(rel)) or "无法判断（无共同刻度）"


EXPERIMENT = {"E/S": "同一对配方，其余条件固定，在两种 E/S 下各测", "锂量": "同一对配方、固定硫载量，在薄锂与厚锂下各测",
              "载量": "同一对配方，固定 E/S 与锂量，在两种硫载量下各测", "倍率": "同一对配方，在两种倍率下各测",
              "温度": "同一对配方，在两种温度下各测"}


def build_pairs(rows):
    weak = [r for r in rows if r["_winner"] == "弱者胜"]
    strong = [r for r in rows if r["_winner"] == "强者胜"]
    pairs, seen = [], set()
    for s in strong:
        cands = []
        for w in weak:
            same = w["_paper"] == s["_paper"]
            if not same and w["_mgroup"] != s["_mgroup"]:
                continue
            kc = key_conditions(w, s)
            known = [k for k in kc if k[3] != "缺"]
            score = sum(1 for k in known if k[3] == "同")
            cands.append((same, score, len(known), w, kc))
        def keep(c):
            w = c[3]
            pw, ps = {w["_A"], w["_B"]}, {s["_A"], s["_B"]}
            return pw == ps or (pw & ps and w["_mgroup"] == s["_mgroup"] and w["_cond"] == s["_cond"])
        within = [c for c in cands if c[0] and keep(c)]
        cross = sorted([c for c in cands if not c[0]], key=lambda c: (-c[1], -c[2]))
        picked, papers = [], set()
        for c in cross:
            key = (frozenset([c[3]["_paper"], s["_paper"]]), s["_mgroup"])
            if c[3]["_paper"] in papers or key in seen:
                continue
            picked.append(c)
            papers.add(c[3]["_paper"])
            seen.add(key)
            if len(picked) == 3:
                break
        for same, score, nknown, w, kc in within + picked:
            cand = []
            if "只有叙述" in (w["_evtype"], s["_evtype"]):
                cand.append("⑤ 只有叙事")
            if w["_c"].get("cathode_type") != s["_c"].get("cathode_type"):
                cand.append("④ 口径不同（正极类型）")
            same_metric = w["_metric"] == s["_metric"]
            if not same_metric:
                cand.append(f"③ 指标不同（{w['_metric']} vs {s['_metric']}）")
            if not same_metric:
                ptchk = "指标不同（见③）"
            elif w["_point"] == s["_point"] and w["_comparable_point"] and s["_comparable_point"]:
                ptchk = "取值点相同"
            else:
                ptchk = f"取值点不同／待核可比性（{w['_point'] or '?'} vs {s['_point'] or '?'}）"
            ir = interval_rel(w, s)
            if "不重叠" in ir:
                cand.append("② 研究区间不同")
            diffs = [k[0] for k in kc if k[3] == "异" and k[0] not in ("正极类型",)]
            if diffs:
                cand.append("① 具体条件差异：" + "、".join(diffs))
            samepair = same and {w["_A"], w["_B"]} == {s["_A"], s["_B"]}
            matched = [k[0] for k in kc if k[3] == "同"]
            unmatched = [k[0] for k in kc if k[3] in ("异", "缺")]
            exps = [EXPERIMENT[d] for d in diffs if d in EXPERIMENT]
            if any(c.startswith("③") for c in cand):
                exps.append("同一组电池同时报告两类指标（容量 / 寿命）")
            pairs.append({
                "对号": "", "稳定键": f"{w['A1 行号']}|{s['A1 行号']}", "行X（弱者胜）": w["A1 行号"], "行Y（强者胜）": s["A1 行号"], "同一篇": "是" if same else "否",
                "同一对配方": "是" if samepair else "否", "指标组相同": "是" if w["_mgroup"] == s["_mgroup"] else "否",
                "具体指标相同": "是" if same_metric else "否", "取值点核对": ptchk,
                "指标 X / Y": f"{w['_metric']} / {s['_metric']}", "比较类型 X / Y": f"{w['_ctype']} / {s['_ctype']}",
                "公共刻度区间关系": ir,
                **{f"{k[0]}（X / Y｜判定）": f"{fmt(k[1]) or '?'} / {fmt(k[2]) or '?'}｜{k[3]}" for k in kc},
                "排序分（仅排核查顺序）": score, "双方已知项数": nknown,
                "候选解释（预填）": "；".join(cand) or "未能归因（关键条件均匹配或未知）",
                "支持程度": ("待核：可能为同篇控制比较" if (samepair and ptchk == "取值点相同")
                             else ("仅发现差异（同一对配方，但指标或取值点不同）" if samepair else "仅发现差异")),
                "控制了的条件": "、".join(matched), "未控制或未知的条件": "、".join(unmatched),
                "能判别的实验": "；".join(dict.fromkeys(exps)) or "先补齐未知条件再定",
                "复核 ✍": "",
                "_papers": {w["_paper"], s["_paper"]}, "_cand": cand,
            })
    pairs.sort(key=lambda p: (p["同一篇"] != "是", -p["排序分（仅排核查顺序）"]))
    for i, p in enumerate(pairs, 1):
        p["对号"] = f"P{i:02d}"
    return pairs


def summary_blocks(rows, pairs, ctrl, sampled, targeted):
    B = []
    npap = len({r["_paper"] for r in rows})
    mainrows = [r for r in rows if r["_main"]]
    nrand = len({r["_paper"] for r in rows if is_random(r["A5 抽样来源"])})
    B.append(("说明", [[f"阶段性汇总：{npap} 篇、{len(rows)} 条比较，其中主层 {len(mainrows)} 条；随机样本 {nrand} 篇，其余按入选标签或定向精读选入（非随机）。两类不能合并做领域推断；下列计数只描述这些论文。"],
                       ["阈值、容差、档位均为方案2 §5.7 的建议值；匹配容差只是候选筛选，可比与否要人工核。"],
                       ["Q1 覆盖以主层比较为分母，只统计锂库存可核实的比较；'锂库存未知'另列。"],
                       [f"E/S 档位来源分直接报告 / 计算 / 可靠换算 / 密度区间推定；'密度区间推定'只表示在假定的 {RHO[0]}–{RHO[1]} g/mL 范围内档位不变，不进入连续值匹配，敏感性分析见 2c。"]]))

    def cover(field, levels):
        tab = defaultdict(lambda: defaultdict(lambda: [set(), 0]))
        for r in mainrows:
            unit = f"{r['_mgroup']}｜{r['_ctype']}"
            lv = r["_cc"][field]
            tab[unit][lv][0].add(r["_paper"])
            tab[unit][lv][1] += 1
        out = [["分析单元（指标组｜比较类型）"] + levels]
        for unit in sorted(tab):
            out.append([unit] + [f"{len(tab[unit][lv][0])} 篇 / {tab[unit][lv][1]} 条" if tab[unit][lv][1] else "" for lv in levels])
        return out
    B.append(("1a 覆盖表 Q1（锂量档；只含主层；含同向、差异小）", cover("li_band", ["薄", "中", "厚", "未报告"])))
    B.append(("1b 覆盖表 Q2（E/S 档；只含主层）", cover("es_band", ["贫液(≤5)", "中等(5–15]", "富液(>15)", "未知"])))
    src_cnt = defaultdict(lambda: [0, set()])
    for r in mainrows:
        k = (r["_cc"]["es_band_src"], r["_cc"]["es_band"])
        src_cnt[k][0] += 1
        src_cnt[k][1].add(r["_paper"])
    B.append(("1c E/S 档位来源（主层）", [["档位来源", "档位", "比较条数", "独立论文数"]]
              + [[k[0], k[1], v[0], len(v[1])] for k, v in sorted(src_cnt.items())]))

    def direction(field, levels, keep=lambda r: True):
        out = [["指标组", "档位", "弱者胜", "强者胜", "差异小", "依据冲突", "无法判断", "弱者胜占比（按论文加权，只算弱/强）", "独立论文数"]]
        grp = defaultdict(list)
        for r in mainrows:
            if keep(r):
                grp[(r["_mgroup"], r["_cc"][field])].append(r)
        for (mg, lv) in sorted(grp):
            rs = grp[(mg, lv)]
            cnt = defaultdict(int)
            for r in rs:
                cnt[r["_winner"]] += 1
            w = defaultdict(float)
            byp = defaultdict(list)
            for r in rs:
                if r["_winner"] in ("弱者胜", "强者胜"):
                    byp[r["_paper"]].append(r["_winner"])
            for p, ws in byp.items():
                for x in ws:
                    w[x] += 1 / len(ws)
            tot = w["弱者胜"] + w["强者胜"]
            out.append([mg, lv, cnt["弱者胜"], cnt["强者胜"], cnt["差异小"], cnt["依据冲突"], cnt["无法判断"],
                        f"{w['弱者胜'] / tot:.0%}" if tot else "", len({r['_paper'] for r in rs})])
        return out
    B.append(("2a 方向分布（按锂量档，主层）", direction("li_band", None)))
    B.append(("2b 方向分布（按 E/S 档，主层）", direction("es_band", None)))
    nint = sum(1 for r in mainrows if r["_cc"]["es_band_src"] == "密度区间推定")
    B.append((f"2c 敏感性：剔除 E/S 档位为'密度区间推定'的 {nint} 条后（按 E/S 档，主层）",
              direction("es_band", None, keep=lambda r: r["_cc"]["es_band_src"] != "密度区间推定")))
    dsrc = defaultdict(lambda: [0, set()])
    for r in mainrows:
        mk = r["_bmap"].split("：")[0].split("（")[0] or "—"
        dsrc[(r["_dsrc"], mk, r["_winner"])][0] += 1
        dsrc[(r["_dsrc"], mk, r["_winner"])][1].add(r["_paper"])
    B.append(("2d 溶剂化方向来源 × 映射到电池电解液 × 胜方（主层）。间接映射 = 依据测的是纯溶剂 / 无盐体系 / 代理，可信度最高'中'",
              [["方向来源", "映射", "胜方", "比较条数", "独立论文数"]]
              + [[k[0], k[1], k[2], v[0], len(v[1])] for k, v in sorted(dsrc.items())]))
    c3 = [["论文", "配方对", "指标", "行（条件，胜方，效应量）", "两两关系"]]
    for pid, pair, metric, rs in ctrl:
        rel = sorted({f"{min(r['A1 行号'], o)}–{max(r['A1 行号'], o)}：{k}" for r in rs for o, k in r.get("_kinds", [])})
        c3.append([pid, " vs ".join(sorted(pair)), metric,
                   "\n".join(f"{r['A1 行号']}（{r['_cond']}，{r['_winner']}，{r['_effect'] or '—'}，A={r['_A']}）" for r in rs), "\n".join(rel)])
    B.append(("3 同篇控制比较候选（同一对配方、同一指标、条件不同）。翻转 = 强弱胜方互换；效应量变化不依赖强弱标签；变量是否单一、取值点是否可比见主表 E3", c3))
    cnt = defaultdict(lambda: [0, set()])
    for p in pairs:
        for c in p["_cand"] or ["未能归因"]:
            k = c.split("：")[0].split("（")[0]
            cnt[k][0] += 1
            cnt[k][1] |= p["_papers"]
    c4 = [["候选解释类别", "矛盾对数（检索用，不代表证据量）", "涉及独立论文数"]]
    for k in sorted(cnt):
        c4.append([k, cnt[k][0], len(cnt[k][1])])
    B.append(("4 候选矛盾清单汇总（同一单元每两篇最多一对）", c4))
    rv = {r["A1 行号"]: r["F4 复核 ✍"] for r in rows}
    B.append(("5a 随机抽查（基线时从 ✅/🧮 行随机抽 20%，固定沿用原 9 条，核对配方—条件对应）",
              [["行号", "核对结果 ✍"]] + [[r, rv.get(r, "")] for r in sampled]))
    B.append(("5b 定向复核（本轮修改涉及的行，与随机抽查分开；理由见主表 F3b 和'新旧对照'页）",
              [["行号", "复核理由", "核对结果 ✍"]] + [[k, v, ""] for k, v in sorted(targeted.items(), key=lambda kv: rid_key(kv[0]))]))
    q1 = sum(1 for r in mainrows if r["_cc"]["li_band"] in ("薄", "中", "厚"))
    pp = {frozenset(p["_papers"]) for p in pairs}
    B.append(("6 阶段计数（不做分支判断）", [
        [f"主层比较 {len(mainrows)} 条；锂库存可核实 {q1} 条（Q1 分母用主层）；同篇控制比较候选 {len(ctrl)} 组。"],
        [f"矛盾对 {len(pairs)} 对（同篇 {sum(1 for p in pairs if p['同一篇'] == '是')} 对），只涉及 {len({x for p in pairs for x in p['_papers']})} 篇独立论文、"
         f"{len(pp)} 个论文组合；同一行会进入多对，矛盾对是核查清单，不是独立证据数。"],
        ["方案2 §5.9 的分支判断已由方案3 取代：先看优势一致性（'优势一致性'页）与机制判别表，规律性推断等定量关口确认后再定。"]]))
    return B


GAPS = [
    "【已定 Q1】策略标签两列：作者原词 + 归类。标签只用于描述、检索、分组，不转成强弱顺序（已删除按标签推断）。强弱方向只来自可比描述符，或点名本行两种电解液的作者原句（低可信）；只有类别称呼（highly / sparingly solvating 等）的不算。",
    f"【已定 Q2】质量基 E/S：记录原始 g/g；有可靠密度时按 体积基 = 质量基 ÷ 密度 换算并注明来源；无可靠密度时只在假定的 {RHO[0]}–{RHO[1]} g/mL 范围内档位不变才给档位，档位来源记'密度区间推定'，不进连续值匹配，汇总另做剔除后敏感性分析（2c）。稀缺性不作为放宽判据的理由。",
    "【已定 Q3】新增平均放电容量（须写平均的圈数范围）、负极预处理、正极孔隙率。图里印着的数字归入'直接（图中印数）'，与读图估读区分。自放电分'容量变化'（正为增加，越高越好）与'容量损失'（越低越好），效应量用差值，只有百分数相减才叫百分点；效应量统一按正 = A 更好。",
    "【已定 Q4】DN 代理本轮不用于跨论文排序或区间判定，保留原文值、定义和来源；S06 的 NMR 法 DN 是本文自定方法，只用于同篇。统一查表后补，且须先核实数据库及测法是否一致。",
    "【已定 Q5】溶解度描述符记录物种、介质（纯溶剂 / 电解液）、测法、浓度定义及单位、温度；跨论文公共刻度要求前四项相同、温差 ≤5 °C，测法未写或含文献值的不进公共刻度。可比性不成立的保留展示，不参与排序。",
    "【已定 Q6】定向读图只补同一对配方跨条件的关键缺点，记图号、曲线、取值点、估读范围；只判方向，不进本轮效应量；估读误差足以改变判断时记无法判断。（本轮未读图）",
    "【已定 Q7】补抽从原两个随机层（弱×常规、强/适中×常规，各 3 篇）的剩余论文中，按统一资格规则用原种子顺序逐篇判，补足各层；记录种子、排除与替补。不从已入选名单凑数。（下一步执行）",
    "【翻转口径】只有强者胜 ↔ 弱者胜叫翻转；一侧变为差异小记'进入差异小区间'，不叫'差异缩小'；无法判断记'方向不明'。配方间差距随条件的变化（如 S09 自放电差值 +13.4 → +20.3 个百分点）看效应量，不依赖强弱标签；该例是代表性电池的描述性差异，不作为已确认的 E/S 单因素效应。",
    "【指标口径】具体指标不同标③；同一指标取值点不同另标'取值点不同／待核可比性'。同篇控制比较须同一对配方、同一指标、同一取值点。'第 N 圈放电容量'记'指定圈数放电容量'，不并入保持率或寿命。观察方向必须有明确的数值或作者比较支持；只有异常曲线、穿梭等现象的，指标方向记无法判断，现象另记。",
    "【映射，10-18 第三轮】'体系内可比'与'能否映射到电池电解液'分开记：描述符标'映射'（直接 = 测的就是电池电解液；间接 = 纯溶剂 / 无盐体系 / 文献代理，按溶剂映射）。间接映射的可信度最高'中'，主表 B9 依据映射列标明；S09-D4、S19-D3 只说明各自测量体系内的次序。",
    "【序号】排序类描述符只表次序，不表距离；保留作者原词（如 supersaturated），作者的解释标为解释。",
    "【协议】取值区间跨变倍率协议（如前 50 圈平均含 0.05C×2 → 0.1C×3 活化段）时记协议，不用单值倍率，倍率记缺、不进匹配。测试温度与静置温度分开记；原文只写 RT 的按 25 °C 归室温档并标'判断'。同篇控制比较在 E3 写明协议是否逐项核过、两行均未报告的条件。",
    "【仍开放】多因素比较很常见：S09 的 F4DMB vs DDN 同时改了溶剂、盐浓度、LiNO₃；S06 的 LiTFSI 版同时改了盐、浓度、稀释剂。",
    "【仍开放】正则标的'跨 E/S / 跨锂厚'多数不是同一对配方换条件。基线 5 篇里真正的同篇控制比较（新论文的见第 4 节）：S24（E/S 10/5/3/2）、S09（日历老化 E/S 5 vs 15；F4DEE vs F4DMB 室温 vs 60 °C 翻转）、S11/S19（同时改了多个条件）。",
    "【仍开放】软包里的'medium'到底是哪种配方，S11 没写；这类'电解液身份不明'需要单独标记。锂厚常不写（S24、S09 扣式、S19 扣式），按规则记未报告。",
]


# ---------- 优势一致性（方案 3 §3） ----------
CAP_TIME = ("首圈容量", "指定圈数放电容量")  # 同一量（放电容量）在不同圈数：可判"随运行时间变化"
FLIP = {"A更好": "A更差", "A更差": "A更好"}


def advantage_consistency(rows):
    """同一篇、同一对配方、同一条件下各指标的赢家是否一致；返回 (明细行, 汇总块)"""
    grp = defaultdict(list)
    for r in rows:
        a, b = sorted([r["_A"], r["_B"]])
        obs = r["D4 观察方向"] if r["_A"] == a else FLIP.get(r["D4 观察方向"], r["D4 观察方向"])
        grp[(r["_paper"], a, b)].append((r, obs))
    out, cats = [], defaultdict(set)
    for (pid, a, b), items in sorted(grp.items(), key=lambda kv: (rid_key(kv[0][0]), kv[0][1], kv[0][2])):
        bycond = defaultdict(list)
        for r, obs in items:
            bycond[r["_cond"]].append((r, obs))
        for cond, its in sorted(bycond.items()):
            main = its[0][0]["_main"]
            det = [(r, o) for r, o in its if o in ("A更好", "A更差")]
            mkey = lambda r: (r["_metric"], r["D2 单位"])  # 同一指标名下单位不同的量（如 ηact / ηconc）算不同指标
            metrics = {mkey(r) for r, _ in its}
            dirs = {o for _, o in det}
            temporal = any(r1["_metric"] in CAP_TIME and r2["_metric"] in CAP_TIME and r1["_point"] != r2["_point"] and o1 != o2
                           for i, (r1, o1) in enumerate(det) for r2, o2 in det[i + 1:])
            cross = any(mkey(r1) != mkey(r2) and o1 != o2 for i, (r1, o1) in enumerate(det) for r2, o2 in det[i + 1:])
            wins = {r["_winner"] for r, o in det}
            if temporal:
                cat = "随运行时间变化"
            elif cross:
                cat = "跨指标权衡"
            elif len(dirs) == 2:
                cat = "同一指标取值点间方向不一（非容量）"
            elif len(metrics) < 2:
                cat = "仅单一指标（不可判定）"
            elif not dirs:
                cat = "无可判定方向"
            elif wins == {"弱者胜"}:
                cat = "一致（弱者全胜）"
            elif wins == {"强者胜"}:
                cat = "一致（强者全胜）"
            else:
                cat = "一致（强弱未定）"
            if main:
                cats[cat].add(pid)
            out.append({"论文": pid, "配方对（A / B）": f"{a} / {b}", "条件": cond, "主层": "是" if main else "否",
                        "判定（初步）": cat, "指标数": len(metrics),
                        "涉及行（指标｜取值点 → A 方向｜胜方）": "\n".join(
                            f"{r['A1 行号']}：{r['_metric']}｜{r['_point']} → {o}｜{r['_winner']}" for r, o in its),
                        "逐对复核 ✍": ""})
        # 跨条件：同一指标（指标名、单位、取值点都相同）在不同条件方向相反
        bym = defaultdict(set)
        for r, o in items:
            if o in ("A更好", "A更差"):
                bym[(r["_metric"], r["D2 单位"], r["_point"])].add((r["_cond"], o))
        for (mt, unit, pt), s in bym.items():
            if len({o for _, o in s}) == 2 and len({c for c, _ in s}) >= 2:
                cats["跨条件变化（同一指标）"].add(pid)
                out.append({"论文": pid, "配方对（A / B）": f"{a} / {b}", "条件": "、".join(sorted({c for c, _ in s})), "主层": "—",
                            "判定（初步）": "跨条件变化（同一指标）", "指标数": 1,
                            "涉及行（指标｜取值点 → A 方向｜胜方）": "\n".join(
                                f"{r['A1 行号']}：{r['_metric']}｜{r['_cond']} → {o}｜{r['_winner']}" for r, o in items
                                if (r["_metric"], r["D2 单位"], r["_point"]) == (mt, unit, pt)),
                            "逐对复核 ✍": ""})
    npap = len({r["_paper"] for r in rows})
    judge = cats["随运行时间变化"] | cats["跨指标权衡"] | {p for k, v in cats.items() if k.startswith("一致") for p in v}
    order = ["随运行时间变化", "跨指标权衡", "同一指标取值点间方向不一（非容量）", "一致（弱者全胜）", "一致（强者全胜）", "一致（强弱未定）", "仅单一指标（不可判定）", "无可判定方向", "跨条件变化（同一指标）"]
    block = [[f"初步结果，待逐对复核。论文共 {npap} 篇；主层里有可判定配对证据（同一条件 ≥2 个指标且有方向）的独立论文 {len(judge)} 篇。一篇可同时落入多类。"],
             ["类别（只计主层条件）", "独立论文数", "论文"]]
    rnd = {r["_paper"] for r in rows if is_random(r["A5 抽样来源"])}
    block[-1] = ["类别（只计主层条件）", "独立论文数", "其中随机样本", "论文"]
    block += [[k, len(cats[k]), "、".join(sorted(cats[k] & rnd, key=rid_key)) or "—", "、".join(sorted(cats[k], key=rid_key))] for k in order]
    block.append([f"说明：首圈 / 倍率 / 极化不都是时间上的'早期'，CE 不自动代表持续性能；'随运行时间变化'只认同一量（放电容量）在不同圈数的反转。{npap} 篇阶段只做描述，不推断随 E/S、锂量的规律；定向精读论文是按机制价值挑的，类别比例只看随机样本列。"])
    return out, block


# ---------- 新旧对照（对基线快照） ----------
CMP_FIELDS = ["D1 指标", "D2 取值点", "D4 观察方向", "B9 溶剂化方向", "B9 方向来源", "D6 胜方", "D7 效应量",
              "C12 倍率（作者）", "C13 统一倍率", "C17 倍率档", "C14 温度 °C", "C17 温度档", "C17 E/S 档", "F2 核验标记"]
OLD_FLIP = {"翻转": "翻转（强弱胜方互换）", "未翻转": "未翻转（胜方相同）", "无法判断": "方向不明"}  # v0 的 E3 写法


def why_solv(r):
    parts = []
    if r["_excluded"]:
        parts.append(f"不可比描述符 {'、'.join(r['_excluded'])} 不参与方向判定")
    if r["_dsrc"] == "无可比依据":
        parts.append("无可比依据")
    return "；".join(parts) or "溶剂化依据变化", "主表 B7"


def tagset(s):
    return {x.split("（")[0].split("：")[0].strip() for x in str(s).split("；") if x.strip()}


def numnorm(v):
    """基线从 Excel 读回时数字会变成 '1.0'，比较前统一格式"""
    s = str(v).strip()
    try:
        return fmt(float(s), 3)
    except ValueError:
        return s


def old_kinds(e3):
    """读基线 E3 里与各行的关系（兼容 v0、v1 两种写法）"""
    k = {m.group(1): m.group(2).strip() for m in re.finditer(r"(S\d+-C\d+)（[^\n]*?→ ([^；\n]+)", e3)}
    for m in re.finditer(r"(S\d+-C\d+)（[^\n]*?胜方(翻转|未翻转|无法判断)；", e3):
        k.setdefault(m.group(1), OLD_FLIP[m.group(2)])
    return k


def is_random(src):
    """抽样来源是随机层（含补抽）的论文"""
    return str(src).startswith(("随机", "补抽"))


def compare_baseline(rows, pairs):
    """返回 (对照行列表, 定向复核 {行号: 理由})"""
    if not os.path.exists(BASELINE):
        return [], {}
    base = json.load(open(BASELINE, encoding="utf-8"))
    out, targeted = [], defaultdict(list)
    base_papers = {k.split("-")[0] for k in base["rows"]}
    newpap = defaultdict(int)
    byid = {r["A1 行号"]: r for r in rows}

    def rec(key, field, old, new, why, loc):
        out.append({"行号 / 稳定键": key, "字段": field, "旧判定": old, "新判定": new, "变更理由": why, "依据位置": loc})

    for rid, r in byid.items():
        old = base["rows"].get(rid)
        if old is None:
            if r["_paper"] not in base_papers:  # 整篇是本轮新提取：不进定向复核，按论文汇成一条
                newpap[r["_paper"]] += 1
                continue
            rec(rid, "（整行）", "（基线无此行）", "新增行", "", "")
            targeted[rid].append("新增行")
            continue
        changed = []
        for f in CMP_FIELDS:
            if f not in old:
                continue
            o, n = old[f], str(r[f] if r[f] is not None else "")
            if numnorm(o) == numnorm(n):
                continue
            if f == "F2 核验标记":
                why, loc = ("胜方变为无法判断，加 🔶" if r["D6 胜方"] == "无法判断" else "核验说明变化"), r["F2 核验说明"]
            elif r["_revise"]:
                why, loc = r["_revise"], r["D8 证据位置"]
            elif f in ("B9 溶剂化方向", "B9 方向来源", "D6 胜方"):
                why, loc = why_solv(r)
            else:
                why, loc = "口径调整", "—"
            rec(rid, f, o, n, why, loc)
            changed.append(f)
        ok_, nk_ = old_kinds(old.get("E3 同篇控制比较", "")), dict(r.get("_kinds", []))
        for o_id in sorted(set(ok_) | set(nk_), key=rid_key):
            a, b = ok_.get(o_id, "（原不同组）"), nk_.get(o_id, "（不再同组）")
            if a == b:
                continue
            orow = byid.get(o_id)
            revs = [x for x in ((orow or {}).get("_revise"), r["_revise"]) if x]
            why = "；".join(dict.fromkeys(revs)) or (
                f"胜方 本行 {old['D6 胜方']}→{r['D6 胜方']}，{o_id} {base['rows'].get(o_id, {}).get('D6 胜方', '?')}→{orow['D6 胜方'] if orow else '?'}")
            rec(rid, f"E3 与 {o_id} 的关系", a, b, why, "主表 E3")
            changed.append("E3")
        if changed:
            targeted[rid].append("本轮修正：" + (r["_revise"] or "、".join(dict.fromkeys(changed)) + " 变化（同组行改动所致，见新旧对照）"))
        if r["_used"] and all(d.get("added") for d in r["_used"]):
            targeted[rid].append(f"方向只靠第二轮新增的 {'、'.join(d['id'] for d in r['_used'])}，现标为{r['_bmap']}")

    ind = [k for k, r in byid.items() if r["_bmap"].startswith("间接")]
    dirc = [k for k, r in byid.items() if r["_bmap"].startswith("直接（")]
    mix = [k for k, r in byid.items() if r["_bmap"].startswith("直接 + 间接")]
    out.insert(0, {"行号 / 稳定键": f"（{len(ind) + len(dirc) + len(mix)} 行）", "字段": "B9 依据映射（新列）+ B7 可信度",
                   "旧判定": "无此列；实验测量且体系内可比的描述符一律记'高'",
                   "新判定": f"间接映射 {len(ind)} 行；直接 {len(dirc)} 行；直接 + 间接 {len(mix)} 行；间接映射的可信度最高记'中'",
                   "变更理由": "体系内可比 ≠ 能映射到电池电解液：纯溶剂 / 无盐体系 / 文献代理测得的方向只能间接映射；判定不变，只改可信度和标注（10-18 第三轮复核）",
                   "依据位置": "主表 B7、B9 依据映射。间接映射行：" + "、".join(ind)})
    out.insert(1, {"行号 / 稳定键": "S09 全部行", "字段": "B7 中 D4 的记录措辞",
                   "旧判定": "饱和 / 未饱和、'0.1 M 全溶'；序号（大=溶解多）",
                   "新判定": "保留原词 supersaturated（DD 为唯一例外），'颜色深 = 浓度高'标为作者解释；序号只表次序、不表距离",
                   "变更理由": "原记录把作者原词改写得过于确定（10-18 第三轮复核）；方向判定不变",
                   "依据位置": "S09 正文 p3 'all the solutions were supersaturated except for the control, DOL/DME'；Fig. 2A 图注"})

    newp = {p["稳定键"]: p for p in pairs}
    for key in sorted(set(base["pairs"]) | set(newp)):
        o, n = base["pairs"].get(key), newp.get(key)
        x, y = key.split("|")
        rx, ry = byid.get(x), byid.get(y)
        revs = "；".join(dict.fromkeys(z["_revise"] for z in (rx, ry) if z and z["_revise"]))
        w_changed = any(base["rows"].get(z, {}).get("D6 胜方") != (rz["D6 胜方"] if rz else None) for z, rz in ((x, rx), (y, ry)))
        wtxt = "；".join(f"{z} 胜方 {base['rows'].get(z, {}).get('D6 胜方', '?')}→{rz['D6 胜方'] if rz else '?'}" for z, rz in ((x, rx), (y, ry)))
        if o and not n:
            rec(key, "矛盾对", f"{o['对号']}：{o['候选解释']}｜{o['支持程度']}", "（不再成对）",
                (wtxt + (f"（{revs}）" if revs else "")) if w_changed else "配对筛选结果变化（同一指标组每两篇最多选一对，原入选行退出或改选）",
                "主表对应行")
        elif n and not o:
            rec(key, "矛盾对", "（基线无）", f"{n['对号']}：{n['候选解释（预填）']}｜{n['支持程度']}",
                wtxt if w_changed else "配对筛选结果变化：两行胜方都没变；同一指标组每两篇最多选一对，原入选行退出后改选本行",
                "主表对应行")
        elif o and n:
            if (o["候选解释"] == n["候选解释（预填）"] and o["支持程度"] == n["支持程度"]
                    and o.get("取值点核对", n["取值点核对"]) == n["取值点核对"] and o["公共刻度区间关系"] == n["公共刻度区间关系"]):
                continue
            why = []
            ot, nt = tagset(o["候选解释"]), tagset(n["候选解释（预填）"])
            for t in sorted((nt - ot) | (ot - nt)):
                why.append(("新增 " if t in nt else "去掉 ") + t)
            o1 = [s for s in o["候选解释"].split("；") if s.startswith("①")]
            n1 = [s for s in n["候选解释（预填）"].split("；") if s.startswith("①")]
            if o1 and n1 and o1 != n1:
                why.append(f"① 内容：{o1[0]} → {n1[0]}")
            if o["支持程度"] != n["支持程度"]:
                why.append(f"支持程度：{o['支持程度']} → {n['支持程度']}")
            if o.get("取值点核对") and o["取值点核对"] != n["取值点核对"]:
                why.append(f"取值点核对：{o['取值点核对']} → {n['取值点核对']}")
            if o["公共刻度区间关系"] != n["公共刻度区间关系"]:
                why.append("区间依据文字变化")
            if revs:
                why.append(f"原因：{revs}")
            rec(key, "矛盾对", f"{o['对号']}：{o['候选解释']}｜{o['支持程度']}",
                f"{n['对号']}：{n['候选解释（预填）']}｜{n['支持程度']}", "；".join(why), n["公共刻度区间关系"])
    for pid, n in sorted(newpap.items(), key=lambda kv: rid_key(kv[0])):
        rec(f"{pid}（整篇）", "（整篇）", "（基线无此论文）", f"新论文，{n} 行", "本轮首次提取（必读 / 定向精读 / 随机样本），不进定向复核", "")
    return out, {k: "；".join(dict.fromkeys(v)) for k, v in targeted.items()}


# ---------- 写 Excel ----------
HEAD_FILL = PatternFill("solid", fgColor="DDE7F3")
MARK_FILL = {"❌": "F4CCCC", "⚠️": "FCE5CD", "🔶": "FFF2CC", "🧮": "E2F0D9", "✅": "FFFFFF"}


def write_sheet(ws, df, widths=None):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for rec in df.itertuples(index=False):
        ws.append(list(rec))
    for col in ws.iter_cols(min_row=2):
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, name in enumerate(df.columns, 1):
        w = (widths or {}).get(name)
        if w is None:
            w = 14 if len(str(name)) < 10 else 18
            if any(k in name for k in ("原句", "依据条目", "配方", "说明", "E3", "候选解释", "实验")):
                w = 48
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"


def main():
    rows = build_rows()
    ctrl = same_paper_controls(rows)
    if os.path.exists(BASELINE):  # 随机抽查固定沿用基线样本，避免重跑后换样
        sampled = json.load(open(BASELINE, encoding="utf-8"))["spot_check"]
    else:
        rng = random.Random(SEED)
        okrows = [r["A1 行号"] for r in rows if r["F2 核验标记"] in ("✅", "🧮")]
        sampled = sorted(rng.sample(okrows, max(1, math.ceil(0.2 * len(okrows))))) if okrows else []
    pairs = build_pairs(rows)
    diffs, targeted = compare_baseline(rows, pairs)
    base_papers = ({k.split("-")[0] for k in json.load(open(BASELINE, encoding="utf-8"))["rows"]}
                   if os.path.exists(BASELINE) else set())
    for r in rows:
        if r["A1 行号"] in sampled:
            r["F3 随机抽查（固定原 9 条）"] = "抽中，待核配方—条件对应"
        if r["A1 行号"] in targeted:
            r["F3b 定向复核（本轮修改涉及）"] = targeted[r["A1 行号"]]
        elif r["_paper"] not in base_papers:
            r["F3b 定向复核（本轮修改涉及）"] = "新论文首次提取（不属定向复核；复核看 🔶/⚠️ 行与优势一致性逐对复核）"
    main_cols = [k for k in rows[0] if not k.startswith("_")]
    dfm = pd.DataFrame([{k: r[k] for k in main_cols} for r in rows])
    dfp = pd.DataFrame([{k: v for k, v in p.items() if not k.startswith("_")} for p in pairs]) if pairs else pd.DataFrame()
    dfd = pd.DataFrame(diffs, columns=["行号 / 稳定键", "字段", "旧判定", "新判定", "变更理由", "依据位置"])
    cons_rows, cons_block = advantage_consistency(rows)
    dfc = pd.DataFrame(cons_rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "主表"
    write_sheet(ws, dfm)
    mi = main_cols.index("F2 核验标记") + 1
    for rr in range(2, ws.max_row + 1):
        c = ws.cell(rr, mi)
        c.fill = PatternFill("solid", fgColor=MARK_FILL.get(c.value, "FFFFFF"))
    ws2 = wb.create_sheet("矛盾对")
    if len(dfp):
        write_sheet(ws2, dfp)
    ws3 = wb.create_sheet("汇总")
    r0 = 1
    for title, block in summary_blocks(rows, pairs, ctrl, sampled, targeted) + [("7 字段与判定规则（10-18 第二轮评审已定 + 仍开放）", [[g] for g in GAPS]),
                                                                              ("8 优势一致性（方案 3 §3；初步，明细见'优势一致性'页）", cons_block)]:
        ws3.cell(r0, 1, title).font = Font(bold=True, size=12)
        r0 += 1
        for line in block:
            for j, v in enumerate(line, 1):
                cell = ws3.cell(r0, j, v)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            r0 += 1
        r0 += 1
    ws3.column_dimensions["A"].width = 60
    for col in "BCDEFGHI":
        ws3.column_dimensions[col].width = 16
    ws5 = wb.create_sheet("优势一致性")
    write_sheet(ws5, dfc, widths={"涉及行（指标｜取值点 → A 方向｜胜方）": 70, "判定（初步）": 20})
    mj = os.path.join(ROOT, "output", "第一阶段_证据表", "机制判别表.json")
    if os.path.exists(mj):  # 方案 3 §4：人工撰写的机制判别表，原样渲染
        M = json.load(open(mj, encoding="utf-8"))
        wsm = wb.create_sheet("机制判别表")
        wsm.append([M["说明"]])
        wsm.append([])
        dfm2 = pd.DataFrame(M["rows"])
        dfm2["复核 ✍"] = ""
        tmp = wb.create_sheet("_tmp")
        write_sheet(tmp, dfm2, widths={c: (10 if c in ("编号", "状态") else 40) for c in dfm2.columns})
        for row in tmp.iter_rows(values_only=True):
            wsm.append(list(row))
        for i, c in enumerate(dfm2.columns, 1):
            wsm.column_dimensions[get_column_letter(i)].width = 10 if c in ("编号", "状态") else 40
        for row in wsm.iter_rows(min_row=3):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in wsm[3]:
            cell.font = Font(bold=True)
            cell.fill = HEAD_FILL
        wb.remove(tmp)
    ws4 = wb.create_sheet("新旧对照")
    write_sheet(ws4, dfd, widths={"行号 / 稳定键": 18, "字段": 18, "旧判定": 36, "新判定": 36, "变更理由": 60, "依据位置": 60})
    wb.save(OUT)

    marks = dfm["F2 核验标记"].value_counts().to_dict()
    print(f"主表 {len(dfm)} 行 → {OUT}")
    print("核验标记：", marks)
    print("胜方：", dfm["D6 胜方"].value_counts().to_dict())
    print("瓶颈层：", dfm["C16 瓶颈层"].value_counts().to_dict())
    print("矛盾对：", len(dfp), "；同篇控制比较组：", len(ctrl), "；随机抽查：", sampled)
    print("新旧对照：", len(dfd), "条；定向复核：", len(targeted), "行")
    bad = dfm[dfm["F2 核验标记"].isin(["❌", "⚠️"])][["A1 行号", "F2 核验标记", "F2 核验说明"]]
    if len(bad):
        print(bad.to_string(index=False))


if __name__ == "__main__":
    main()
