"""Claude v4 调试集 53 篇：按方案 §5.7 用程序拼“一篇一条结论（固定四行）”，出复核用 Excel 和 md。

只读 output/Claude_v4_调试集53篇_2026-10-25/independent_predictions/（冻结结果，不改）与 validation/issues.json；
不调用任何模型。输出到 output/云端_2026-10-26_Claude调试集53篇_一句话结论/。

四行（§5.7）：
  【文献】对照（对照类型）；梯度：各档（溶剂化由强到弱）；另有：…；尺子🧮
  性质：沿梯度的数值或排序（只记高低）；未测：…
  表现：每个测试一行：条件数值｜各指标方向（作者定性另标）（未报告：…）
  机制：强侧短板 / 弱侧短板 / 其他说法〔过程大类〕；作者主张
本版不算尺子（每个主盐 Li⁺ 分到的配位溶剂分子数）：53 篇组分单位写法不统一、溶剂表只收了四篇表样的溶剂，换算口径待用户定。
"""
import collections
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'output' / 'Claude_v4_调试集53篇_2026-10-25'
OUT = ROOT / 'output' / '云端_2026-10-26_Claude调试集53篇_一句话结论'
LIS_CORE = [('首圈容量', {'首圈容量'}), ('寿命或保持率', {'寿命（循环圈数）', '寿命', '保持率', '指定圈容量', '衰减率'}),
            ('库仑效率', {'库仑效率'}), ('过电位', {'过电位'})]
COND_SHOW = [('sulfur_loading', '硫载量'), ('ES', 'E/S'), ('Li_thickness', '锂厚'), ('rate', '倍率'),
             ('temperature', '温度'), ('LiNO3', 'LiNO₃')]
SOLV_KEYS = [('Li⁺ 溶剂化', ('Li⁺ 溶剂化', 'Li+ 溶剂化', '锂离子溶剂化')),
             ('多硫化物溶剂化', ('多硫化物溶剂化',)), ('多硫化物溶解度', ('多硫化物溶解度', '溶解度'))]
SIDE_NAME = {'强侧': '强溶剂化一侧的短板', '弱侧': '弱溶剂化一侧的短板', '其他': '其他说法'}
SRC_SHORT = {'本文结论': '', '本文结论（与他文一致）': '，与他文一致', '本文推测': '，推测', '转述他文或通用说法': '，转述'}

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


def reported(v):
    return v not in (None, '', '未报告', '原文未写', '不适用') and not str(v).startswith('未报告')


def short(s, n=60):
    s = str(s)
    return s if len(s) <= n else s[:n] + '…'


def cond_text(c):
    out = []
    for k, label in COND_SHOW:
        v = c.get(k)
        if not reported(v):
            continue
        if k == 'ES':
            u = c.get('ES_unit', '')
            u = re.split(r'[（(]', str(u))[0].strip() if reported(u) else ''
            v = f'{v} {u}'.strip()
        out.append(f'{label} {short(v, 40)}')
    return '；'.join(out)


def val(x):
    return '—' if x is None else str(x)


def result_text(r):
    d = r['direction_text']
    tail = ''
    if r['basis'] == 'numeric' and (r.get('value_A') is not None or r.get('value_B') is not None):
        tail = f"（{val(r.get('value_A'))} vs {val(r.get('value_B'))} {r.get('unit') or ''}".rstrip() + '）'
    elif r['basis'] == 'qualitative':
        tail = '（作者定性）'
    elif r['basis'] == 'figure_only':
        tail = '（仅图，未读）'
    if not r.get('same_conditions', True) and d == '不判':
        tail += '（条件不同）'
    return d + tail


def test_line(i, t):
    by_metric = collections.OrderedDict()
    for r in t['results']:
        by_metric.setdefault(r['metric'], []).append(r)
    parts = []
    for m, rs in by_metric.items():
        groups = collections.OrderedDict()
        for r in rs:
            groups.setdefault((r['B'], result_text(r)), []).append(r['A'])
        parts.append(m + '：' + '；'.join(f"{'、'.join(As)} 相对 {B} {txt}" for (B, txt), As in groups.items()))
    line = f"  {'①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳'[i] if i < 20 else f'({i + 1})'} {t['name']}"
    ct = cond_text(t['conditions'])
    if ct:
        line += f'｜{ct}'
    line += '：' + ('；'.join(parts) if parts else '无 A/B 比较（只测了一个配方）')
    if 'Li‖S' in t.get('cell', '') or 'Li||S' in t.get('cell', '') or '软包' in t.get('cell', ''):
        have = {r['metric'] for r in t['results']}
        miss = [name for name, keys in LIS_CORE if not (have & keys)]
        if miss:
            line += f"（未报告：{'、'.join(miss)}）"
    return line


def prop_text(p, order):
    vals = p.get('values') or {}
    name = p['name']
    if isinstance(vals, dict) and vals:
        keys = [k for k in order if k in vals] + [k for k in vals if k not in order]
        body = '、'.join(f'{k} {short(vals[k], 30)}' for k in keys[:8]) + ('…' if len(keys) > 8 else '')
        s = f'{name}：{body}'
    else:
        s = name
    if reported(p.get('direction')):
        s += f"（{short(p['direction'], 50)}）"
    return s


def conclusion(d):
    pid = d['paper_id']
    F = {f['id']: f for f in d['formulations']}
    ctrl = d.get('control_id')
    l1 = f"【{pid}】对照：" + (f"{ctrl}（{short(F[ctrl].get('raw_formula') or F[ctrl].get('name'), 80)}）" if ctrl in F else '无明确对照')
    l1 += f"，{d['control_type']}"
    in_grad = set()
    for g in d.get('gradients', []):
        l1 += f"；梯度：{g['name']} " + ' / '.join(f'{o}' + (f'（{v}）' if v and str(v) != o else '')
                                                for o, v in zip(g['order'], g.get('values') or [None] * len(g['order'])))
        l1 += '（溶剂化由强到弱）'
        in_grad |= set(g['order'])
    others = [k for k in F if k != ctrl and k not in in_grad]
    if others:
        l1 += '；另有：' + '、'.join(others)
    l1 += '；尺子🧮：本版未算'

    order = d['gradients'][0]['order'] if d.get('gradients') else list(F)
    props = [prop_text(p, order) for p in d.get('properties', [])]
    names = ' '.join(p['name'] for p in d.get('properties', []))
    untested = [k for k, keys in SOLV_KEYS if not any(x in names for x in keys)]
    l2 = '性质：' + ('；'.join(props) if props else '无') + (f"；未测：{'、'.join(untested)}" if untested else '')

    l3 = ['表现：'] + [test_line(i, t) for i, t in enumerate(d.get('tests', []))]

    l4 = ['机制：']
    by_side = collections.OrderedDict((s, []) for s in ['强侧', '弱侧', '其他'])
    for m in d.get('mechanisms', []):
        by_side.setdefault(m.get('side', '其他'), []).append(m)
    for s, ms in by_side.items():
        if not ms:
            continue
        cats = sorted({c for m in ms for c in m.get('categories', [])})
        items = '；'.join(f"{m['summary']}（{m['id'].split('-')[-1]}{SRC_SHORT.get(m.get('source_type'), '')}）" for m in ms)
        l4.append(f"  {SIDE_NAME.get(s, s)}〔{'、'.join(cats) or '未标大类'}〕：{items}")
    claims = '；'.join(c['summary'] for c in d.get('design_claims', []))
    l4.append(f"  作者主张：{claims or '未写'}")
    return '\n'.join([l1, l2] + l3 + l4)


def look_items(d):
    """要你看：同一测试、同一对配方、同一指标，作者定性方向与文字数值方向相反。"""
    out = []
    flip = {'变好': '变坏', '变坏': '变好'}
    for t in d.get('tests', []):
        rs = t['results']
        for q in rs:
            if q['basis'] != 'qualitative' or q['direction_text'] not in flip:
                continue
            for n in rs:
                if n['basis'] == 'numeric' and n['metric'] == q['metric'] and \
                        {n['A'], n['B']} == {q['A'], q['B']}:
                    nd = n['direction_text'] if n['A'] == q['A'] else flip.get(n['direction_text'], n['direction_text'])
                    if nd == flip[q['direction_text']]:
                        out.append([d['paper_id'], t['id'], q['metric'], f"{q['A']} 相对 {q['B']}",
                                    f"作者定性：{q['direction_text']}（{q['id']}）", f"数值：{nd}（{n['id']}）", ''])
    return out


def main():
    issues = json.loads((SRC / 'validation' / 'issues.json').read_text(encoding='utf-8'))
    papers = [json.loads(p.read_text(encoding='utf-8')) for p in sorted((SRC / 'independent_predictions').glob('*.json'))]
    assert len(papers) == 53, len(papers)
    rows = {k: [] for k in ['结论', '要看', '配方', '测试', '方向', '机制', '核对']}
    md = ['# Claude 调试集 53 篇：一句话结论（程序拼，方案 §5.7）', '',
          '> 由 `scripts/step9c_conclusions_debug53.py` 从冻结结果拼出，未调用模型。尺子🧮本版未算。'
          '调试集按 §14.1a 不逐篇人工复核；AI 提取，用户尚未逐项人工确认。', '']
    for d in papers:
        pid = d['paper_id']
        c = conclusion(d)
        look = look_items(d)
        rows['结论'].append([pid, d['doi'], c, len(look), '', ''])
        rows['要看'] += look
        md += [f'## {pid}', '', '```', c, '```', '']
        for f in d['formulations']:
            comps = '；'.join(f"{x.get('name')} {x.get('amount')} {x.get('unit')}（{x.get('role')}）" for x in f.get('components', []))
            rows['配方'].append([pid, f['id'], f.get('name', ''), comps, f.get('raw_formula', ''), '对照' if f['id'] == d.get('control_id') else '',
                               ' '.join(f.get('evidence_ids', []))])
        for t in d.get('tests', []):
            c9 = t['conditions']
            rows['测试'].append([pid, t['id'], t['name'], t.get('cell', '')] + [c9.get(k, '') for k in
                              ['sulfur_loading', 'ES', 'ES_unit', 'Li_thickness', 'rate', 'formation', 'temperature', 'LiNO3', 'one_C']] +
                              ['、'.join(t.get('formulations', [])), ' '.join(t.get('evidence_ids', []))])
            for r in t['results']:
                rows['方向'].append([pid, r['id'], r['metric'], r['A'], r['B'], val(r.get('value_A')), val(r.get('value_B')),
                                   r.get('unit') or '', r.get('point', ''), r['basis'], r['direction_text'],
                                   '是' if r.get('same_conditions', True) else '否', ' '.join(r.get('evidence_ids', [])), r.get('reason', '')])
        for m in d.get('mechanisms', []):
            rows['机制'].append([pid, m['id'], m.get('side', ''), '、'.join(m.get('categories', [])), m.get('source_type', ''),
                               m['summary'], m.get('A', ''), m.get('B', ''), m.get('property', ''),
                               ' '.join(m.get('evidence_ids', [])), '、'.join(m.get('test_ids', []))])
    for i in issues:
        rows['核对'].append([i['paper_id'], i['code'], i['unit'], str(i['detail'])])

    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    sheets = [
        ('说明', ['项目', '内容'], [
            ['你只看两页', '①一句话结论：一篇一格，固定四行（§5.7）。调试集不逐篇人工复核（§14.1a）；错提/漏提两列留给测试集用。②要你看：同一测试、同一对配方、同一指标，作者定性方向与文字数值方向相反的条目。其余明细页和“程序核对记录”要查时再翻。'],
            ['来源', 'output/Claude_v4_调试集53篇_2026-10-25/independent_predictions/（53 篇冻结结果，按旧规程 extraction_spec.md 提取）；程序核对记录来自同目录 validation/issues.json（step9a validate）。'],
            ['程序拼，不调模型', '结论每一个字都来自冻结 JSON 的字段；机制概括是提取时模型写的 summary，括号里是说法编号和来源（无标注 = 本文结论）。'],
            ['尺子🧮', '本版未算：53 篇组分单位写法不统一（v/v、vol%、摩尔比、未报告…），溶剂表只收了四篇表样的溶剂；换算口径待用户定。'],
            ['表现行', '只写报告了的条件；每个指标写“A 相对 B 方向（数值或作者定性）”；Li‖S/软包测试没报告的核心指标（首圈容量、寿命或保持率、库仑效率、过电位）写“未报告”。“仅图，未读”= 登记了读图请求，本轮不读图。'],
            ['状态', 'AI 提取与程序拼装，用户尚未逐项人工确认；调试集结果不是准确率。'],
        ]),
        ('①一句话结论', ['文献', 'DOI', '一句话结论', '要你看（条）', '错提 ✍', '漏提 ✍'], rows['结论']),
        ('②要你看', ['文献', '测试', '指标', '比较', '作者定性', '数值', '备注'], rows['要看']),
        ('明细-配方', ['文献', '配方', '作者原名', '组分（名称 用量 单位（角色））', '原文写法', '对照', '原句'], rows['配方']),
        ('明细-测试', ['文献', '测试', '名称', '电池', '硫载量', 'E/S', 'E/S 单位', '锂厚', '倍率', '化成', '温度', 'LiNO₃', '1C',
                     '参与配方', '原句'], rows['测试']),
        ('明细-方向', ['文献', '结果', '指标', 'A', 'B', 'A 值', 'B 值', '单位', '取值点', '依据', '方向', '同条件', '原句', '说明'], rows['方向']),
        ('明细-机制', ['文献', '说法', '侧', '过程大类', '来源', '概括', 'A', 'B', '性质', '原句', '落到测试'], rows['机制']),
        ('程序核对记录', ['文献', '类型', '位置', '说明'], rows['核对']),
    ]
    WIDE = {'一句话结论': 150, '内容': 120, '说明': 60, '概括': 70, '组分（名称 用量 单位（角色））': 60, '原文写法': 40}
    for name, head, data in sheets:
        ws = wb.create_sheet(name)
        ws.append(head)
        for r in data:
            ws.append(r)
        for j, h in enumerate(head, 1):
            ws.cell(1, j).font = Font(bold=True)
            ws.cell(1, j).fill = PatternFill('solid', fgColor='DDEBF7')
            ws.column_dimensions[ws.cell(1, j).column_letter].width = WIDE.get(h, 14)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
    wb.save(OUT / '一句话结论_Claude调试集53篇.xlsx')
    (OUT / '一句话结论_Claude调试集53篇.md').write_text('\n'.join(md), encoding='utf-8')
    print(f"53 篇结论；要你看 {len(rows['要看'])} 条；配方 {len(rows['配方'])}、测试 {len(rows['测试'])}、"
          f"方向 {len(rows['方向'])}、机制 {len(rows['机制'])}、程序核对 {len(rows['核对'])}")


if __name__ == '__main__':
    main()
