# -*- coding: utf-8 -*-
"""
论断式提取表样 v2（2026-10-23，方案 §5.7 修订）。

和 v1（step6c，10-22）的区别：
  - 一篇一个 JSON，里面分“配方 / 梯度 / 性质 / 测试 / 机理测量 / 作者归因”，不再按“一对配方一条”存；
  - 配方拆成组分（名称、角色、原文用量、单位），程序换算能精确换算的统一单位（🧮）；
  - 测试 = 条件全部相同的一组实验；每个结果挂在它的测试下，只在同一测试内判好坏；
  - 梯度上每一档都记，程序另出“梯度排序”，看最优点在梯度中间还是一端；
  - N/P 按 §7.4 分设计值（硫理论容量）和实测值（对照 B 的实测首圈容量）。

流程：json/*.json → 原句到全文缓存核对 → 判方向 → 拼一句话结论、梯度排序 → Excel（只供查看复核）。
"""
import json
import math
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "output" / "text_cache"
OUT_DIR = ROOT / "output" / "论断式提取表样_v2_2026-10-23"
JSON_DIR = OUT_DIR / "json"
OUT_FILE = OUT_DIR / "论断式提取表样_v2_S11_S21_S33_S39_含读图.xlsx"

# ---------- 1. 指标：哪个方向算“变好” ----------
BETTER = {
    "首圈容量": "↑", "首圈能量密度": "↑", "初期容量": "↑", "倍率容量": "↑", "寿命": "↑", "容量保持率": "↑", "指定圈数容量": "↑",
    "库仑效率": "↑", "过电位": "↓", "阻抗": "↓", "自放电": "↓", "衰减率": "↓", "穿梭电流": "↓",
}
LIFE = {"寿命", "容量保持率", "衰减率", "指定圈数容量"}
CORE = [("首圈容量", {"首圈容量", "首圈能量密度", "初期容量"}), ("寿命或保持率", LIFE), ("库仑效率", {"库仑效率"}), ("过电位", {"过电位"})]
CE_METRICS = {"库仑效率"}          # §7.11：库仑效率不设自动“差异小”阈值
SMALL = 0.05                       # §7.11：|A−B| ÷ max(A, B) < 5% 记差异小


def better_dir(res):
    if res.get("越大越好") is True:
        return "↑"
    name = res["指标"]
    for k, v in BETTER.items():
        if name.startswith(k):
            return v
    raise KeyError(f"指标没有“变好”方向：{name}")


# ---------- 2. 数值解析与比较（支持 ≈x、>x、<x） ----------
def parse(v):
    """返回 (下限, 上限, 显示文字)；无法解析返回 None。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return (float(v), float(v), f"{v:g}")
    s = str(v).strip()
    m = re.fullmatch(r"([<>≈]?)\s*(-?\d+(?:\.\d+)?)", s)
    if not m:
        return None
    sign, x = m.group(1), float(m.group(2))
    if sign == ">":
        return (x, math.inf, s)
    if sign == "<":
        return (-math.inf, x, s)
    return (x, x, s)


def compare(a, b, is_ce):
    """数值方向：↑ / ↓ / 差异小 / 无法判断（区间重叠）。"""
    pa, pb = parse(a), parse(b)
    if pa is None or pb is None:
        return None
    if pa[0] > pb[1]:
        d = "↑"
    elif pa[1] < pb[0]:
        d = "↓"
    elif pa[0] == pa[1] == pb[0] == pb[1]:
        return "差异小"
    else:
        return "无法判断"
    if pa[0] == pa[1] and pb[0] == pb[1] and not is_ce:
        if abs(pa[0] - pb[0]) / max(abs(pa[0]), abs(pb[0])) < SMALL:
            return "差异小"
    return d


def good_bad(direction, better):
    if direction in ("差异小",):
        return "差异小"
    if direction in ("↑", "↓"):
        return "变好" if direction == better else "变坏"
    return "无法判断"


# ---------- 3. 原句核对（与 step6c 相同） ----------
def norm(s):
    s = unicodedata.normalize("NFKC", s)
    s = re.sub(r"[\x00-]", "", s)
    return re.sub(r"\s+", "", s)


_cache = {}


def fulltext(stem, part):
    key = (stem, part)
    if key not in _cache:
        p = CACHE / f"{stem}.{part}.txt"
        if p.exists():
            raw = p.read_text(encoding="utf-8", errors="ignore")
            _cache[key] = (norm(raw), norm(re.sub(r"-[ \t]*\n", "", raw)))
        else:
            _cache[key] = None
    return _cache[key]


def check(stem, part, quote):
    if not quote:
        return "—"
    text = fulltext(stem, part)
    if text is None:
        return "❌"
    return "✅" if all(any(norm(q) in t for t in text) for q in quote.split("‖")) else "❌"


# ---------- 4. 成分：可读写法与统一换算 ----------
SOLVENT_ROLES = ("溶剂化溶剂", "弱或非溶剂化溶剂")


def unit_key(u):
    """比较单位时去掉括号里的注释，并把 mol/L 与 M 视为同一单位。"""
    u = re.sub(r"（.*?）|\(.*?\)", "", u).strip()
    return {"mol/L": "M"}.get(u, u)


def comp_text(f):
    comps = f["组分"]
    if not comps:
        return f.get("弱或非溶剂化溶剂体积分数", "成分未报告")
    solv_all = [c for c in comps if c["角色"] in SOLVENT_ROLES]
    if len(solv_all) == 1 and solv_all[0]["单位"] == "唯一溶剂":
        rest = "，".join(f"{c['名称']} {c['用量']} {c['单位']}" for c in comps if c["角色"] not in SOLVENT_ROLES)
        return f"{solv_all[0]['名称']}，{rest}"
    if all(c["单位"] == "摩尔比" for c in comps):
        return " : ".join(c["名称"] for c in comps) + " = " + " : ".join(c["用量"] for c in comps) + "（摩尔比）"
    parts_vol = [c for c in comps if c["角色"] in SOLVENT_ROLES and c["单位"].startswith("体积份")]
    pct = [c for c in comps if c["角色"] in SOLVENT_ROLES and c["单位"] == "vol%"]
    other_solv = [c for c in comps if c["角色"] in SOLVENT_ROLES and c not in parts_vol and c not in pct]
    out = []
    if parts_vol:
        out.append("/".join(c["名称"] for c in parts_vol) + " " + ":".join(c["用量"] for c in parts_vol) + " v/v")
    for c in other_solv:
        out.append(f"{c['名称']}（{c['用量']}{c['单位']}）")
    for c in pct:
        out.append(f"+ {c['用量']} vol% {c['名称']}")
    solv = " ".join(out)
    rest = "，".join(f"{c['名称']} {c['用量']} {c['单位']}" for c in comps if c["角色"] not in SOLVENT_ROLES)
    return "，".join(x for x in (solv, rest) if x)


def unify(f):
    """只做不需要额外假设的换算；返回 🧮 文字或说明为什么没算。"""
    comps = f["组分"]
    solv = [c for c in comps if c["角色"] in SOLVENT_ROLES]
    if not solv:
        return "成分未报告"
    if any(c["用量"] == "未报告" for c in comps):
        return "未算：有组分用量未报告"
    # 摩尔比：直接得到溶剂摩尔分数、盐与溶剂化溶剂摩尔比
    if all(c["单位"] == "摩尔比" for c in comps):
        tot = sum(float(c["用量"]) for c in solv)
        solvating = sum(float(c["用量"]) for c in solv if c["角色"] == "溶剂化溶剂")
        frac = "、".join(f"{c['名称']} {float(c['用量']) / tot * 100:.1f}%" for c in solv)
        salts = "；".join(f"{c['名称']} : 溶剂化溶剂 = {float(c['用量']) / solvating:.3f}"
                          for c in comps if c["角色"].startswith("盐"))
        return f"🧮 溶剂摩尔分数：{frac}；{salts}（摩尔比）"
    # 体积份 + vol%：得到溶剂体积分数
    parts = [c for c in solv if c["单位"].startswith("体积份")]
    pct = [c for c in solv if c["单位"] == "vol%"]
    if len(parts) + len(pct) == len(solv) and parts:
        try:
            p_tot = sum(float(c["用量"]) for c in parts)
            p_left = 100 - sum(float(c["用量"]) for c in pct)
        except ValueError:
            return "未算：有组分用量未报告"
        vols = [(c["名称"], float(c["用量"]) / p_tot * p_left) for c in parts] + \
               [(c["名称"], float(c["用量"])) for c in pct]
        txt = "、".join(f"{n} {v:.0f}%" for n, v in vols)
        salt_note = "；盐与溶剂的摩尔比需溶液密度，未算" if any(
            c["单位"] in ("M", "mol/L") for c in comps) else ""
        return f"🧮 溶剂体积分数：{txt}{salt_note}"
    return "未算：单位为 mol/L 或 m，换算需溶液密度或单位含义不明"


def changed(a, b):
    """A 相对 B 改变了什么（按组分逐项比）。"""
    ca = {c["名称"]: c for c in a["组分"]}
    cb = {c["名称"]: c for c in b["组分"]}
    if not ca:
        return "A 成分未报告"
    out = []
    common = [n for n in ca if n in cb]
    diff_unit = [n for n in common if unit_key(ca[n]["单位"]) != unit_key(cb[n]["单位"])]
    if diff_unit:
        ua = "、".join(sorted({unit_key(ca[n]["单位"]) for n in diff_unit} - {"唯一溶剂"}))
        ub = "、".join(sorted({unit_key(cb[n]["单位"]) for n in diff_unit} - {"唯一溶剂"}))
        out.append(f"{'、'.join(diff_unit)} 的用量写法不同（A 用 {ua}，B 用 {ub}），未换算，不能直接说增减")
    for n in common:
        if n in diff_unit:
            continue
        if ca[n]["用量"] != cb[n]["用量"]:
            out.append(f"{n} {cb[n]['用量']} → {ca[n]['用量']} {unit_key(ca[n]['单位'])}")
    for n, c in ca.items():
        if n not in cb:
            out.append(f"加入 {n}（{c['用量']} {unit_key(c['单位'])}）")
    for n in cb:
        if n not in ca:
            out.append(f"去掉 {n}")
    out.append(f"弱或非溶剂化溶剂体积分数 {b.get('弱或非溶剂化溶剂体积分数', '?')} → {a.get('弱或非溶剂化溶剂体积分数', '?')}")
    return "；".join(out)


# ---------- 5. 性质：A 相对 B ----------
WORDS = {"Li⁺ 溶剂化": ("更强", "更弱"), "多硫化物溶解度": ("更高", "更低"),
         "离子电导率": ("更高", "更低"), "Li⁺ 扩散系数": ("更高", "更低"),
         "黏度": ("更高", "更低"), "离子迁移数": ("更高", "更低")}


def rank_of(order, code):
    for i, item in enumerate(order):
        if code in item.split("≈"):
            return i
    return None


def prop_vs(p, a, b):
    hi, lo = WORDS.get(p["性质"], ("↑", "↓"))
    if "数值" in p:
        va, vb = p["数值"].get(a), p["数值"].get(b)
        if va is not None and vb is not None:
            d = compare(va, vb, False)
            word = {"↑": hi, "↓": lo, "差异小": "相当"}.get(d, "无法判断")
            return f"{word}（{va} vs {vb} {p.get('单位', '')}）"
        if va is not None:
            return f"B 未测（A {va} {p.get('单位', '')}）"
        return "未报告"
    if "排序" in p:
        ra, rb = rank_of(p["排序"], a), rank_of(p["排序"], b)
        if ra is None or rb is None:
            return "未报告"
        return "相当" if ra == rb else (lo if ra > rb else hi)
    d = p.get("方向")
    if isinstance(d, dict):
        return d.get(a, "未报告")
    return d or "未报告"


# ---------- 6. 测试条件与 N/P（§7.4） ----------
def num(x):
    return x if isinstance(x, (int, float)) else None


def np_values(t, b_code):
    li, s = num(t.get("锂厚")), num(t.get("硫载量"))
    design = f"{li * 0.2 / (s * 1.675):.2f}" if li and s else "缺失"
    measured = "缺失"
    if li and s:
        for r in t["结果"]:
            if r["指标"] == "首圈容量" and num((r.get("数值") or {}).get(b_code)):
                q = num(r["数值"][b_code]) * s * 0.001
                measured = f"{li * 0.2 / q:.2f}（按对照首圈 {r['数值'][b_code]} mAh/g）"
                break
    return design, measured


def show(x):
    if isinstance(x, dict):
        return "；".join(f"{k} {v}" for k, v in x.items())
    return "未报告" if x in (None, "") else str(x)


def cond_text(t, b_code):
    design, measured = np_values(t, b_code)
    es = t.get("E/S")
    es_txt = show(es) + (f" {t['E/S单位']}" if t.get("E/S单位") and es not in ("不适用",) else "")
    li = t.get("锂厚")
    li_txt = f"{li} µm" if num(li) else show(li)
    parts = [t["电池"], f"硫 {show(t.get('硫载量'))}" + (" mg/cm²" if num(t.get("硫载量")) else ""),
             f"E/S {es_txt}", f"锂 {li_txt}"]
    if num(li) and num(t.get("硫载量")):
        parts.append(f"设计 N/P {design}🧮、实测 N/P {measured}")
    parts += [f"{t['倍率']}（1C {t.get('1C定义', '未报告')}；化成 {t.get('化成', '未报告')}）",
              f"温度 {t.get('温度', '未报告')}", f"LiNO₃ {show(t.get('LiNO₃'))}"]
    return "，".join(parts)


def inconsistent(t, a, b):
    """条件字段若按配方分别给值且 A、B 不同，返回说明。"""
    out = []
    for k in ("E/S", "硫载量", "锂厚", "温度", "倍率"):
        v = t.get(k)
        if isinstance(v, dict) and a in v and b in v and v[a] != v[b]:
            out.append(f"{k} A {v[a]} vs B {v[b]}")
    return "；".join(out)


# ---------- 7. 单个结果在一对配方上的方向 ----------
def pair_result(res, a, b):
    better = better_dir(res)
    is_ce = res["指标"] in CE_METRICS
    vals = res.get("数值") or {}
    unit = res.get("单位", "")
    if a in vals and b in vals:
        d = compare(vals[a], vals[b], is_ce)
        detail = f"{vals[a]} vs {vals[b]} {unit}".strip()
        if res.get("来源") == "读图":
            detail += f"（读图，{res.get('图', '')}）"
        if is_ce and parse(vals[a]) and parse(vals[b]) and unit == "%":
            detail += f"，差 {parse(vals[a])[0] - parse(vals[b])[0]:+.2f} 个百分点"
        return d, good_bad(d, better), res.get("来源", "数值"), detail
    for q in res.get("作者定性", []):
        if q["A"] == a and q["B"] == b:
            d = q["方向"]
            note = "作者定性" + (f"：{q['说明']}" if q.get("说明") else "")
            return d, good_bad(d, better), "作者定性", note
    return None


# ---------- 8. 读入并展开 ----------
papers = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(JSON_DIR.glob("S*.json"))]
solvents = json.loads((JSON_DIR / "溶剂表.json").read_text(encoding="utf-8"))

sentence_rows, gradient_rows, formula_rows, cond_rows = [], [], [], []
detail_rows, prop_rows, mech_rows, attr_rows, review_rows = [], [], [], [], []


def review(pid, stem, item, what):
    if not item or not item.get("原句"):
        return
    mark = check(stem, item.get("部分", "main"), item["原句"])
    review_rows.append([pid, what, item["原句"], item.get("位置", ""), mark, ""])


for P in papers:
    pid, stem, B = P["文献编号"], P["全文缓存"], P["对照"]
    F = {f["代号"]: f for f in P["配方"]}

    # 配方表
    for f in P["配方"]:
        formula_rows.append([pid, f["代号"], f["作者原名"],
                             "\n".join(f"{c['名称']}｜{c['角色']}｜{c['用量']} {c['单位']}" for c in f["组分"]) or "未报告",
                             comp_text(f), unify(f), f.get("弱或非溶剂化溶剂体积分数", ""), f.get("备注", "")])
        review(pid, stem, f.get("原句"), f"成分：{f['代号']}")
    for g in P.get("梯度", []):
        review(pid, stem, g.get("原句"), f"梯度：{g['名称']}")
    for p in P["性质"]:
        val = p.get("数值") or p.get("排序") or p.get("方向")
        prop_rows.append([pid, p["性质"], p.get("对象", ""), json.dumps(val, ensure_ascii=False) if not isinstance(val, str) else val,
                          p.get("单位", ""), p.get("排序含义", ""), p.get("依据", ""), p.get("测量体系", "")])
        review(pid, stem, p, f"性质：{p['性质']}")
    for m in P.get("机理测量", []):
        mech_rows.append([pid, m["名称"], json.dumps(m["数值"], ensure_ascii=False), m.get("单位", ""), m.get("说明", "")])
        review(pid, stem, m, f"机理：{m['名称']}")
    for at in P["作者归因"]:
        attr_rows.append([pid, at["对象"], at["概括"], at["证据等级"], at["依据"], ""])
        review(pid, stem, at, f"归因：{at['对象']}")

    # 测试条件表
    for t in P["测试"]:
        design, measured = np_values(t, B)
        cond_rows.append([pid, t["编号"], t["名称"], t["电池"], show(t.get("硫载量")), show(t.get("E/S")),
                          t.get("E/S单位", ""), show(t.get("锂厚")), design, measured, t.get("作者N/P", ""),
                          t["倍率"], t.get("1C定义", ""), t.get("化成", ""), t.get("电压", ""), t.get("温度", ""),
                          show(t.get("LiNO₃")), "、".join(t["配方"]),
                          "；".join(x for x in (t.get("硫载量说明"), t.get("E/S说明"), t.get("锂厚说明"), t.get("条件说明")) if x)])
        for c in t.get("条件原句", []):
            review(pid, stem, c, f"条件：{t['编号']}")
        for r in t["结果"]:
            if r.get("来源") == "读图":
                review_rows.append([pid, f"结果：{t['编号']} {r['指标']}（{r['取值点']}）",
                                    json.dumps(r["数值"], ensure_ascii=False), r.get("图", ""), "🖼 读图", ""])
            else:
                review(pid, stem, r, f"结果：{t['编号']} {r['指标']}")

    # 比较对：默认每个配方对对照；测试另给“比较对”时用它
    pairs = []
    for t in P["测试"]:
        tp = t.get("比较对") or [[a, B] for a in t["配方"] if a != B and B in t["配方"]]
        for pr in tp:
            if pr not in pairs:
                pairs.append(pr)

    for a, b in pairs:
        fa, fb = F[a], F[b]
        props = "；".join(f"{p['性质']}：{prop_vs(p, a, b)}" for p in P["性质"])
        lines = [f"【{pid}】{fa['作者原名']} 对比 {fb['作者原名']}",
                 f"成分：A = {comp_text(fa)}；B = {comp_text(fb)}",
                 f"改变：{changed(fa, fb)}",
                 f"性质（A 相对 B）：{props}"]
        for t in P["测试"]:
            if a not in t["配方"] or b not in t["配方"]:
                continue
            got = {"变好": [], "变坏": [], "差异小": [], "无法判断": []}
            seen = set()
            for r in t["结果"]:
                pr = pair_result(r, a, b)
                metric = f"{r['指标']}（{r['取值点']}）"
                if pr is None:
                    continue
                d, gb, src, detail = pr
                seen.add(r["指标"])
                got[gb].append(f"{metric} {detail}".strip())
                detail_rows.append([pid, t["编号"], a, b, r["指标"], r["取值点"],
                                    show((r.get("数值") or {}).get(a)), show((r.get("数值") or {}).get(b)),
                                    d, gb, src, r.get("备注", "")])
            missing = []
            if t["电池"].startswith("Li‖S") and any(k in t["名称"] for k in ("循环", "软包", "贫液", "首两圈")):
                for label, group in CORE:
                    if not any(any(m.startswith(g) for g in group) for m in seen):
                        missing.append(label)
            bad_cond = inconsistent(t, a, b)
            head = f"{t['编号']} {t['名称']}｜{cond_text(t, b)}"
            if not seen:
                notes = "；".join(r["说明"] for r in t["结果"] if r.get("说明"))
                body = "本测试没有这对配方的可比结果" + (f"（{notes}）" if notes else "（数值只在图中或未比较这一对）")
            elif bad_cond:
                body = "⚠️ A、B 条件不同（" + bad_cond + "），不算同一测试，只列出：" + \
                       "；".join(sum(got.values(), []))
            else:
                body = "；".join(f"{k}：{'、'.join(v)}" for k, v in got.items() if v)
                if missing:
                    body += f"；未报告：{'、'.join(missing)}"
            lines.append(f"{head}\n    → {body}")
        sentence_rows.append([pid, f"{a} vs {b}", "\n".join(lines), ""])

    # 梯度排序：梯度上 ≥3 档有数值，或有作者给的最优/排序
    for g in P.get("梯度", []):
        order = g["顺序"]
        head = [f"【{pid}】梯度：{g['名称']}（" + " / ".join(f"{c} {v}" for c, v in zip(order, g["取值"])) + "）",
                f"说明：{g.get('说明', '')}"]
        for p in P["性质"]:
            if "数值" in p and sum(c in p["数值"] for c in order) >= 2:
                head.append(f"性质 {p['性质']}：" + " / ".join(show(p["数值"].get(c, "—")) for c in order) + f" {p.get('单位', '')}")
            elif "排序" in p:
                head.append(f"性质 {p['性质']}（{p.get('排序含义', '')}）：" + " > ".join(p["排序"]))
        for t in P["测试"]:
            for r in t["结果"]:
                vals = r.get("数值") or {}
                n = sum(c in vals for c in order)
                txt = None
                if n >= 3:
                    better = better_dir(r)
                    cand = [(c, parse(vals[c])) for c in order if c in vals and parse(vals[c])]
                    best = [c for c, pc in cand if all(
                        (pc[0] >= po[1] if better == "↑" else pc[1] <= po[0]) for o, po in cand if o != c)]
                    txt = " / ".join(show(vals.get(c, "—")) for c in order) + f" {r.get('单位', '')}" + \
                          (f"，最优 {best[0]}" if len(best) == 1 else "，最优不确定")
                elif r.get("最优"):
                    txt = f"最优 {r['最优']['配方']}（{r['最优']['说明']}）"
                elif r.get("排序") and len(r.get("作者定性", [])) >= 2:
                    txt = f"作者排序：{r['排序']}"
                if txt:
                    head.append(f"{t['编号']} {t['名称']}｜{r['指标']}（{r['取值点']}）：{txt}")
                    gradient_rows.append([pid, g["名称"], t["编号"], t["名称"], cond_text(t, B), r["指标"], r["取值点"], txt, ""])
        for m in P.get("机理测量", []):
            if sum(c in m["数值"] for c in order) >= 3:
                head.append(f"机理 {m['名称']}：" + " / ".join(show(m["数值"].get(c, "—")) for c in order) + f" {m.get('单位', '')}")
        sentence_rows.append([pid, "梯度排序", "\n".join(head), ""])

solvent_rows = []
for s in solvents["溶剂"]:
    mark = check(s["全文缓存"], s["部分"], s["原句"]) if s.get("原句") else "—"
    if mark == "✅" and s.get("核验说明", "").startswith("⚠️"):
        mark = "⚠️"
    solvent_rows.append([s["简称"], s["全称"], s["DN"], s["介电常数"], s["来源"], s.get("原句", ""), mark, s.get("核验说明", "")])

# ---------- 9. 说明 ----------
guide = [
    ["这是什么", "方案 §5.7 修订版（10-23）的论断式提取表样：S11、S21、S33、S39 四篇已精读论文。AI 辅助提取与核对，用户尚未逐项人工确认。不替代 10-22 表样、23 篇证据表、30 篇试提，只展示新格式。"],
    ["主存储", "json/ 里每篇一个 JSON；本 Excel 由 scripts/step6d_claim_records_v2.py 生成，只供查看复核。改数据改 JSON 再重新生成。"],
    ["测试", "条件（电池形式、硫载量、E/S、锂厚或 N/P、倍率及化成、温度、LiNO₃）全部相同的一组实验算一个测试；有一项不同就拆开。每个结果只挂在它的测试下，好坏只在同一测试内判，不把不同测试的指标拼在一起。"],
    ["一句话结论", "每对配方一格：成分 → 改变了什么 → 性质（A 相对 B）→ 每个测试一行“条件数值：变好/变坏/差异小/未报告”。有梯度的论文另有一格“梯度排序”。"],
    ["方向怎么判", "有数值：|A−B| ÷ max(A,B) < 5% 记差异小（§7.11）；库仑效率不设阈值，写百分点差。只有作者文字：标“作者定性”，不和数值方向混。≈、>、< 的值按区间比，区间重叠记无法判断。"],
    ["N/P", "§7.4：锂面容量 = 锂厚 × 0.2。设计 N/P 用硫理论容量（硫载量 × 1.675）；实测 N/P 用同一测试里对照 B 的实测首圈容量，原文没给数值就记缺失。"],
    ["成分换算 🧮", "只做不需要额外假设的换算：摩尔比 → 溶剂摩尔分数、盐与溶剂化溶剂摩尔比；体积份/vol% → 溶剂体积分数。mol/L 换摩尔比需溶液密度，不算。"],
    ["读图", "只读两类图：Li‖S 循环类测试里文字缺的同一测试容量侧（长循环图），以及梯度上只在图里有数的档。值标“读图”和图号，只用于判方向、不进效应量；被遮挡、标记重叠的标 🔶。复核表里读图项核验列为 🖼，需人工对照原图。"],
    ["证据等级", "直接：在电池里直接测到作者说的那个东西；间接：测了相关的东西，或在简化体系里测；只是说法：没有对应测量。"],
    ["复核", "带 ✍ 的列由用户填写，AI 不改。🔶 = 需人工确认；⚠️ = 原文自身有矛盾或单位问题。"],
]

# ---------- 10. 写 Excel ----------
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
USER_FILL = PatternFill("solid", fgColor="FFF2CC")
WIDE = {"一句话结论": 140, "说明": 110, "原句": 70, "组分（名称｜角色｜用量 单位）": 42, "可读写法": 45, "统一换算": 50,
        "条件": 60, "结果": 45, "概括": 45, "依据": 55, "备注": 40, "数值/排序/方向": 45, "条件说明": 45, "核验说明": 45}


def add_sheet(wb, title, head, rows):
    ws = wb.create_sheet(title)
    ws.append(head)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = USER_FILL if "✍" in h else HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = WIDE.get(h, 14)
        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if "✍" in h:
                cell.fill = USER_FILL
    ws.freeze_panes = "B2"


wb = Workbook()
wb.remove(wb.active)
add_sheet(wb, "说明", ["项目", "说明"], guide)
add_sheet(wb, "一句话结论", ["文献", "比较", "一句话结论", "复核 ✍"], sentence_rows)
add_sheet(wb, "梯度排序", ["文献", "梯度", "测试", "测试名", "条件", "指标", "取值点", "结果", "复核 ✍"], gradient_rows)
add_sheet(wb, "配方与成分", ["文献", "代号", "作者原名", "组分（名称｜角色｜用量 单位）", "可读写法", "统一换算",
                         "弱或非溶剂化溶剂体积分数", "备注"], formula_rows)
add_sheet(wb, "测试条件", ["文献", "测试", "测试名", "电池", "硫载量 mg/cm²", "E/S", "E/S 单位", "锂厚 µm",
                       "设计 N/P 🧮", "实测 N/P 🧮", "作者 N/P", "倍率", "1C 定义", "化成", "电压", "温度",
                       "LiNO₃", "参与配方", "条件说明"], cond_rows)
add_sheet(wb, "结果明细", ["文献", "测试", "A", "B", "指标", "取值点", "A 值", "B 值", "方向", "好坏", "来源", "备注"], detail_rows)
add_sheet(wb, "性质", ["文献", "性质", "对象", "数值/排序/方向", "单位", "排序含义", "依据", "测量体系"], prop_rows)
add_sheet(wb, "机理测量", ["文献", "名称", "数值", "单位", "说明"], mech_rows)
add_sheet(wb, "作者归因", ["文献", "对象", "概括", "证据等级", "依据", "复核 ✍"], attr_rows)
add_sheet(wb, "复核表", ["文献", "项目", "原句", "位置", "核验", "复核 ✍"], review_rows)
add_sheet(wb, "溶剂表", ["简称", "全称", "DN", "介电常数", "来源", "原句", "核验", "核验说明"], solvent_rows)
wb.save(OUT_FILE)

marks = [r[4] for r in review_rows]
print("写出：", OUT_FILE.relative_to(ROOT))
print("读图项：", marks.count("🖼 读图"))
print("一句话结论格数：", len(sentence_rows), " 梯度排序行：", len(gradient_rows), " 结果明细：", len(detail_rows))
print("原句数：", len(review_rows) - marks.count("🖼 读图"), " ✅", marks.count("✅"), " ❌", marks.count("❌"))
for r in review_rows:
    if r[4] == "❌":
        print("  找不到：", r[0], r[1], r[2][:90])
