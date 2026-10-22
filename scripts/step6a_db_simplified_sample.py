# -*- coding: utf-8 -*-
"""
精简版数据库表样（S39）：5 张录入表 + 复核表 + 词表 + 说明。

规则（2026-10-22 与用户讨论确定，写入方案 §5.6.10）：
- 文献编号沿用 S/P 编号；不造配方/体系/测试/观测/证据代号。
- 配方名 = 作者原名；同名不同成分时括号写"给哪个实验配的"。
- 电池类型拆成"电极组合""电池形式"两列，放测试条件表。
- 简化溶液（表征、溶解度测试用）在配方表单独一行，写明对应电池配方及差别。
- 原文有就填原文值；没有且能确定计算时才算，格内标 🧮 并写公式与输入。
- 每个事实只记一次：组成→配方表，电池条件→测试条件表，测得的性质/性能→结果表。

数据范围与 output/统一数据库设计/S39_关联示例.json 相同（局部示例），只换格式，不新增提取。
"""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 仓库根目录：脚本在 scripts/ 下，往上一层就是根目录（本机、云端都适用）
ROOT = Path(__file__).resolve().parents[1]
TEXT = ROOT / "output" / "text_cache" / "10.1021_acsenergylett.5c02200"
OUT_DIR = ROOT / "output" / "云端_2026-10-22_数据库精简"
OUT_FILE = OUT_DIR / "锂硫电解液数据库_精简表样_S39.xlsx"

# ---------- 1. 原句：从全文缓存里核对 ----------
# 每条原句记"正文/SI"，后面用代码去全文里找，找到才打 ✅
QUOTES = {
    "制备(1)(4)": ("si", "The solvent mixtures of DOL/DME (1:1 v/v), DOL/F4DEE (4:1 v/v) and DOL/TTE (4:1 v/v) were prepared in advance. Four kinds of electrolyte solutions were then prepared: (1) For Li+ solvation study, 1.0 M LiTFSI was dissolved in each single/mixed solvent;"),
    "制备(3)": ("si", "(3) For Li||Carbon paper cell construction, 1.0 M LiTFSI and 0.1 M LiNO3 were dissolved in the solutions prepared in (2);"),
    "制备(4)": ("si", "(4) For Li||S cell construction, 1.0 M LiTFSI and 0.2 M LiNO3 were dissolved in varied solvents."),
    "制备(2)": ("si", "(2) For LiPS solvation study, Li2S and S8 were mixed with a molar ratio 8:5 and then dissolved in varied solvents to reach Li2S6 nominal concentration of 0.1 M;"),
    "溶解度": ("main", "the DOL/F4DEE mixed solvent shows a moderate Li2S6 solubility of 0.38 M, in sharp contrast to 1.0 M in DOL/ DME and 0.22 M in DOL/TTE."),
    "GITT条件": ("si", "For Li||Carbon paper (CP) GITT, 2032-coin cells were constructed with Li metal as anode, carbon paper (GDL340, punched into disks of 14 mm in diameter) as cathode and W-scope as separator. Electrolyte used in Li||CP GITT test contained 1.0 M LiTFSI, 0.1 M LiNO3 and 0.1 M Li2S6 with varied solvents. 50 μL electrolyte was added in each cell. The GITT step includes a 6-minute galvanostatic discharging (or charging) at 0.12 mA cm-2, followed by a 30-minute rest. The cells were cycled first with GITT discharging steps until 1.6 V, then with GITT charging steps until 2.6 V,"),
    "第II段容量": ("main", "The DOL/DME cell delivers 900 mAh g−1 on stage II, in sharp contrast to 500 and 400 mAh g−1 in the DOL/F4DEE and DOL/TTE cells."),
    "循环协议": ("si", "For Li||S coin cells long-term cycling tests, the galvanostatic cycles include two 0.05 C formation cycles (1 C = 1200 mA gS-1), followed by long-term cycling at 0.1 C with 1.6-2.6 V voltage range."),
    "S32印字": ("si", "0.1 C cycling S loading: 3.5 mg cm-2 E/S ratio: 10 L mg-1 1.6-2.6 V"),
    "保持率基准": ("main", "Long-term cycling performance at 0.1 C, with capacity retention calculated as the specific discharge capacity normalized to its maximum value."),
}


def norm(s):
    """去掉空白和换行再比较，避免 PDF 转文本时的断行影响匹配。
    µ 在 PDF 转出的文本里有时变成私用区字符，一并去掉。"""
    s = re.sub(r"[-]", "", s)
    return re.sub(r"\s+", "", s)


texts = {p: norm((TEXT.parent / f"{TEXT.name}.{p}.txt").read_text(encoding="utf-8")) for p in ("main", "si")}


def check(key):
    """在全文缓存里找原句：找到返回 ✅，找不到返回 ❌。"""
    part, q = QUOTES[key]
    return "✅" if norm(q) in texts[part] else "❌"


def q(*keys):
    """多条原句用 ‖ 隔开放在一个格子里。"""
    return " ‖ ".join(QUOTES[k][1] for k in keys)


# ---------- 2. 五张录入表的内容 ----------
# 表头里带 ✍ 的列由用户填写，脚本一律留空

papers_head = ["文献编号", "DOI", "第一作者 年份 期刊", "来源批次", "数据集用途", "提取进度", "备注"]
papers = [
    ["S39", "10.1021/acsenergylett.5c02200", "Yang J. et al., 2025, ACS Energy Lett. 10, 5674",
     "23篇精读", "开发集", "局部示例：只含格式演示用的几条记录；全文未重提",
     "电导率、黏度等 P1 物性尚未逐项检查"],
]

form_head = ["文献编号", "配方名（作者原名＋给哪个实验配的）", "体系类型", "溶剂及比例", "盐",
             "LiNO₃", "其他添加", "对应电池配方及差别", "原句", "位置", "核验", "复核 ✍"]
forms = [
    ["S39", "DOL/F4DEE（Li‖S 电池用）", "电池电解液", "DOL:F4DEE = 4:1 v/v", "1.0 M LiTFSI",
     "0.2 M", "未报告", "—（本身就是电池配方）", q("制备(4)"), "SI §2", check("制备(4)"), ""],
    ["S39", "DOL/F4DEE（Li⁺ 溶剂化表征用）", "简化溶液", "DOL:F4DEE = 4:1 v/v", "1.0 M LiTFSI",
     "未报告", "未报告", "对应 DOL/F4DEE（Li‖S 电池用）；原文未写 LiNO₃，电池配方有 0.2 M",
     q("制备(1)(4)"), "SI §2", check("制备(1)(4)"), ""],
    ["S39", "DOL/F4DEE（Li₂S₆ 溶解度测试用）", "简化溶液", "DOL:F4DEE = 4:1 v/v", "未报告",
     "未报告", "Li₂S₆（饱和溶解试验，投料量未提取）",
     "对应 DOL/F4DEE（Li‖S 电池用）；原文称 mixed solvent，未写盐和 LiNO₃；测的是饱和溶液，不是电池内浓度",
     q("制备(1)(4)", "溶解度"), "SI §2；正文 Fig. 1d 讨论", "✅" if check("制备(1)(4)") == check("溶解度") == "✅" else "❌", ""],
    ["S39", "DOL/F4DEE（Li‖碳纸 GITT 用）", "简化溶液", "DOL:F4DEE = 4:1 v/v", "1.0 M LiTFSI",
     "0.1 M", "0.1 M Li₂S₆（名义浓度）",
     "对应 DOL/F4DEE（Li‖S 电池用）；LiNO₃ 0.1 M（电池 0.2 M）；预加 Li₂S₆；碳纸代替硫正极",
     q("制备(2)", "制备(3)"), "SI §2", "✅" if check("制备(2)") == check("制备(3)") == "✅" else "❌", ""],
    ["S39", "DOL/DME（Li‖碳纸 GITT 用）", "简化溶液", "DOL:DME = 1:1 v/v", "1.0 M LiTFSI",
     "0.1 M", "0.1 M Li₂S₆（名义浓度）",
     "对应 DOL/DME 的 Li‖S 电池配方（本表样未收）；LiNO₃ 0.1 M（电池 0.2 M）；预加 Li₂S₆",
     q("制备(2)", "制备(3)"), "SI §2", "✅" if check("制备(2)") == check("制备(3)") == "✅" else "❌", ""],
]

test_head = ["文献编号", "图表位置", "测试类型", "用到的配方", "电极组合", "电池形式", "正极",
             "硫载量", "E/S", "电解液体积", "隔膜", "倍率（作者）", "1C 定义", "电流密度",
             "电压窗口", "温度", "其他协议", "原句", "位置", "核验", "复核 ✍"]
tests = [
    ["S39", "Fig. 1d；SI Fig. S8", "Li₂S₆ 溶解度测试", "DOL/F4DEE（Li₂S₆ 溶解度测试用）",
     "—（非电池）", "—", "—", "—", "—", "—", "—", "—", "—", "—", "—", "", "饱和判据、平衡时间待查",
     q("溶解度"), "正文", check("溶解度"), ""],
    ["S39", "Fig. 3a", "Li‖碳纸 GITT", "DOL/F4DEE（Li‖碳纸 GITT 用）；DOL/DME（Li‖碳纸 GITT 用）",
     "Li‖碳纸", "扣式 2032", "碳纸 GDL340，直径 14 mm", "—（无硫正极）", "—", "50 µL", "W-scope",
     "—", "—", "0.12 mA/cm²", "1.6–2.6 V", "", "6 min 恒流 / 30 min 静置；先放电至 1.6 V，再充电至 2.6 V",
     q("GITT条件"), "SI §4", check("GITT条件"), ""],
    ["S39", "Fig. 5d；SI Fig. S32", "Li‖S 长期循环", "DOL/F4DEE（Li‖S 电池用）",
     "Li‖S", "扣式（型号未报告）", "", "3.5 mg/cm²", "10 µL/mg", "", "",
     "0.1 C", "1200 mA/g（按硫）", "", "1.6–2.6 V", "",
     "先 2 圈 0.05 C 化成；保持率 = 放电比容量 ÷ 其最大值",
     q("循环协议", "S32印字", "保持率基准"), "SI §4；SI Fig. S32 图内文字；正文 Fig. 5 图注",
     "🔶", ""],
]

res_head = ["文献编号", "配方名", "图表位置", "指标", "取值点", "数值", "单位", "来源类型",
            "原句", "位置", "核验", "复核 ✍"]
results = [
    ["S39", "DOL/F4DEE（Li₂S₆ 溶解度测试用）", "Fig. 1d；SI Fig. S8", "Li₂S₆ 溶解度", "—", "0.38", "M",
     "文献实测", q("溶解度"), "正文 Fig. 1d 讨论", check("溶解度"), ""],
    ["S39", "DOL/F4DEE（Li‖碳纸 GITT 用）", "Fig. 3a", "第 II 段放电比容量", "首次 GITT 放电", "500",
     "mAh/g", "文献实测", q("第II段容量"), "正文 Fig. 3a 讨论", check("第II段容量"), ""],
    ["S39", "DOL/DME（Li‖碳纸 GITT 用）", "Fig. 3a", "第 II 段放电比容量", "首次 GITT 放电", "900",
     "mAh/g", "文献实测", q("第II段容量"), "正文 Fig. 3a 讨论", check("第II段容量"), ""],
]

comp_head = ["文献编号", "A 配方", "B 配方", "B 是什么", "改变了什么", "指标", "取值点", "A 值",
             "B 值", "方向", "可比性与条件差异", "作者怎么说（原句）", "核验", "复核 ✍"]


def lookup(form, metric, point):
    """A/B 值不手抄：从结果表里按"配方名＋指标＋取值点"找出来。"""
    for r in results:
        if r[1] == form and r[3] == metric and r[4] == point:
            return f"{r[5]} {r[6]}"
    return "结果表中未找到"


comps = []
for a, b in [("DOL/F4DEE（Li‖碳纸 GITT 用）", "DOL/DME（Li‖碳纸 GITT 用）")]:
    comps.append(["S39", a, b, "作者比较（同一系列参照）", "共溶剂 F4DEE→DME；溶剂比例 4:1→1:1",
                  "第 II 段放电比容量", "首次 GITT 放电",
                  lookup(a, "第 II 段放电比容量", "首次 GITT 放电"),
                  lookup(b, "第 II 段放电比容量", "首次 GITT 放电"),
                  "A 更低", "同一 GITT 系列、条件相同；温度未查；只代表第 II 段，不能当作 Li‖S 电池容量",
                  q("第II段容量"), check("第II段容量"), ""])

# ---------- 3. 复核表：逐行列出核验结果和问题 ----------
review_head = ["表", "行（配方名或图表位置）", "核验", "原句", "说明 / 问题"]
review = []
for r in forms:
    review.append(["配方", r[1], r[10], r[8], "组成按 SI §2 原句；未写的组分记“未报告”，不填 0"])
for r in tests:
    note = "条件按原句"
    if r[1].startswith("Fig. 5d"):
        note = ("🔶 Fig. 5d 与 SI Fig. S32 按同一 0.1 C 循环系列合并（沿用旧示例），"
                "原文没有明说是同一批电池，需人工确认；正极直径、电解液体积未查，"
                "所以不计算电解液体积（规则：输入不全不算）")
    elif r[1].startswith("Fig. 1d"):
        note = "测试温度、饱和判据还没查，温度格留空（空白 = 还没查）"
    review.append(["测试条件", r[1], r[19], r[17], note])
for r in results:
    review.append(["结果", f"{r[1]}｜{r[3]}", r[10], r[8], "数值与原句一致"])
for r in comps:
    review.append(["A/B 比较", f"{r[1]} vs {r[2]}", r[12], r[11],
                   "A/B 值由程序从结果表带入；方向为数值比较"])

# ---------- 4. 词表与说明 ----------
vocab_head = ["项目", "可选写法", "含义"]
vocab = [
    ["体系类型", "电池电解液", "真正装进电池、性能结果对应的那瓶电解液"],
    ["体系类型", "简化溶液", "为表征、溶解度、碳纸 GITT 等单独配的溶液；结果只能当间接证据"],
    ["电极组合", "Li‖S、Li‖Cu、Li‖Li、Li‖碳纸、其他（写清）", "不用“全电池/半电池”，各文献叫法不统一；作者叫法留在原句里"],
    ["电池形式", "扣式 2032、扣式（型号未报告）、软包、其他（写清）", ""],
    ["来源类型", "文献实测、文献模拟（MD/DFT）、本课题计算", "提取自论文不等于实验值"],
    ["缺失写法", "空白", "还没查"],
    ["缺失写法", "未报告", "查过正文和 SI，没写"],
    ["缺失写法", "未加入", "原文明确说没加（例如 without LiNO3）"],
    ["缺失写法", "仅图未读", "数值只在图里，没读出来"],
    ["计算值", "🧮 数值（公式与输入）", "原文没有、输入确定时才算，例：14.5 µL/mg（🧮 25 µL ÷ 1.72 mg）"],
    ["核验", "✅", "原句在全文里找到，内容一致"],
    ["核验", "🧮", "本课题计算，输入均已核对"],
    ["核验", "🔶", "能找到依据，但有需要人工判断的地方"],
    ["核验", "⚠️", "原文本身有歧义或前后不一"],
    ["核验", "❌", "原句找不到或与记录不符"],
    ["指标（示例，持续补充）", "首圈放电比容量、第 n 圈放电比容量、容量保持率、平均库仑效率、循环寿命、"
     "第 II 段放电比容量、Li₂S₆ 溶解度、离子电导率、黏度、Li⁺ 配位数……", "从这里选，不自造英文代码"],
]

guide_head = ["项目", "说明"]
guide = [
    ["这是什么", "按 10-22 讨论定下的精简格式，把旧表样里的 S39 局部记录重新排了一遍。只换格式，不新增提取；不是人工 gold。"],
    ["怎么读", "从左往右：文献 → 配方 → 测试条件 → 结果 → A/B 比较。表与表之间只靠“文献编号＋配方名”“文献编号＋图表位置”连接，都是原文里本来就有的名字。"],
    ["每个事实只记一次", "组成只写在配方表；电池条件只写在测试条件表；结果表只放测出来的性质和性能。"],
    ["一行是什么", "配方表：一瓶溶液；测试条件表：一组实验（按原文图表）；结果表：一个数；比较表：同篇同条件的一对配方的一个指标。"],
    ["同一个配方多个数", "100 圈、200 圈容量分两行；0 °C、25 °C 黏度分两行。不加“描述符1/2/3”列。"],
    ["简化溶液", "表征、溶解度测试单独配的溶液，在配方表里单独一行，“体系类型”写简化溶液，并写对应电池配方和差别。"],
    ["计算值", "原文有就填原文值；没有且输入确定时才算，格内写 🧮 和公式。统一倍率、单位硫 LiNO₃ 投料等分析用的派生量由程序在分析时生成，不在录入表里手填。"],
    ["人工列", "带 ✍ 的列由用户填写，AI 不改。"],
    ["与旧表样对照", "旧表样 10 张表约 300 列 → 本表 5 张录入表 66 列。删去：数据库主表（分析时程序生成）、实际体系表（并入配方表）、配方组分表（并入配方表）、证据表与证据ID（原句放在当行）、论断表（作者解释进比较表，研究者假设移到 Q2）、各表复核人/时间等管理列、误差 8 列（写进数值格，如 0.38 ± 0.02 (SD, n=3)）。"],
    ["旧表样中不再单列的记录", "LiNO₃ 浓度、E/S、倍率不再作为“观测”重复记录（已在配方表/测试条件表）；保持率基准写进测试条件“其他协议”；“电解液体积/电导率 未提取”这类空记录不再占行，空白即还没查；研究者假设 H-map 移到 Q2 机制阶段。"],
    ["状态", "AI 辅助核对，用户尚未逐项人工确认。"],
]

# ---------- 5. 写 Excel ----------
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
USER_FILL = PatternFill("solid", fgColor="FFF2CC")  # 黄色 = 用户填写


def add_sheet(wb, title, head, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(head)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = USER_FILL if "✍" in h else HEAD_FILL
        long = h in ("原句", "作者怎么说（原句）", "说明", "说明 / 问题", "对应电池配方及差别")
        w = (widths or {}).get(h, 50 if long else 16)
        ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
            if "✍" in h:
                row[0].fill = USER_FILL
    ws.freeze_panes = "B2"
    return ws


wb = Workbook()
wb.remove(wb.active)
add_sheet(wb, "说明", guide_head, guide, {"项目": 18, "说明": 100})
add_sheet(wb, "①文献", papers_head, papers, {"第一作者 年份 期刊": 36, "提取进度": 36, "备注": 30})
add_sheet(wb, "②配方", form_head, forms, {"配方名（作者原名＋给哪个实验配的）": 30, "溶剂及比例": 20})
add_sheet(wb, "③测试条件", test_head, tests, {"用到的配方": 30, "其他协议": 30})
add_sheet(wb, "④结果", res_head, results, {"配方名": 30})
add_sheet(wb, "⑤A／B比较", comp_head, comps, {"A 配方": 28, "B 配方": 28, "可比性与条件差异": 36})
add_sheet(wb, "复核表", review_head, review, {"行（配方名或图表位置）": 36})
add_sheet(wb, "词表", vocab_head, vocab, {"可选写法": 50, "含义": 60})

OUT_DIR.mkdir(parents=True, exist_ok=True)
wb.save(OUT_FILE)

n_cols = sum(len(h) for h in (papers_head, form_head, test_head, res_head, comp_head))
print("写出：", OUT_FILE.relative_to(ROOT))
print("录入表列数合计：", n_cols)
print("核验：", {k: check(k) for k in QUOTES})
