# -*- coding: utf-8 -*-
"""
论断式提取表样 v4（2026-10-24 用户定）：一篇一条结论；复核只给用户看两页。

和 v3（step6g，已移入 不适用/）的区别：
  - 一句话结论改为一篇一格、固定四行：对照与梯度 / 性质 / 各测试的表现（只写方向）/ 机制（两侧短板 + 作者主张）；
  - 性能数值只用来判方向，不进分析；删去证据等级、机制链的节点与边、跨文献三页（跨文献比较等全部论文提完后单独做）；
  - 机制一句一行：整句原文 + 来源 + 这句提到的过程大类（正极反应 / 多硫化物穿梭 / 负极锂 / 电解液消耗，可多标）；
  - 尺子：每个主盐 Li⁺ 分到的配位溶剂分子数（配方 + 手册分子量、密度，不用 DFT/MD；稀释剂、LiNO₃ 不计）；
  - 复核分工（10-24 用户同意）：能让程序判的不给人判——数字必须出现在它挂的原句里、原句必须在全文里；
    原文自己没写清的按固定规则处理（不判 / 原文未写 / 不换算）；读图只用锚点核对通过的；
    用户只看两页：①一句话结论（测试集里填错提、漏提）和②要你看（作者说法与本篇数据对不上的）。
    其余核对结果全部进“程序核对记录”，用户不用看。

流程：json/*.json → 原句、数字到全文缓存核对 → 判方向 → 拼结论 → Excel。
"""
import json
import math
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "output" / "text_cache"
OUT_DIR = ROOT / "output" / "论断式提取表样_v4_2026-10-24"
JSON_DIR = OUT_DIR / "json"
OUT_FILE = OUT_DIR / "论断式提取表样_v4_S11_S21_S33_S39.xlsx"

# ---------- 1. 指标：哪个方向算“变好” ----------
BETTER = {
    "首圈容量": "↑", "首圈能量密度": "↑", "初期容量": "↑", "倍率容量": "↑", "寿命": "↑", "容量保持率": "↑", "指定圈数容量": "↑",
    "库仑效率": "↑", "过电位": "↓", "阻抗": "↓", "自放电": "↓", "衰减率": "↓", "穿梭电流": "↓",
}
LIFE = {"寿命", "容量保持率", "衰减率", "指定圈数容量"}
CORE = [("首圈容量", {"首圈容量", "首圈能量密度", "初期容量"}), ("寿命或保持率", LIFE), ("库仑效率", {"库仑效率"}), ("过电位", {"过电位"})]
PERF = {"首圈/初期容量": ("首圈容量", "首圈能量密度", "初期容量"), "寿命/容量保持": tuple(LIFE),
        "库仑效率": ("库仑效率",), "倍率/过电位": ("倍率容量", "过电位"), "自放电": ("自放电",),
        "阻抗": ("阻抗",), "穿梭电流": ("穿梭电流",)}
CE_METRICS = {"库仑效率"}          # §7.11：库仑效率不设自动“差异小”阈值
SMALL = 0.05                       # §7.11：|A−B| ÷ max(A, B) < 5% 记差异小


def better_dir(res):
    if res.get("越大越好") is True:
        return "↑"
    for k, v in BETTER.items():
        if res["指标"].startswith(k):
            return v
    raise KeyError(f"指标没有“变好”方向：{res['指标']}")


# ---------- 2. 数值解析与比较（支持 ≈x、>x、<x） ----------
def parse(v):
    """返回 (下限, 上限, 显示文字)；无法解析返回 None。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return (float(v), float(v), f"{v:g}")
    m = re.fullmatch(r"([<>≈]?)\s*(-?\d+(?:\.\d+)?)", str(v).strip())
    if not m:
        return None
    sign, x = m.group(1), float(m.group(2))
    if sign == ">":
        return (x, math.inf, str(v))
    if sign == "<":
        return (-math.inf, x, str(v))
    return (x, x, str(v))


def compare(a, b, is_ce):
    """A 相对 B：↑ / ↓ / 差异小 / 无法判断（区间重叠）。"""
    pa, pb = parse(a), parse(b)
    if pa is None or pb is None:
        return None
    if pa[0] > pb[1]:
        d = "↑"
    elif pa[1] < pb[0]:
        d = "↓"
    elif pa[0] == pa[1] == pb[0] == pb[1]:
        return "差异小"
    else:
        return "无法判断"
    if pa[0] == pa[1] and pb[0] == pb[1] and not is_ce:
        if abs(pa[0] - pb[0]) / max(abs(pa[0]), abs(pb[0])) < SMALL:
            return "差异小"
    return d


def good_bad(direction, better):
    if direction == "差异小":
        return "差异小"
    if direction in ("↑", "↓"):
        return "变好" if direction == better else "变坏"
    return "无法判断"


def show(x):
    if isinstance(x, dict):
        return "；".join(f"{k} {v}" for k, v in x.items())
    return "未报告" if x in (None, "") else str(x)


def num(x):
    return x if isinstance(x, (int, float)) else None


def reported(x):
    return x not in (None, "", "未报告", "不适用")


# ---------- 3. 原句与数字核对 ----------
def norm(s):
    s = unicodedata.normalize("NFKC", s)
    s = re.sub("[\x00-]", "", s)
    return re.sub(r"\s+", "", s)


_cache = {}


def fulltext(stem, part):
    key = (stem, part)
    if key not in _cache:
        p = CACHE / f"{stem}.{part}.txt"
        if p.exists():
            raw = p.read_text(encoding="utf-8", errors="ignore")
            _cache[key] = (norm(raw), norm(re.sub(r"-[ \t]*\n", "", raw)))
        else:
            _cache[key] = None
    return _cache[key]


def found(stem, part, quote):
    """原句（可用 ‖ 分段）是否都在全文里。"""
    if not quote:
        return False
    text = fulltext(stem, part)
    return text is not None and all(any(norm(q) in t for t in text) for q in quote.split("‖"))


WORDNUM = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
           "eleven": 11, "twelve": 12, "fifteen": 15, "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "hundred": 100,
           "first": 1, "second": 2, "third": 3}


def num_in_quote(v, quote):
    """数值是否出现在它挂的原句里（英文数字词也认，如 eight cycles）。非数值不查。"""
    x = re.sub(r"^[≈<>~]\s*", "", str(v)).strip()
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", x):
        return True
    s = norm(re.sub(r"\b(" + "|".join(WORDNUM) + r")\b", lambda m: str(WORDNUM[m.group(1).lower()]), quote or "", flags=re.I))
    cands = {x, x.rstrip("0").rstrip(".") if "." in x else x}
    return any(re.search(rf"(?<![\d.]){re.escape(c)}(?!\d)", s) for c in cands)


SRC = ["本文结论", "本文结论（与他文一致）", "本文推测", "转述他文或通用说法"]
CITED = SRC[-1]
CITE_RE = [r"\[\d{1,3}[a-z]?(?:\s*[,–−-]\s*\d{1,3}[a-z]?)*\]",
           r"(?:(?<=[a-z\)])|(?<=[a-z\)][\.,]))\d{1,3}(?:\s*[,−–]\s*\d{1,3})*(?=[\s\.,;:)]|$)"]  # 上标引文；排除 0.2、2.072 这类小数
PHRASE_RE = r"\b(it is reported|has been reported|have been reported|were reported|was reported|as reported|previous(?:ly)? (?:stud|report|work)|reported previously|according to|in accordance with|aligns? well with|agrees? (?:well )?with|recent studies)"
HEDGE_RE = r"\b(may|might|likely|possibly|presumably|probably|can be attributed)\b"
_located = {}


def _mapped(stem, part, join_hyphen):
    key = (stem, part, join_hyphen)
    if key not in _located:
        p = CACHE / f"{stem}.{part}.txt"
        raw = p.read_text(encoding="utf-8", errors="ignore") if p.exists() else ""
        if join_hyphen:
            raw = re.sub(r"-[ \t]*\n", "", raw)
        chars, pos = [], []
        for i, c in enumerate(raw):
            for n in unicodedata.normalize("NFKC", c):
                if not n.isspace() and ord(n) >= 32:
                    chars.append(n)
                    pos.append(i)
        _located[key] = (raw, "".join(chars), pos)
    return _located[key]


def source_flags(stem, part, quote):
    """原句放回原文，查本句和下一句的引文编号、转述说法、推测词（只作提示）。"""
    frags = [f for f in quote.split("‖") if f.strip()]
    for join_hyphen in (True, False):
        raw, nt, pos = _mapped(stem, part, join_hyphen)
        spans = []
        for f in frags:
            nf = "".join(ch for ch in unicodedata.normalize("NFKC", f) if not ch.isspace() and ord(ch) >= 32)
            k = nt.find(nf)
            if k < 0:
                break
            spans.append((pos[k], pos[k + len(nf) - 1] + 1))
        if len(spans) == len(frags):
            break
    else:
        return "原文中找不到"
    b = spans[-1][1]
    tail = raw[b:b + 40]
    nxt_end = raw.find(". ", b + 2)
    nxt = raw[b:nxt_end + 1] if nxt_end > 0 else tail
    sent = " ".join(frags) + tail.split(". ")[0]
    hits = []
    for name, txt in (("本句", sent), ("下一句", nxt)):
        cites = [m.group(0) for r_ in CITE_RE for m in re.finditer(r_, txt)]
        cites = [c for c in cites if not re.search(r"S\d|Li\d", c)]
        if cites:
            hits.append(f"{name}有引文编号 {' '.join(cites)}")
        ph = re.findall(PHRASE_RE, txt, flags=re.I)
        if ph:
            hits.append(f"{name}有转述说法 “{'、'.join(ph)}”")
        if name == "本句":
            hd = re.findall(HEDGE_RE, txt, flags=re.I)
            if hd:
                hits.append(f"本句有推测词 “{'、'.join(hd)}”")
    return "；".join(hits) or "无"


def source_check(src, flags):
    """AI 判的来源与程序痕迹对不上时返回说明，对得上返回空串。"""
    own = "本句有引文编号" in flags or "本句有转述说法" in flags
    nxt = "下一句有引文编号" in flags or "下一句有转述说法" in flags
    if src == SRC[0] and (own or "推测词" in flags):
        return f"判为本文结论，但{flags}"
    if src == SRC[1] and not (own or nxt):
        return "判为与他文一致，但本句和下一句都没查到引用"
    if src == CITED and not own:
        return "判为转述，但本句没查到引用痕迹（可能是通用说法）"
    return ""


# ---------- 4. 过程大类：原句提到哪类就标哪类，程序用关键词核对 ----------
CATS = {"正极反应": ("改善", "变差", "不变"), "多硫化物穿梭": ("减轻", "加重"),
        "负极锂": ("改善", "劣化"), "电解液消耗": ("减少", "加快")}
KW = {
    "正极反应": r"cathode|sulfur (?:redox|reduction|reversib|utili|conversion)|sulfur active material|li2s (?:deposition|nucleation|precipitation)"
               r"|deposition of li2s|conversion|redox kinetics|oxidation|reduction|disproportionation|chain-shortening|discharge plateau"
               r"|capacity activation|insulating li2s|passivat",
    "多硫化物穿梭": r"shuttl",
    "负极锂": r"anode|li metal|lithium metal|\blma\b|li deposition|li-deposition|li growth|plating|stripping|\bsei\b"
             r"|solid electrolyte interphase|dendrit|corros|remaining li|li cyclability",
    "电解液消耗": r"electrolyte (?:depletion|decomposition|consumption|dry-out|exchange)|decomposition of dme|solvent depletion|dry-out|depletion",
}
SOLVATION = {"Li⁺ 溶剂化", "多硫化物溶剂化", "多硫化物溶解度", "多硫化物团簇化"}


def mention_check(labels, sentence):
    s = unicodedata.normalize("NFKC", sentence).lower()
    hit = [c for c, p in KW.items() if re.search(p, s)]
    out = []
    miss = [c for c in labels if c not in hit]
    extra = [c for c in hit if c not in labels]
    if miss:
        out.append(f"标了{'、'.join(miss)}，但句中没找到相关词")
    if extra:
        out.append(f"句中有{'、'.join(extra)}相关词，但没标")
    return "；".join(out)


# ---------- 5. 成分：可读写法、🧮 换算、类别与尺子 ----------
SOLVENT_ROLES = ("溶剂化溶剂", "弱或非溶剂化溶剂")
solvents = json.loads((JSON_DIR / "溶剂表.json").read_text(encoding="utf-8"))
M_NOTE = "原文小写 m（mol/kg 还是 mol/L 不明），按规则不换算"


def base(name):
    return re.sub(r"（.*?）|\(.*?\)", "", name).strip()


SOLV = {}
for _s in solvents["溶剂"]:
    for _k in [_s["简称"]] + _s.get("别名", []):
        SOLV[base(_k)] = _s


def unit_key(u):
    """比较单位时去掉括号里的注释，并把 mol/L 与 M 视为同一单位。"""
    u = re.sub(r"（.*?）|\(.*?\)", "", u).strip()
    return {"mol/L": "M"}.get(u, u)


def comp_text(f):
    comps = f["组分"]
    if not comps:
        return f.get("弱或非溶剂化溶剂体积分数", "成分未报告")
    solv_all = [c for c in comps if c["角色"] in SOLVENT_ROLES]
    if len(solv_all) == 1 and solv_all[0]["单位"] == "唯一溶剂":
        rest = "，".join(f"{c['名称']} {c['用量']} {c['单位']}" for c in comps if c["角色"] not in SOLVENT_ROLES)
        return f"{solv_all[0]['名称']}，{rest}"
    if all(c["单位"] == "摩尔比" for c in comps):
        return " : ".join(c["名称"] for c in comps) + " = " + " : ".join(c["用量"] for c in comps) + "（摩尔比）"
    parts_vol = [c for c in comps if c["角色"] in SOLVENT_ROLES and c["单位"].startswith("体积份")]
    pct = [c for c in comps if c["角色"] in SOLVENT_ROLES and c["单位"] == "vol%"]
    other = [c for c in comps if c["角色"] in SOLVENT_ROLES and c not in parts_vol and c not in pct]
    out = []
    if parts_vol:
        out.append("/".join(c["名称"] for c in parts_vol) + " " + ":".join(c["用量"] for c in parts_vol) + " v/v")
    for c in other:
        out.append(f"{c['名称']}（{c['用量']}{c['单位']}）")
    for c in pct:
        out.append(f"+ {c['用量']} vol% {c['名称']}")
    rest = "，".join(f"{c['名称']} {c['用量']} {c['单位']}" for c in comps if c["角色"] not in SOLVENT_ROLES)
    return "，".join(x for x in (" ".join(out), rest) if x)


def vol_fractions(f):
    """溶剂体积分数（0–1）；比例未报告或单位不是体积时返回 None。"""
    solv = [c for c in f["组分"] if c["角色"] in SOLVENT_ROLES]
    if len(solv) == 1 and solv[0]["单位"] == "唯一溶剂":
        return {solv[0]["名称"]: 1.0}
    parts = [c for c in solv if c["单位"].startswith("体积份")]
    pct = [c for c in solv if c["单位"] == "vol%"]
    if not solv or len(parts) + len(pct) != len(solv):
        return None
    try:
        left = 1 - sum(float(c["用量"]) for c in pct) / 100
        out = {c["名称"]: float(c["用量"]) / 100 for c in pct}
        if parts:
            tot = sum(float(c["用量"]) for c in parts)
            out.update({c["名称"]: float(c["用量"]) / tot * left for c in parts})
        return out
    except ValueError:
        return None


def has_m(f):
    return any(unit_key(c["单位"]) == "m" for c in f["组分"])


def salt_ratio(f):
    """不需要额外假设的“盐 : 溶剂化溶剂”摩尔比（只用摩尔比配方）；算不了返回 None。"""
    comps = f["组分"]
    solv = [c for c in comps if c["角色"] in SOLVENT_ROLES]
    salts = [c for c in comps if c["角色"].startswith("盐")]
    if not solv or not salts or not all(c["单位"] == "摩尔比" for c in comps):
        return None
    try:
        solvating = sum(float(c["用量"]) for c in solv if c["角色"] == "溶剂化溶剂")
        return {c["名称"]: float(c["用量"]) / solvating for c in salts}
    except ValueError:
        return None


def unify(f):
    """🧮 换算：只做不需要额外假设的换算；返回文字。"""
    comps = f["组分"]
    if not [c for c in comps if c["角色"] in SOLVENT_ROLES]:
        return "成分未报告"
    if any(c["用量"] == "未报告" for c in comps):
        return "未算：有组分用量未报告"
    out = []
    fr = vol_fractions(f)
    if fr and len(fr) > 1:
        out.append("溶剂体积分数 " + "、".join(f"{base(n)} {v * 100:.0f}%" for n, v in fr.items()))
    if all(c["单位"] == "摩尔比" for c in comps):
        solv = [c for c in comps if c["角色"] in SOLVENT_ROLES]
        tot = sum(float(c["用量"]) for c in solv)
        out.append("溶剂摩尔分数 " + "、".join(f"{base(c['名称'])} {float(c['用量']) / tot * 100:.1f}%" for c in solv))
    sr = salt_ratio(f)
    if sr:
        out.append("；".join(f"{n} : 溶剂化溶剂 = {v:.3f}" for n, v in sr.items()))
    if has_m(f):
        out.append(M_NOTE)
    return "🧮 " + "；".join(out) if out else "—"


def coord_info(f):
    """返回 (电解液类别, 尺子文字, 尺子数值或 None)。
    尺子 = 每个主盐 Li⁺ 分到的配位溶剂分子数：配位溶剂的物质的量 ÷ 主盐的物质的量。
    摩尔比：精确；M（mol/L）：按“溶剂体积 ≈ 溶液体积”粗算；小写 m 不换算。稀释剂、LiNO₃ 不计；含配位能力不明的溶剂不进尺子。"""
    comps = f["组分"]
    solv = [c for c in comps if c["角色"] in SOLVENT_ROLES]
    if not solv:
        return "成分未报告", "未算（成分未报告）", None
    info = []
    for c in solv:
        names = base(c["名称"]).split("/")
        ss = [SOLV.get(n) for n in names]
        ab = {s.get("配位能力") if s else "不明" for s in ss}
        info.append((c, ss[0] if len(ss) == 1 else None, ab.pop() if len(ab) == 1 else "不明"))
    unknown = [base(c["名称"]) for c, s, ab in info if ab == "不明"]
    coord = [base(c["名称"]).replace("/", "+") + ("（比例未报告）" if "/" in base(c["名称"]) else "")
             for c, s, ab in info if ab == "配位"]
    diluent = [base(c["名称"]) for c, s, ab in info if ab == "不配位"]
    salts = [c for c in comps if c["角色"] == "盐"]
    salt = salts[0] if len(salts) == 1 else None
    conc = parse(salt["用量"])[0] if salt and parse(salt["用量"]) and unit_key(salt["单位"]) in ("M", "m") else None
    if unknown:
        cat = f"含配位能力不明的溶剂（{'、'.join(unknown)}）"
    elif diluent:
        cat = f"加稀释剂（{'、'.join(diluent)}）"
    elif conc is not None and conc >= 3:
        cat = "高浓度"
    else:
        cat = "常规"
    cat += f"｜配位溶剂：{'+'.join(coord) or '无'}"
    if salt is None:
        return cat, "未算（主盐不唯一或未报告）", None
    unit, exact = unit_key(salt["单位"]), True
    try:
        if unit == "摩尔比" and all(c["单位"] == "摩尔比" for c in solv):
            n = sum(float(c["用量"]) for c, s, ab in info if ab == "配位") / float(salt["用量"])
        elif unit == "m":
            return cat, f"未算（{M_NOTE}）", None
        elif unit == "M" and conc:
            fr = vol_fractions(f)
            if fr is None:
                return cat, "未算（溶剂比例未报告）", None
            tot = 0.0
            for c, s, ab in info:
                if ab != "配位":
                    continue
                if not s or not s.get("密度") or not s.get("分子量"):
                    return cat, f"未算（{base(c['名称'])} 缺密度或分子量）", None
                tot += fr[c["名称"]] * 1000 * s["密度"] / s["分子量"]
            n, exact = tot / conc, False
        else:
            return cat, "未算（换算要额外假设）", None
    except ValueError:
        return cat, "未算（用量未报告）", None
    if unknown:
        return cat, f"未进尺子（{'、'.join(unknown)} 配位能力不明）；只计 {'+'.join(coord)}：{n:.1f}", None
    return cat, f"{n:.1f}" + ("" if exact else "（粗算）"), n


# ---------- 6. 测试条件 ----------
def np_values(t, b_code):
    li, s = num(t.get("锂厚")), num(t.get("硫载量"))
    design = f"{li * 0.2 / (s * 1.675):.2f}" if li and s else "缺失"
    measured = "缺失"
    if li and s:
        for r in t["结果"]:
            if r["指标"] == "首圈容量" and num((r.get("数值") or {}).get(b_code)):
                measured = f"{li * 0.2 / (num(r['数值'][b_code]) * s * 0.001):.2f}"
                break
    return design, measured


def compact_cond(t):
    out = [t["电池"]]
    s, es, li = t.get("硫载量"), t.get("E/S"), t.get("锂厚")
    if reported(s):
        out.append(f"S {show(s)} mg/cm²")
    if reported(es):
        out.append(f"E/S {show(es)} {t.get('E/S单位', '')}".strip())
    if num(li):
        out.append(f"Li {li} µm")
    out.append(t["倍率"])
    if t.get("温度") not in (None, "", "未报告"):
        out.append(t["温度"])
    out.append(f"LiNO₃ {show(t.get('LiNO₃'))}")
    return "，".join(out)


def inconsistent(t, a, b):
    out = []
    for k in ("E/S", "硫载量", "锂厚", "温度", "倍率"):
        v = t.get(k)
        if isinstance(v, dict) and a in v and b in v and v[a] != v[b]:
            out.append(f"{k} A {v[a]} vs B {v[b]}")
    return "；".join(out)


# ---------- 7. 方向（含固定规则） ----------
def fig_ok(res):
    """读图值只用锚点核对通过的；没有锚点或核不上，整张图不用。"""
    return res.get("来源") != "读图" or str(res.get("锚点", "")).startswith("通过")


def pair_result(res, a, b):
    """返回 (↑/↓/差异小/无法判断, 变好/变坏/差异小/不判（理由）, 来源, 说明)；没有这对的数据或读图不用时返回 None。"""
    if not fig_ok(res):
        return None
    better = better_dir(res)
    is_ce = res["指标"] in CE_METRICS
    vals = res.get("数值") or {}
    unit = res.get("单位", "")
    if a in vals and b in vals:
        d = compare(vals[a], vals[b], is_ce)
        detail = f"{vals[a]} vs {vals[b]} {unit}".strip()
        if is_ce and parse(vals[a]) and parse(vals[b]) and unit == "%":
            detail += f"，差 {parse(vals[a])[0] - parse(vals[b])[0]:+.1f} 个百分点"
        gb = "不判（A、B 的数来自不同图，原文没说是同一批电池）" if res.get("跨图") else good_bad(d, better)
        return d, gb, res.get("来源", "数值"), detail
    for q in res.get("作者定性", []):
        if q["A"] == a and q["B"] == b:
            gb = "不判（作者说法前后不一）" if q.get("说法不一") else good_bad(q["方向"], better)
            return q["方向"], gb, "作者定性", "作者定性" + (f"：{q['说明']}" if q.get("说明") else "")
    return None


WORDS = {"Li⁺ 溶剂化": ("更强", "更弱"), "多硫化物溶剂化": ("更强", "更弱"), "多硫化物团簇化": ("增多", "减少")}


def prop_vs(p, a, b):
    hi, lo = WORDS.get(p["性质"], ("更高", "更低"))
    if "数值" in p:
        va, vb = p["数值"].get(a), p["数值"].get(b)
        if va is not None and vb is not None:
            return {"↑": hi, "↓": lo, "差异小": "相当"}.get(compare(va, vb, False), "无法判断")
        return "无法判断（对照未测）" if va is not None else None
    if "排序" in p:
        order = [x.split("≈") for x in p["排序"]]
        ra = next((i for i, x in enumerate(order) if a in x), None)
        rb = next((i for i, x in enumerate(order) if b in x), None)
        if ra is None or rb is None:
            return None
        return "相当" if ra == rb else (lo if ra > rb else hi)
    d = p.get("方向")
    if isinstance(d, dict) and a in d:
        return {"↑": hi, "↓": lo}.get(d[a], d[a])
    return None


STEP_WORD = {"↑": 1, "↓": -1, "差异小": 0}


def shape(raw, better="↑", is_ce=False):
    """沿梯度相邻两档比较 → 形状（按“好”的方向）。raw 为按梯度顺序、已去掉缺失的原始值；库仑效率不设 5% 阈值。"""
    sg = 1 if better == "↑" else -1
    steps = []
    for x, y in zip(raw, raw[1:]):
        d = compare(y, x, is_ce)
        steps.append(STEP_WORD[d] * sg if d in STEP_WORD else 0)
    nz = [s for s in steps if s]
    if not nz:
        return "持平"
    if all(s >= 0 for s in steps):
        return "逐级升" if all(s > 0 for s in steps) else "升"
    if all(s <= 0 for s in steps):
        return "逐级降" if all(s < 0 for s in steps) else "降"
    changes = sum(1 for a, b in zip(nz, nz[1:]) if a != b)
    if changes == 1:
        return "先升后降" if nz[0] > 0 else "先降后升"
    return "有起伏"


PERF_WORD = {"逐级升": "逐级变好", "升": "随梯度变好", "逐级降": "逐级变坏", "降": "随梯度变坏",
             "先升后降": "先变好后变坏", "先降后升": "先变坏后变好", "有起伏": "有起伏", "持平": "各档差异小"}
PROP_WORD = {"逐级升": "逐级升高", "升": "随梯度升高", "逐级降": "逐级降低", "降": "随梯度降低",
             "先升后降": "先升后降", "先降后升": "先降后升", "有起伏": "有起伏", "持平": "各档相近"}


def best_of(order, vals, better):
    cand = [(c, parse(vals[c])) for c in order if c in vals and parse(vals[c])]
    best = [c for c, pc in cand if all(
        (pc[0] > po[1] if better == "↑" else pc[1] < po[0]) or o == c for o, po in cand)]
    return best[0] if len(best) == 1 else None


def clean(x):
    return re.sub(r"（🔶.*?）", "", x).strip()


def strip_marks(s):
    return re.sub(r"[🔶⚠️🧮]\s*", "", s).strip()


def check_landing(P, x, a, b):
    """机制落到的表现：到本篇同一测试的结果里核对 A 相对 B 的好坏。返回 (是否一致, 说明)。"""
    metrics = [x["指标"]] if x.get("指标") else PERF[x["表现"]]
    want = {"变好": "变好", "变坏": "变坏", "相当": "差异小"}[x["方向"]]
    hits = []
    t = next((t for t in P["测试"] if t["编号"] == x["测试"]), None)
    for r in (t["结果"] if t else []):
        if not any(r["指标"].startswith(m) for m in metrics):
            continue
        pr, inv = pair_result(r, a, b), False
        if pr is None:
            pr, inv = pair_result(r, b, a), True
        if pr is None or pr[1].startswith("不判"):
            continue
        gb = {"变好": "变坏", "变坏": "变好"}.get(pr[1], pr[1]) if inv else pr[1]
        hits.append((gb, f"{r['指标']}（{clean(r['取值点'])}）{gb}"))
    if not hits:
        return None, "本篇这一测试没有可核对的数据"
    ok = all(h[0] == want for h in hits)
    return ok, f"作者说 {x['表现']}{x['方向']}；数据：" + "；".join(h[1] for h in hits)


# ---------- 8. 读入并展开 ----------
papers = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(JSON_DIR.glob("S*.json"))]
SHEET = {"结论": "①一句话结论", "要看": "②要你看", "配方": "明细-配方", "测试": "明细-测试", "方向": "明细-方向",
         "梯度": "明细-梯度趋势", "机制": "明细-机制", "溶剂": "明细-溶剂表", "记录": "程序核对记录"}
pages = {k: [] for k in SHEET}
stats = {"原句": 0, "原句找到": 0, "数值": 0, "数值在原句": 0, "反算": 0}


def add(page, row):
    pages[page].append(row)
    return len(pages[page]) + 1          # Excel 行号（第 1 行是表头）


def where(page, row):
    return f"{SHEET[page]} 第 {row} 行" if row else SHEET[page]


def log(pid, page, row, content, action, basis=""):
    pages["记录"].append([pid, where(page, row), content, action, basis])


def todo(pid, page, row, content, why, basis):
    pages["要看"].append([len(pages["要看"]) + 1, pid, where(page, row), content, why, basis, ""])


def notes(d, skip=()):
    """提取时写下的原文疑点（🔶/⚠️）：按规则处理，记进程序核对记录。"""
    out = []
    for k, v in d.items():
        if k in skip:
            continue
        s = v if isinstance(v, str) else json.dumps(v, ensure_ascii=False)
        if "🔶" in s or "⚠️" in s:
            out.append(f"{k}：{strip_marks(s)}")
    return "；".join(out)


for P in papers:
    pid, stem, B = P["文献编号"], P["全文缓存"], P["对照"]
    F = {f["代号"]: f for f in P["配方"]}
    g = P["梯度"][0] if P.get("梯度") else None
    order = g["顺序"] if g else []
    best_codes = {c["最优"] for c in P.get("设计主张", [])}
    n_todo0 = len(pages["要看"])

    def disp(code):
        if code in order:
            v = g["取值"][order.index(code)]
            return v if v.endswith("%") else code
        return code

    # 明细-配方
    ruler = {}
    for f in P["配方"]:
        c = f["代号"]
        role = "对照" if c == B else ("作者最优" if c in best_codes or "最优" in f["作者原名"] else
                                     ("梯度档" if c in order else "其他"))
        cat, rtxt, rval = coord_info(f)
        ruler[c] = (rtxt, rval)
        comps = f["组分"]
        q = f.get("原句") or {}
        row = add("配方", [pid, c, f["作者原名"], role, g["取值"][order.index(c)] if c in order else "—",
                          "；".join(f"{x['名称']} {x['用量']} {x['单位']}" for x in comps if x["角色"] in SOLVENT_ROLES) or "未报告",
                          "；".join(f"{x['名称']} {x['用量']} {x['单位']}" for x in comps if x["角色"].startswith("盐")) or "—",
                          "；".join(f"{x['名称']} {x['用量']} {x['单位']}" for x in comps if x["角色"] == "添加剂") or "—",
                          unify(f), cat, rtxt, q.get("原句", "")])
        if q.get("原句") and not found(stem, q.get("部分", "main"), q["原句"]):
            log(pid, "配方", row, f"配方 {c} 原句", "原句在全文里找不到 → 正式跑时让模型重提；仍找不到记缺失", q["原句"])
            print("  ❌ 原句找不到：", pid, c)
        n = notes(f, skip=("原句",))
        if n:
            log(pid, "配方", row, f"配方 {c}", "按规则：原文没写的记“原文未写”，结论里照原文写", n)
        if has_m(f):
            log(pid, "配方", row, f"配方 {c} 用量单位", f"按规则：{M_NOTE}，尺子不算", comp_text(f))

    # 明细-测试
    for t in P["测试"]:
        design, measured = np_values(t, B)
        li = t.get("锂厚")
        li_txt = f"{li}（设计 N/P {design}；实测 N/P {measured}）" if num(li) and num(t.get("硫载量")) else show(li)
        quotes = t.get("条件原句", [])
        joined = " ".join(x["原句"] for x in quotes)
        row = add("测试", [pid, t["编号"], t["名称"], t["电池"], show(t.get("硫载量")),
                          f"{show(t.get('E/S'))} {t.get('E/S单位', '')}".strip(), li_txt,
                          f"{t['倍率']}（1C {t.get('1C定义', '未报告')}；化成 {t.get('化成', '未报告')}）",
                          show(t.get("温度")), show(t.get("LiNO₃")), "、".join(t["配方"]),
                          "\n".join(f"{x['原句']}（{x.get('位置', '')}）" for x in quotes)])
        bad_q = [x["原句"] for x in quotes if not found(stem, x.get("部分", "main"), x["原句"])]
        if bad_q:
            log(pid, "测试", row, f"测试 {t['编号']} 条件原句", "原句在全文里找不到 → 正式跑时重提", "\n".join(bad_q))
            print("  ❌ 条件原句找不到：", pid, t["编号"])
        paper_q = " ".join(x["原句"] for t2 in P["测试"] for x in t2.get("条件原句", []))
        for k in ("硫载量", "E/S", "锂厚"):
            v = t.get(k)
            if not reported(v) or isinstance(v, dict):
                continue
            mm = k == "锂厚" and num(v) and re.search(rf"(?<![\d.]){v / 1000:g}(?:\.0)?mm", norm(joined + " " + paper_q))
            if num_in_quote(v, joined):
                continue
            if str(t.get(f"{k}说明", "")).startswith("🧮"):
                log(pid, "测试", row, f"测试 {t['编号']} {k} {v}", f"🧮 推算：{strip_marks(t[f'{k}说明'])}", joined)
            elif num_in_quote(v, paper_q) or mm:
                log(pid, "测试", row, f"测试 {t['编号']} {k} {v}",
                    "数在本篇的条件原句里（同一套电池共用条件" + ("；原文按 mm 写" if mm else "") + "）", "")
            else:
                log(pid, "测试", row, f"测试 {t['编号']} {k} {v}", "这个数不在本篇条件原句里 → 正式跑时重提", joined)
                print("  ⚠ 条件数不在原句：", pid, t["编号"], k, v)
        n = notes(t, skip=("结果", "条件原句"))
        if n:
            log(pid, "测试", row, f"测试 {t['编号']}", "按规则：原文没写清的照原文记、标“推定/未写”，不交人判", n)

    pairs = []
    for t in P["测试"]:
        for pr in (t.get("比较对") or [[a, B] for a in t["配方"] if a != B and B in t["配方"]]):
            if pr not in pairs:
                pairs.append(pr)

    # 明细-方向：性质行 + 表现行
    first_row = {}
    for a, b in pairs:
        for p in P["性质"]:
            d = prop_vs(p, a, b)
            if d is None:
                continue
            vals = p.get("数值") or {}
            row = add("方向", [pid, f"性质测量（{p.get('测量体系', p.get('依据', ''))}）", f"{a} vs {b}", "性质",
                              p["性质"] + (f"（{p['对象']}）" if p.get("对象") else ""),
                              show(vals.get(a, "—")), show(vals.get(b, "—")), d, p.get("位置", "")])
            first_row.setdefault(("性质", p["性质"]), row)
        for t in P["测试"]:
            if a not in t["配方"] or b not in t["配方"]:
                continue
            bad = inconsistent(t, a, b)
            for r in t["结果"]:
                pr = pair_result(r, a, b)
                if pr is None:
                    continue
                d, gb, src, detail = pr
                vals = r.get("数值") or {}
                fig = r.get("来源") == "读图"
                direction = "不判（A、B 条件不同）" if bad else gb + ("（作者定性）" if src == "作者定性" else "")
                row = add("方向", [pid, f"{t['编号']} {t['名称']}", f"{a} vs {b}", "表现",
                                  f"{r['指标']}（{clean(r['取值点'])}）", show(vals.get(a, "—")), show(vals.get(b, "—")),
                                  direction, r.get("位置", "") + (f"；读图 {r.get('图', '')}" if fig else "")])
                key = (t["编号"], r["指标"], r["取值点"])
                first_row.setdefault(key, row)
                if bad and key + ("bad",) not in first_row:
                    first_row[key + ("bad",)] = row
                    log(pid, "方向", row, f"{t['编号']} {a} vs {b}", "按规则：A、B 条件不同，不算同一测试，不判", bad)
    # 结果的数字核对、读图、规则处理（每个结果记一次）
    for t in P["测试"]:
        for r in t["结果"]:
            row = first_row.get((t["编号"], r["指标"], r["取值点"]))
            label = f"{t['编号']} {r['指标']}（{clean(r['取值点'])}）"
            vals = r.get("数值") or {}
            if r.get("来源") == "读图":
                log(pid, "方向", row, f"{label} 读图值 {json.dumps(vals, ensure_ascii=False)}",
                    "锚点核对通过，用于判方向" if fig_ok(r) else "没有可核的锚点 → 按规则这张图不用，记“仅图未读”",
                    f"{r.get('图', '')}；{r.get('锚点', '无锚点')}")
                continue
            if r.get("原句") and not found(stem, r.get("部分", "main"), r["原句"]):
                log(pid, "方向", row, f"{label} 原句", "原句在全文里找不到 → 正式跑时重提", r["原句"])
                print("  ❌ 原句找不到：", pid, label)
            for c, v in vals.items():
                stats["数值"] += 1
                if (r.get("计算") or {}).get(c):
                    stats["反算"] += 1
                    log(pid, "方向", row, f"{label} {c} = {v}", f"🧮 反算：{r['计算'][c]}", r.get("原句", ""))
                elif num_in_quote(v, r.get("原句", "")):
                    stats["数值在原句"] += 1
                else:
                    log(pid, "方向", row, f"{label} {c} = {v}", "这个数不在它挂的原句里 → 正式跑时重提", r.get("原句", "") or "（没挂原句）")
                    print("  ⚠ 数值不在原句：", pid, label, c, v)
            if r.get("跨图"):
                log(pid, "方向", row, label, "按规则：A、B 的数来自不同图、原文没说同一批电池，不判", r.get("备注", ""))
            for q in r.get("作者定性", []):
                if q.get("说法不一"):
                    log(pid, "方向", row, f"{label} {q['A']} vs {q['B']}", "按规则：作者说法前后不一，不判（有锚点的图可读时改用读图）",
                        strip_marks(q.get("说明", "")))
            n = notes(r, skip=("原句", "作者定性"))
            if n:
                log(pid, "方向", row, label, "按规则处理，不交人判", n)
    for p in P["性质"]:
        row = first_row.get(("性质", p["性质"]))
        if p.get("原句") and not found(stem, p.get("部分", "main"), p["原句"]):
            log(pid, "方向", row, f"性质 {p['性质']} 原句", "原句在全文里找不到 → 正式跑时重提", p["原句"])
            print("  ❌ 性质原句找不到：", pid, p["性质"])
        for c, v in (p.get("数值") or {}).items():
            stats["数值"] += 1
            if num_in_quote(v, p.get("原句", "")):
                stats["数值在原句"] += 1
            else:
                log(pid, "方向", row, f"性质 {p['性质']} {c} = {v}", "这个数不在它挂的原句里 → 正式跑时重提", p.get("原句", ""))
                print("  ⚠ 性质数值不在原句：", pid, p["性质"], c, v)
        n = notes(p, skip=("原句",))
        if n:
            log(pid, "方向", row, f"性质 {p['性质']}", "按规则处理，不交人判", n)

    # 明细-梯度趋势 + ① 表现行
    test_lines = []
    for t in P["测试"]:
        segs, seen = [], set()           # 每项：[指标, [取值点], 是否读图, 分隔符, 结论]
        tp = t.get("比较对") or [[a, B] for a in t["配方"] if a != B and B in t["配方"]]
        bad_all = all(inconsistent(t, a, b) for a, b in tp) if tp else False
        for r in t["结果"]:
            if not fig_ok(r):
                continue
            vals = r.get("数值") or {}
            better = better_dir(r)
            fig = r.get("来源") == "读图"
            point = clean(r["取值点"])
            metric = f"{r['指标']}（{point}）" + ("（读图）" if fig else "")
            src = r.get("位置", "") + (f"；读图 {r.get('图', '')}" if fig else "")
            present = [c for c in order if c in vals and parse(vals[c])]
            tail, sep = None, "："
            if len(present) >= 3:
                sh = shape([vals[c] for c in present], better, r["指标"] in CE_METRICS)
                bc = best_of(present, vals, better)
                absent = [disp(c) for c in order if c not in present]
                tail, sep = PERF_WORD[sh] + (f"，{disp(bc)} 最好" if bc and sh in ("先升后降", "先降后升", "有起伏") else "") + \
                    (f"（{'、'.join(absent)} 未报告）" if absent else ""), ""
                add("梯度", [pid, f"{t['编号']} {t['名称']}", metric, " / ".join(show(vals.get(c, "—")) for c in order) + f" {r.get('单位', '')}",
                            PERF_WORD[sh], disp(bc) if bc else "不唯一", src])
            else:
                items = [(a, b, pair_result(r, a, b)) for a, b in tp if not inconsistent(t, a, b)]
                items = [x for x in items if x[2]]
                seg = []
                groups = []              # 相对对照的各档：按（好坏, 是否作者定性）归组，保持顺序
                for a, b, pr in items:
                    if b != B:
                        continue
                    key = (pr[1], pr[2] == "作者定性" and not pr[1].startswith("不判"))
                    if groups and groups[-1][0] == key:
                        groups[-1][1].append(disp(a))
                    else:
                        groups.append((key, [disp(a)]))
                n_ctrl = sum(len(v) for _, v in groups)
                if len(groups) == 1 and n_ctrl > 1:
                    seg.append(f"各档都{groups[0][0][0]}" + ("（作者定性）" if groups[0][0][1] else ""))
                elif groups and len({k[0] for k, _ in groups}) == 1 and n_ctrl > 1:
                    seg.append(f"各档都{groups[0][0][0]}（部分为作者定性）")
                elif groups:
                    seg.append("，".join(f"{'、'.join(v)} {k[0]}" + ("（作者定性）" if k[1] else "") for k, v in groups))
                for a, b, pr in items:
                    if b != B:
                        seg.append(f"{disp(a)} 比 {disp(b)} {pr[1]}" + ("（作者定性）" if pr[2] == "作者定性" and not pr[1].startswith("不判") else ""))
                if seg:
                    tail = "；".join(seg)
                if r.get("排序"):
                    tail = (tail or "") + f"（作者排序：{r['排序']}）"
                    add("梯度", [pid, f"{t['编号']} {t['名称']}", metric, f"作者排序：{r['排序']}", "见作者排序", "—", src])
                if r.get("最优"):
                    tail = (tail or "") + f"，作者称 {disp(r['最优']['配方'])} 最好"
                    add("梯度", [pid, f"{t['编号']} {t['名称']}", metric, r["最优"]["说明"], "见作者说法", disp(r["最优"]["配方"]), src])
            if tail:
                if segs and segs[-1][0] == r["指标"] and segs[-1][2] == fig and segs[-1][3:] == [sep, tail]:
                    segs[-1][1].append(point)            # 同一指标、同样结论的几个取值点合并
                else:
                    segs.append([r["指标"], [point], fig, sep, tail])
                seen.add(r["指标"])
        parts = [f"{n}（{'、'.join(ps)}）" + ("（读图）" if fig else "") + f"{sep}{tail}" for n, ps, fig, sep, tail in segs]
        if bad_all and tp:
            parts = [f"不判（A、B 条件不同：{inconsistent(t, *tp[0])}）"]
        if not parts:
            continue
        miss = []
        if t["电池"].startswith("Li‖S") and any(k in t["名称"] for k in ("循环", "软包", "贫液", "首两圈")):
            miss = [lab for lab, grp in CORE if not any(any(m.startswith(x) for x in grp) for m in seen)]
        test_lines.append(f"{t['名称']}｜{compact_cond(t)}：{'；'.join(parts)}" + (f"（未报告：{'、'.join(miss)}）" if miss else ""))

    # 明细-机制：一句一行
    def strength(code):
        return order.index(code) if code in order else None

    def side(m):
        props = [p for p in m["性质"] if p["名称"] in SOLVATION]
        outs = [x["方向"] for x in m["落到"] if x["方向"] != "相当"]
        if not props or not outs:
            return "其他"
        ra, rb = strength(m["A"]), strength(m["B"])
        if ra is not None and rb is not None:
            a_strong = ra < rb
        else:
            p0 = props[0]
            a_strong = p0["方向"] in ("更强", "更高") or (p0["名称"] == "多硫化物团簇化" and p0["方向"] == "减少")
        return "强" if a_strong == (outs[0] == "变坏") else "弱"

    for m in P["机制"]:
        for p in m["过程"]:
            if p["大类"] not in CATS or p["方向"] not in CATS[p["大类"]]:
                raise KeyError(f"{pid}-{m['编号']}：过程大类或方向不在词表里：{p}")
        land = "\n".join(f"{x['测试']} {x.get('指标', x['表现'])} {x['方向']}" for x in m["落到"]) or "（不落到电池测试）"
        mentioned = {c for q in m["原句"] for c in q["提到"]}
        miss_cat = [p["大类"] for p in m["过程"] if p["大类"] not in mentioned]
        only_cited = all(q["来源"] == CITED for q in m["原句"])
        m["_side"] = side(m)
        first_row_m = None
        for i, q in enumerate(m["原句"]):
            if q["来源"] not in SRC:
                raise KeyError(f"{pid}-{m['编号']}-{q['id']}：来源不在词表里：{q['来源']}")
            ok = found(stem, q.get("部分", "main"), q["原句"])
            stats["原句"] += 1
            stats["原句找到"] += ok
            first = i == 0
            row = add("机制", [pid, m["编号"] if first else "",
                              (f"{m['A']} vs {m['B']}" + ("（作者说法覆盖全梯度）" if m.get("梯度段") == "全梯度" else "")) if first else "",
                              ("；".join(f"{p['名称']} {p['方向']}" for p in m["性质"]) or "—") if first else "",
                              ("；".join(f"{p['大类']} {p['方向']}" for p in m["过程"]) or "—") if first else "",
                              land if first else "",
                              (m["概括"] + ("（只有转述句支撑，不算本篇结论）" if only_cited else "")) if first else "",
                              q["id"], q["原句"].replace("‖", " … "), "、".join(q["提到"]) or "—", q["来源"]])
            first_row_m = first_row_m or row
            if not ok:
                log(pid, "机制", row, f"{m['编号']} {q['id']} 原句", "原句在全文里找不到 → 正式跑时重提", q["原句"])
                print("  ❌ 机制原句找不到：", pid, m["编号"], q["id"])
            sc = source_check(q["来源"], source_flags(stem, q.get("部分", "main"), q["原句"]))
            if sc:
                log(pid, "机制", row, f"{m['编号']} {q['id']} 来源（AI 判：{q['来源']}）",
                    "留到分析阶段：这条说法进入共识/矛盾清单时再看原句", sc)
            mc = mention_check(q["提到"], q["原句"])
            if mc:
                log(pid, "机制", row, f"{m['编号']} {q['id']} 提到的大类（AI 标：{'、'.join(q['提到']) or '无'}）",
                    "留到分析阶段：这条说法进入共识/矛盾清单时再看原句", mc)
        if miss_cat:
            log(pid, "机制", first_row_m, f"{m['编号']} 过程大类", "留到分析阶段", f"说法里的 {'、'.join(miss_cat)} 在它的原句里都没提到")
        for x in m["落到"]:
            ok, msg = check_landing(P, x, m["A"], m["B"])
            if ok is False:
                todo(pid, "机制", first_row_m, f"{m['编号']}：{m['概括']}", "作者说的变好/变坏和本篇同一测试的数据对不上——可能是提错了，也可能是作者说错了", msg)
            else:
                log(pid, "机制", first_row_m, f"{m['编号']} 落到 {x['测试']} {x['表现']}",
                    "篇内核对：一致" if ok else "篇内核对：没有可核对的数据", msg)
    for c in P.get("设计主张", []):
        q = c["原句"]
        ok = found(stem, q.get("部分", "main"), q["原句"])
        stats["原句"] += 1
        stats["原句找到"] += ok
        row = add("机制", [pid, f"设计主张（{c['类型']}）", f"最优 {c['最优']}", "—", "—", "—",
                          f"作者主张{c['类型']}，最优 {disp(c['最优'])}，依据 {'、'.join(c['两侧依据'])}", "—", q["原句"],
                          "、".join(q.get("提到", [])) or "—", "本文结论"])
        if not ok:
            log(pid, "机制", row, f"设计主张 {c['类型']} 原句", "原句在全文里找不到 → 正式跑时重提", q["原句"])
            print("  ❌ 设计主张原句找不到：", pid, c["类型"])
        mc = mention_check(q.get("提到", []), q["原句"])
        if mc:
            log(pid, "机制", row, f"设计主张 {c['类型']} 提到的大类", "留到分析阶段", mc)

    # ① 一句话结论（一篇一格，四行）
    def ruler_seq():
        seq = [ruler[c][1] for c in order]
        if not any(v is not None for v in seq):
            return ""
        txt = " / ".join(f"{v:.1f}" if v is not None else "—" for v in seq)
        why = sorted({re.sub(r"^未算（|^未进尺子（|）.*$", "", ruler[c][0]) for c in order if ruler[c][1] is None})
        return f"；每个主盐 Li⁺ 分到的配位溶剂分子数🧮：{txt}" + (f"（— = {'；'.join(why)}）" if why else "")

    def salt_seq():
        rs = [salt_ratio(F[c]) for c in order]
        if sum(r is not None for r in rs) < 2:
            return ""
        names = sorted({n for r in rs if r for n in r})
        main = [n for n in names if not n.startswith("LiNO")] or names
        vals = [r.get(main[0]) if r else None for r in rs]
        known = [v for v in vals if v is not None]
        if max(known) / min(known) < 1.05:
            return ""
        return f"；{main[0]} : 溶剂化溶剂🧮：" + " / ".join(f"{v:.3f}" if v is not None else "—" for v in vals)

    line1 = f"【{pid}】对照：{comp_text(F[B])}（{P['对照类型']}）"
    if g:
        line1 += f"；梯度：{g['名称']} " + " / ".join(disp(c) for c in order) + \
                 ("（溶剂化由强到弱）" if g.get("溶剂化方向") == "由强到弱" and "由强到弱" not in g["名称"] else "")
    others = [c for c in F if c != B and c not in order]
    if others:
        line1 += "；另有：" + "、".join(f"{c}（{F[c]['作者原名']}）" for c in others)
    line1 += ruler_seq() + salt_seq()

    props, missing = [], []
    for p in P["性质"]:
        name = p["性质"] + (f"（{p['对象']}）" if p.get("对象") else "")
        vals = p.get("数值")
        if vals:
            present = [c for c in order if c in vals and parse(vals[c])]
            if len(present) >= 3:
                props.append(f"{name} " + " / ".join(show(vals.get(c, "—")) for c in order) +
                             f" {p.get('单位', '')}，{PROP_WORD[shape([vals[c] for c in present])]}")
            else:
                props.append(f"{name} " + "、".join(f"{disp(c)} {v}" for c, v in vals.items()) + f" {p.get('单位', '')}")
        elif p.get("排序"):
            props.append(f"{name}（{p.get('排序含义', '')}）：{' > '.join(p['排序'])}")
        elif isinstance(p.get("方向"), dict):
            props.append(f"{name}：" + "、".join(f"{disp(c)} " + {"↑": "更高", "↓": "更低"}.get(d, d)
                                                for c, d in p["方向"].items()) + f"（相对 {disp(B)}）")
        elif p.get("方向") in (None, "", "未报告"):
            missing.append(p["性质"])
        else:
            props.append(f"{name}：{p['方向']}")
    line2 = "性质：" + "；".join(props) + (f"；未测：{'、'.join(missing)}" if missing else "") + \
            (f"。说明：{g['说明']}" if g and g.get("说明") else "")
    line3 = "表现：\n" + "\n".join(f"  {chr(0x2460 + i)} {x}" for i, x in enumerate(test_lines))
    mech = []
    for k, label in (("强", "强溶剂化一侧的短板"), ("弱", "弱溶剂化一侧的短板"), ("其他", "其他说法")):
        ms = [m for m in P["机制"] if m["_side"] == k]
        if not ms:
            continue
        cats = []
        for m in ms:
            for p in m["过程"]:
                if p["大类"] not in cats:
                    cats.append(p["大类"])
        mech.append(f"  {label}" + (f"〔{'、'.join(cats)}〕" if cats else "") + "：" +
                    "；".join(f"{m['概括']}（{m['编号']}）" for m in ms))
    if P.get("设计主张"):
        mech.append("  作者主张：" + "；".join(f"{c['类型']}，最优 {disp(c['最优'])}" for c in P["设计主张"]))
    line4 = "机制：\n" + "\n".join(mech)
    pages["结论"].append([pid, "\n".join((line1, line2, line3, line4)), len(pages["要看"]) - n_todo0, "", ""])

dens = [s for s in solvents["溶剂"] if s.get("密度")]
log("全部", "溶剂", None, "密度（" + "、".join(s["简称"] for s in dens) + "）", "尺子用到的密度，已按供应商产品页核对",
    "；".join(s.get("物性来源", f"{s['简称']} {s['密度']} g/mL") for s in dens))
for s in solvents["溶剂"]:
    add("溶剂", [s["简称"], s["全称"], s.get("配位能力", "未定"), s.get("配位依据", ""), s.get("分子量", ""), s.get("密度", ""),
                s.get("物性来源", ""), s["DN"], s["介电常数"], s["来源"]])
if not pages["要看"]:
    pages["要看"].append(["—", "—", "—", "本批没有要你看的条目", "作者说法与本篇数据都对得上；原文没写清的已按规则处理（见“程序核对记录”）", "", ""])

# ---------- 9. 说明 ----------
guide = [
    ["这是什么", "论断式提取表样 v4（10-24）：S11、S21、S33、S39 四篇（都在调试集里）。AI 辅助提取与核对，用户尚未逐项人工确认。"],
    ["你只看两页", "①一句话结论：一篇一格。测试集里在“错提”“漏提”两列写问题，没错留空。②要你看：作者说的变好/变坏和本篇数据对不上的条目（本批 0 条）。其余各页（明细-…、程序核对记录）不用看，要查原文时再翻。"],
    ["错提、漏提怎么写", "错提：结论里哪条不对，写编号（如“② 寿命”“M2”）和一句理由；包括“原句是真的，但原文没这个意思”（概括夸大、把别人的结论当成本文的、数挂错了配方或测试）。漏提：原文有、结论里没写的，写一句。没写的就算对。召回率、精确率由 Claude 按你写的统计，你确认。"],
    ["程序已经替你核过的", "每个数字都在它挂的原句里（英文数字词也认）、每句原句都在全文里；变好/变坏、梯度趋势、N/P、尺子由程序算；读图只用锚点（正文写明的数）核对通过的图；原句来源和“提到的大类”用关键词核过，对不上的留到分析阶段再看。"],
    ["原文没写清的固定规则", "原文没写 → 记“原文未写”；A、B 条件不同 → 不判；A、B 的数来自不同图且原文没说同一批 → 不判；作者说法前后不一 → 不判（有锚点的图可读时改用读图）；单位含义不明（如小写 m）→ 不换算、尺子不算；从“长 25%”“超过 6 倍”这类说法反算的数 → 写明公式照用；读图没有锚点或核不上 → 整张图不用。全部记在“程序核对记录”。"],
    ["一句话结论四行", "① 对照（是什么、对照类型）与梯度、尺子 ② 性质（只记高低，不判好坏）③ 每个测试一行：条件数值 + 各指标变好/变坏（梯度论文写趋势和最优档）；Li‖S 循环类测试没报告的核心指标写“未报告” ④ 机制：哪一侧的短板、在哪类过程（正极反应 / 多硫化物穿梭 / 负极锂 / 电解液消耗，可多标），以及作者主张（折中/解耦）。"],
    ["尺子", "每个主盐 Li⁺ 分到的配位溶剂分子数 = 配位溶剂的物质的量 ÷ 主盐的物质的量。稀释剂（不配位）和 LiNO₃ 不计；摩尔比精确，mol/L 按“溶剂体积 ≈ 溶液体积”粗算；含配位能力不明的溶剂不进尺子。稀释剂不进溶剂化层，但会减少每个 Li⁺ 分到的溶剂，所以“盐固定、稀释剂顶替溶剂”时尺子变小。尺子只代表 Li⁺ 一侧，多硫化物一侧看实测溶解度。"],
    ["方向怎么判", "|A−B| ÷ max(A,B) < 5% 记差异小；库仑效率写百分点差；只有作者文字记“作者定性”；≈、>、< 按区间比，重叠记无法判断。"],
    ["主存储", "json/ 里每篇一个 JSON；本 Excel 由 scripts/step6h_claim_records_v4.py 生成。改数据改 JSON 再重新生成。"],
]

# ---------- 10. 写 Excel ----------
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
USER_FILL = PatternFill("solid", fgColor="FFF2CC")
WIDE = {"说明": 120, "一句话结论": 140, "作者原名": 30, "溶剂": 34, "盐": 20, "添加剂": 14, "🧮 换算": 44, "电解液类别": 30,
        "每个主盐 Li⁺ 分到的配位溶剂分子数 🧮": 26, "原句": 60, "测试名": 22, "电池形式": 20, "锂厚 µm（N/P 🧮）": 30,
        "倍率及化成": 34, "LiNO₃": 18, "条件原句": 70, "测试": 24, "指标（取值点）": 30, "方向": 22, "出处": 26,
        "各档值（按梯度顺序，溶剂化由强到弱）": 34, "趋势": 14, "比较": 18, "性质变化": 26, "过程（大类·怎么变）": 24,
        "落到（测试·表现）": 24, "说法概括（AI）": 44, "整句原文": 70, "这句提到的大类": 18, "来源": 16, "在哪": 22,
        "内容": 36, "为什么要看": 40, "原句或依据": 60, "程序怎么处理": 44, "依据": 60, "全称": 30, "配位依据": 40,
        "物性来源": 22, "错提 ✍": 30, "漏提 ✍": 30, "你的判断 ✍": 30}


def add_sheet(wb, title, head, rows):
    ws = wb.create_sheet(title)
    ws.append(head)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = USER_FILL if "✍" in h else HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = WIDE.get(h, 12)
        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if "✍" in h:
                cell.fill = USER_FILL
    ws.freeze_panes = "B2"


wb = Workbook()
wb.remove(wb.active)
add_sheet(wb, "说明", ["项目", "说明"], guide)
add_sheet(wb, SHEET["结论"], ["文献", "一句话结论", "要你看（条）", "错提 ✍", "漏提 ✍"], pages["结论"])
add_sheet(wb, SHEET["要看"], ["序号", "文献", "在哪", "内容", "为什么要看", "原句或依据", "你的判断 ✍"], pages["要看"])
add_sheet(wb, SHEET["配方"], ["文献", "代号", "作者原名", "角色", "梯度值", "溶剂", "盐", "添加剂", "🧮 换算", "电解液类别",
                             "每个主盐 Li⁺ 分到的配位溶剂分子数 🧮", "原句"], pages["配方"])
add_sheet(wb, SHEET["测试"], ["文献", "测试", "测试名", "电池形式", "硫载量 mg/cm²", "E/S", "锂厚 µm（N/P 🧮）", "倍率及化成", "温度",
                             "LiNO₃", "参与配方", "条件原句"], pages["测试"])
add_sheet(wb, SHEET["方向"], ["文献", "测试", "比较", "类型", "指标（取值点）", "A 值", "对照值", "方向", "出处"], pages["方向"])
add_sheet(wb, SHEET["梯度"], ["文献", "测试", "指标（取值点）", "各档值（按梯度顺序，溶剂化由强到弱）", "趋势", "最优档", "出处"], pages["梯度"])
add_sheet(wb, SHEET["机制"], ["文献", "说法", "比较", "性质变化", "过程（大类·怎么变）", "落到（测试·表现）", "说法概括（AI）",
                             "原句编号", "整句原文", "这句提到的大类", "来源"], pages["机制"])
add_sheet(wb, SHEET["溶剂"], ["简称", "全称", "配位能力", "配位依据", "分子量", "密度 g/mL", "物性来源", "DN（论文值）",
                             "介电常数（论文值）", "来源"], pages["溶剂"])
add_sheet(wb, SHEET["记录"], ["文献", "在哪", "内容", "程序怎么处理", "依据"], pages["记录"])
wb.save(OUT_FILE)

print("写出：", OUT_FILE.relative_to(ROOT))
print({SHEET[k]: len(v) for k, v in pages.items()})
print(f"机制原句 {stats['原句']}，找到 {stats['原句找到']}；正文数值 {stats['数值']}，在原句里 {stats['数值在原句']}，反算 {stats['反算']}")
print("要你看：", sum(1 for r in pages["要看"] if r[0] != "—"))
